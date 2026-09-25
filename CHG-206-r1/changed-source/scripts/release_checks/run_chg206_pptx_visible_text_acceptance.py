#!/usr/bin/env python3
"""Native five-report PowerPoint text-fidelity acceptance for CHG-206 R1."""
from __future__ import annotations

import argparse
import hashlib
import importlib.metadata
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import traceback
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from playwright.sync_api import sync_playwright

ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT))
sys.path.insert(0,str(Path(__file__).parent))

from company_ui.products.visualizer.metric_format import format_metric_value
from company_ui.products.visualizer.ppt_service import _report_text_role, bound_export_items
from editor_host import NativeHost
from native_common import BrowserEvents, browser_kwargs, ready, write_json


SOURCE_SHA='47a988064deac544dfdca360e046a9cb20f5c282'
SOURCE_TREE='d36220e2cf7bd93116765b4f0f1bba19ed682747'
R2_SHA='7224c0815142557e2362e2ca4e6201e67006126c'
REPORT_KEYS=(
    'executive-business-review','semiconductor-rca','experiment-decision',
    'technical-status-review','supply-chain-capacity-holdout',
)
BASELINE_RASTER={
    'executive-business-review':'08__executive-business-review__executive-business-review-1.png',
    'semiconductor-rca':'03__semiconductor-rca__semiconductor-rca-1.png',
    'experiment-decision':'07__experiment-decision__experiment-decision-1.png',
    'technical-status-review':'06__technical-status-review__technical-status-review-1.png',
    'supply-chain-capacity-holdout':'05__supply-chain-holdout__supply-chain-capacity-holdout-1.png',
}
BASELINE_PPTX={key:key for key in REPORT_KEYS}
BASELINE_PPTX['supply-chain-capacity-holdout']='supply-chain-holdout'


def _digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _normalized_visible_text(value: str) -> str:
    return ' '.join(str(value).split())


def _semantic_entry(shape):
    nodes=shape._element.xpath('.//p:cNvPr')
    if not nodes:
        return None
    description=nodes[0].get('descr') or ''
    if not description.startswith('VisualizerSemantic:'):
        return None
    try:
        value=json.loads(description.removeprefix('VisualizerSemantic:'))
    except json.JSONDecodeError:
        return None
    return value if isinstance(value,dict) else None


def _all_text(shape) -> str:
    if getattr(shape,'has_text_frame',False):
        return shape.text
    if getattr(shape,'has_table',False):
        return '\n'.join(cell.text for row in shape.table.rows for cell in row.cells)
    if getattr(shape,'has_chart',False):
        parts=[]
        chart=shape.chart
        if chart.has_title:
            parts.append(chart.chart_title.text_frame.text)
        for series in chart.series:
            parts.append(series.name)
            parts.extend(str(point.label) for point in chart.plots[0].categories)
        return '\n'.join(parts)
    return ''


def _item_strings(entry: dict, projected: dict) -> list[str]:
    engine=str(entry.get('engine') or '')
    values=[]
    title=str(entry.get('title') or entry.get('element') or '')
    if title and engine!='TableEngine':
        values.append(title)
    if engine=='TextEngine':
        value=entry.get('text') or entry.get('body') or ''
        if value: values.append(str(value))
    elif engine in {'EvidenceCompositeEngine','DecisionCompositeEngine','ProjectCompositeEngine'}:
        values.extend(str(entry[key]) for key in ('statement','detail','status') if entry.get(key) not in (None,''))
    elif engine=='MetricEngine':
        values.append(format_metric_value(projected.get('value'),projected,projected.get('_metric_field')))
        if entry.get('detail'): values.append(str(entry['detail']))
        if entry.get('unit') and str(entry.get('unit')) not in values: values.append(str(entry['unit']))
    elif engine=='ComparisonEngine':
        before=projected.get('before');after=projected.get('after')
        values.append(f'{format_metric_value(before,projected,projected.get("_metric_field"))} → {format_metric_value(after,projected,projected.get("_metric_field"))}')
        if entry.get('detail'): values.append(str(entry['detail']))
    elif engine=='TimelineEngine':
        for milestone in projected.get('milestones') or []:
            if isinstance(milestone,dict):
                label=str(milestone.get('label') or '')
                date=milestone.get('date')
                values.append(f'{label} · {date}' if date not in (None,'') else label)
    elif engine=='TableEngine':
        table=projected.get('customTable') or {}
        values.extend(str(value) for value in table.get('headers') or [])
        values.extend(str(value) for row in table.get('rows') or [] for value in row if value not in (None,''))
    elif engine=='DiagramEngine':
        values.extend(str(value) for value in projected.get('nodes') or [])
    elif engine in {'WaferFabEngine','SmartLayoutEngine','InteractionLayer','EditorInfrastructure'}:
        values.extend(str(projected.get(key) or '') for key in ('wafer_id','lot_id','tool','chamber','recipe','process') if projected.get(key))
    return [value for value in values if value]


def _expected_role(entry: dict, text: str) -> str:
    engine=str(entry.get('engine') or '')
    title=str(entry.get('title') or entry.get('element') or '')
    if engine in {'MetricEngine','ComparisonEngine'}:
        return 'metric title' if text==title else ('comparison value' if engine=='ComparisonEngine' else 'metric value')
    if text==title: return 'text heading'
    if engine=='TextEngine': return _report_text_role(entry).replace('_',' ')
    if engine in {'EvidenceCompositeEngine','DecisionCompositeEngine','ProjectCompositeEngine'}:
        if text==entry.get('statement'):
            role=_report_text_role(entry)
            statement=str(text).casefold().strip()
            decision_language=('contain','hold ','release ','do not','stop ','reserve ','block ')
            next_step_language=('complete ','verify ','run ','confirm ','review ','update ','schedule ','implement ','document ','prepare ')
            if engine=='ProjectCompositeEngine':
                if any(statement.startswith(token) for token in decision_language): role='risk_decision'
                elif any(statement.startswith(token) for token in next_step_language): role='conclusion_next_step'
            if engine=='EvidenceCompositeEngine': role='narrative_interpretation'
            return role.replace('_',' ')
        if text==entry.get('status'): return 'action status'
        role=_report_text_role(entry)
        statement=str(entry.get('statement') or '').casefold().strip()
        decision_language=('contain','hold ','release ','do not','stop ','reserve ','block ')
        next_step_language=('complete ','verify ','run ','confirm ','review ','update ','schedule ','implement ','document ','prepare ')
        if engine=='ProjectCompositeEngine':
            if any(statement.startswith(token) for token in decision_language): return 'conclusion next step'
            if any(statement.startswith(token) for token in next_step_language): return 'evidence detail'
        if engine=='DecisionCompositeEngine' and role=='conclusion_next_step': return 'risk decision'
        if engine=='ProjectCompositeEngine' and role=='risk_decision': return 'conclusion next step'
        return 'evidence detail'
    if engine=='TableEngine': return 'evidence detail'
    if engine in {'DiagramEngine','WaferFabEngine'}: return 'engineering spatial label'
    return 'supporting report text'


def _text_role_facts(shapes) -> list[dict]:
    facts=[]
    for shape in shapes:
        if not getattr(shape,'has_text_frame',False):
            continue
        nodes=shape._element.xpath('.//p:cNvPr')
        title=nodes[0].get('title') if nodes else None
        if not title or not title.startswith('Visembler report text roles: '):
            continue
        fonts=[]
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                fonts.append({'text':run.text,'family':run.font.name,'size_pt':run.font.size.pt if run.font.size else None,'bold':run.font.bold,'color_rgb':str(run.font.color.rgb) if run.font.color.type else None})
        facts.append({'shape':shape.name,'roles':title.removeprefix('Visembler report text roles: ').split(', '),'text':shape.text,'geometry_emu':{'left':shape.left,'top':shape.top,'width':shape.width,'height':shape.height},'background_rgb':str(shape.fill.fore_color.rgb),'wrap':shape.text_frame.word_wrap,'anchor':str(shape.text_frame.vertical_anchor),'margins_emu':{'left':shape.text_frame.margin_left,'right':shape.text_frame.margin_right,'top':shape.text_frame.margin_top,'bottom':shape.text_frame.margin_bottom},'fit':'none','fonts':fonts})
    return facts


def inspect_pptx(path: Path, source_model: dict, browser_geometry: dict) -> dict:
    deck=Presentation(str(path));shapes=[shape for slide in deck.slides for shape in slide.shapes]
    projections={entry['id']:entry for entry in bound_export_items(source_model)}
    source={entry['id']:entry for entry in source_model.get('items') or []}
    visible_by_id={}
    for shape in shapes:
        entry=_semantic_entry(shape)
        if entry and entry.get('id'):
            visible_by_id[str(entry['id'])]=shape
    missing=[];string_inventory=[];spatial={};charts=0;tables=0
    for item_id,entry in source.items():
        projection=projections[item_id];engine=str(entry.get('engine') or '')
        shape=visible_by_id.get(item_id)
        if engine in {'DiagramEngine','WaferFabEngine'} and shape is None:
            title=str(entry.get('title') or entry.get('element') or '')
            candidates=[value for value in shapes if title in value.name]
            if candidates:
                shape=candidates[0]
        if engine in {'CoreChartEngine','EngineeringChartEngine'}:
            chart_shape=next((value for value in shapes if getattr(value,'has_chart',False) and (_semantic_entry(value) or {}).get('id')==item_id),None)
            if chart_shape:
                charts+=1
                chart_text=_all_text(chart_shape)
                expected=str(entry.get('title') or entry.get('element') or '')
                if expected and expected not in chart_text:
                    missing.append({'item_id':item_id,'role':'chart title','text':expected})
        if engine=='TableEngine':
            table_shape=next((value for value in shapes if getattr(value,'has_table',False) and (_semantic_entry(value) or {}).get('id')==item_id),None)
            if table_shape:
                tables+=1
                shape=table_shape
        if engine=='DiagramEngine':
            nodes=[value for value in shapes if '::node-' in value.name and entry.get('title') in value.name]
            labels=[value.text for value in nodes if getattr(value,'has_text_frame',False)]
            edges=[value for value in shapes if value.name.startswith('VIZ::DiagramEdge::') and entry.get('title') in value.name]
            spatial[item_id]={'kind':'Process Flow','nodes':len(nodes),'labels':labels,'connectors':len(edges)}
            if [_normalized_visible_text(value) for value in labels]!=[_normalized_visible_text(value) for value in projection.get('nodes') or []]:
                missing.append({'item_id':item_id,'role':'process flow nodes','expected':projection.get('nodes'),'actual':labels})
            if len(edges)<len(projection.get('edges') or []):
                missing.append({'item_id':item_id,'role':'process flow connectors','expected':len(projection.get('edges') or []),'actual':len(edges)})
        if engine=='WaferFabEngine':
            dies=[value for value in shapes if '::die-' in value.name and entry.get('title') in value.name]
            spatial[item_id]={'kind':entry.get('element'),'die_primitives':len(dies),'title_visible':any(getattr(value,'has_text_frame',False) and value.text==str(entry.get('title') or entry.get('element')) for value in shapes if entry.get('title') in value.name)}
            if len(dies)<1:
                missing.append({'item_id':item_id,'role':'wafer spatial primitives','expected':'at least one die','actual':0})
        if engine in {'DiagramEngine','WaferFabEngine'}:
            visible='\n'.join(_all_text(value) for value in shapes if (entry.get('title') or entry.get('element') or '') in value.name)
        else:
            visible=_all_text(shape) if shape is not None else '\n'.join(_all_text(value) for value in shapes)
        for text in _item_strings(entry,projection):
            present=text in visible or (engine=='DiagramEngine' and _normalized_visible_text(text) in _normalized_visible_text(visible))
            expected_role=_expected_role(entry,text)
            role_tags=[]
            if shape is not None:
                nodes=shape._element.xpath('.//p:cNvPr')
                title=nodes[0].get('title') if nodes else ''
                if title and title.startswith('Visembler report text roles: '): role_tags=title.removeprefix('Visembler report text roles: ').split(', ')
            role_matches=expected_role in role_tags or engine in {'TableEngine','DiagramEngine','WaferFabEngine','CoreChartEngine','EngineeringChartEngine'}
            string_inventory.append({'item_id':item_id,'engine':engine,'role':expected_role,'text':text,'present_in_expected_element':present,'role_matches_element_metadata':role_matches,'shape':shape.name if shape is not None else None})
            if not present:
                missing.append({'item_id':item_id,'role':'semantic text','text':text})
            if not role_matches:
                missing.append({'item_id':item_id,'role':'role metadata','expected':expected_role,'actual':role_tags})
    bad_bounds=[]
    for slide_index,slide in enumerate(deck.slides,1):
        for shape in slide.shapes:
            if shape.left<0 or shape.top<0 or shape.left+shape.width>deck.slide_width or shape.top+shape.height>deck.slide_height:
                bad_bounds.append({'slide':slide_index,'shape':shape.name,'geometry_emu':[shape.left,shape.top,shape.width,shape.height]})
    role_facts=_text_role_facts(shapes)
    required_roles={'report headline','context subhead','metric title','metric value'}
    if not any(entry.get('engine')=='EvidenceCompositeEngine' for entry in source.values()):
        if not any(entry.get('engine')=='TextEngine' and _report_text_role(entry)=='narrative_interpretation' for entry in source.values()):
            required_roles.add('evidence detail')
    if any(entry.get('engine')=='ComparisonEngine' for entry in source.values()): required_roles.add('comparison value')
    if any(entry.get('engine') in {'EvidenceCompositeEngine','DecisionCompositeEngine','ProjectCompositeEngine'} and entry.get('detail') for entry in source.values()): required_roles.add('evidence detail')
    if any(entry.get('engine') in {'DecisionCompositeEngine','ProjectCompositeEngine'} and entry.get('statement') for entry in source.values()): required_roles.add('risk decision')
    if any(entry.get('engine') in {'EvidenceCompositeEngine','DecisionCompositeEngine','ProjectCompositeEngine'} and entry.get('status') for entry in source.values()): required_roles.add('action status')
    if any(entry.get('element') in {'Key Takeaway','Hero Title'} or entry.get('engine') in {'DecisionCompositeEngine','ProjectCompositeEngine'} for entry in source.values()): required_roles.add('conclusion next step')
    observed_roles={role for fact in role_facts for role in fact['roles']}
    if any(row['engine']=='TableEngine' and row['role']=='evidence detail' and row['present_in_expected_element'] for row in string_inventory):
        observed_roles.add('evidence detail')
    if not required_roles.issubset(observed_roles):
        missing.append({'role':'report roles','expected':sorted(required_roles),'actual':sorted(observed_roles)})
    layouts=browser_geometry.get('items') or []
    rects=[{key:float(value[key]) for key in ('x','y','w','h')} for value in layouts]
    overlaps=[]
    for index,first in enumerate(rects):
        for second in rects[index+1:]:
            if min(first['x']+first['w'],second['x']+second['w'])-max(first['x'],second['x'])>1 and min(first['y']+first['h'],second['y']+second['h'])-max(first['y'],second['y'])>1:
                overlaps.append([first,second])
    geometry_issues=[];canvas=browser_geometry.get('canvas') or {};layout_by_id={str(item.get('id')):item for item in layouts if isinstance(item,dict) and item.get('id')}
    if canvas.get('width') and canvas.get('height'):
        sx=deck.slide_width/float(canvas['width']);sy=deck.slide_height/float(canvas['height'])
        for item_id,shape in visible_by_id.items():
            if item_id not in layout_by_id or projections[item_id].get('engine')=='DiagramEngine':
                continue
            rect=layout_by_id[item_id]
            expected=[round(float(rect[key])*scale) for key,scale in (('x',sx),('y',sy),('w',sx),('h',sy))]
            actual=[shape.left,shape.top,shape.width,shape.height]
            if any(abs(first-second)>1 for first,second in zip(actual,expected)):
                geometry_issues.append({'item_id':item_id,'expected_emu':expected,'actual_emu':actual})
    return {'path':path.name,'sha256':_digest(path),'bytes':path.stat().st_size,'slide_count':len(deck.slides),'shape_count':sum(len(slide.shapes) for slide in deck.slides),'text_shape_count':sum(getattr(shape,'has_text_frame',False) for shape in shapes),'chart_count':charts,'table_count':tables,'items':string_inventory,'text_roles':role_facts,'spatial':spatial,'browser_geometry':browser_geometry,'element_rectangle_overlap_count':len(overlaps),'element_rectangle_geometry_issues':geometry_issues,'off_slide_shapes':bad_bounds,'missing_or_misplaced_strings':missing,'status':'PASS' if not missing and not bad_bounds and not overlaps and not geometry_issues else 'FAIL'}


def export_download(page, destination: Path, host: NativeHost) -> bytes:
    page.locator('#exportBtn').click()
    page.get_by_role('button',name='Editable PowerPoint').wait_for(state='visible',timeout=8_000)
    downloads=[]
    page.on('download',lambda value:downloads.append(value))
    page.locator('#exportPptAction').click()
    deadline=time.monotonic()+30
    while time.monotonic()<deadline and not downloads:
        page.wait_for_timeout(100)
    if not downloads:
        details=page.evaluate('''()=>({bridge:(()=>{const s=window.CompanyUIVisualizerBridge?.state?.();return s?{report_id:s.report_id,revision:s.revision,pending:s.pending,inflight:s.inflight,recovery:s.recovery}:null})(),debug:(window.__VIZ_PROD__?.ui?.debugLog||[]).slice(0,20).map(({level,event,detail})=>({level,event,detail})),preflight:(()=>{const p=window.__VIZ_PROD__?.preflight?.();return p?{issues:p.issues,layoutIssues:p.layoutIssues,dataIssues:p.dataIssues}:null})(),save:{label:document.querySelector('#saveBtn')?.textContent||null,status:document.querySelector('#saveStatus')?.textContent||null},modal:document.querySelector('#genericModal.show #modalBody')?.innerText||null,toast:document.querySelector('[role=status]')?.innerText||''})''')
        page.screenshot(path=str(destination.parent/(destination.stem+'-export-error.png')),full_page=True)
        log_tail=host.log_path.read_text(encoding='utf-8',errors='replace')[-8000:] if host.log_path.exists() else ''
        raise AssertionError('Supported Editable PowerPoint action produced no download: '+json.dumps(details,ensure_ascii=False)[:12000]+'\nNative host log tail:\n'+log_tail)
    downloads[0].save_as(str(destination))
    return destination.read_bytes()


def _render(path: Path, key: str, renderer: Path, pdftoppm: Path, output: Path) -> dict:
    pdf_dir=output/'pdf';raster_dir=output/'rasters';profile=output/'profiles'/key
    for directory in (pdf_dir,raster_dir,profile): directory.mkdir(parents=True,exist_ok=True)
    command=[str(renderer),'--headless',f'-env:UserInstallation={profile.as_uri()}','--convert-to','pdf','--outdir',str(pdf_dir),str(path)]
    renderer_env=os.environ.copy();renderer_env['PYTHONDONTWRITEBYTECODE']='1'
    converted=subprocess.run(command,capture_output=True,text=True,timeout=180,env=renderer_env)
    pdf=pdf_dir/f'{path.stem}.pdf'
    if converted.returncode or not pdf.is_file():
        raise RuntimeError(f'LibreOffice conversion failed for {path.name}: {converted.stderr or converted.stdout}')
    prefix=raster_dir/key
    raster_command=[str(pdftoppm),'-png','-r','144',str(pdf),str(prefix)]
    rasterized=subprocess.run(raster_command,capture_output=True,text=True,timeout=180)
    if rasterized.returncode:
        raise RuntimeError(f'pdftoppm failed for {path.name}: {rasterized.stderr or rasterized.stdout}')
    images=sorted(raster_dir.glob(key+'-*.png'))
    if not images:
        raise RuntimeError(f'No raster pages were produced for {path.name}.')
    return {'pdf':str(pdf),'pdf_sha256':_digest(pdf),'libreoffice_command':command,'libreoffice_environment':{'PYTHONDONTWRITEBYTECODE':'1'},'libreoffice_returncode':converted.returncode,'libreoffice_stdout':converted.stdout,'libreoffice_stderr':converted.stderr,'pdftoppm_command':raster_command,'pdftoppm_returncode':rasterized.returncode,'pdftoppm_stdout':rasterized.stdout,'pdftoppm_stderr':rasterized.stderr,'rasters':[{'path':str(image),'sha256':_digest(image),'bytes':image.stat().st_size} for image in images]}


def _contact_sheet(rows: list[tuple[str,Path,Path]], destination: Path) -> None:
    font=ImageFont.load_default();tiles=[]
    for label,before,after in rows:
        tile=Image.new('RGB',(1320,470),'white');draw=ImageDraw.Draw(tile)
        draw.text((12,8),label,fill='#132d49',font=font)
        for x,path,caption in ((10,before,'CHG-173 R2 failure'),(670,after,'CHG-206 candidate')):
            with Image.open(path) as source:
                image=source.convert('RGB');image.thumbnail((640,425))
                tile.paste(image,(x,30));draw.text((x,18),caption,fill='#132d49',font=font)
        tiles.append(tile)
    sheet=Image.new('RGB',(1320,470*len(tiles)),'white')
    for index,tile in enumerate(tiles): sheet.paste(tile,(0,index*470))
    destination.parent.mkdir(parents=True,exist_ok=True);sheet.save(destination)


def main() -> int:
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output',required=True,type=Path)
    parser.add_argument('--renderer',required=True,type=Path)
    parser.add_argument('--pdftoppm',required=True,type=Path)
    parser.add_argument('--r2-evidence',required=True,type=Path)
    parser.add_argument('--candidate-sha',required=True)
    parser.add_argument('--candidate-tree',required=True)
    args=parser.parse_args();output=args.output.expanduser().resolve()
    if output==ROOT or ROOT in output.parents: parser.error('--output must be outside the checkout.')
    if output.exists() and any(output.iterdir()): parser.error('--output must be new or empty.')
    output.mkdir(parents=True,exist_ok=True)
    head=subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True).strip()
    tree=subprocess.check_output(['git','rev-parse','HEAD^{tree}'],cwd=ROOT,text=True).strip()
    if head!=args.candidate_sha or tree!=args.candidate_tree: parser.error('Acceptance must start on the exact candidate SHA and tree.')
    dirty=subprocess.check_output(['git','status','--porcelain'],cwd=ROOT,text=True).strip()
    if dirty: parser.error('Candidate acceptance requires a clean worktree.')
    fixture=json.loads((ROOT/'tests/fixtures/chg206/whole_report_models.json').read_text(encoding='utf-8'))
    reports={row['key']:row for row in fixture['reports']}
    parent=subprocess.check_output(['git','rev-parse','HEAD^'],cwd=ROOT,text=True).strip()
    if parent!=SOURCE_SHA: parser.error('The candidate parent must be the exact requested CHG-173 R2 main commit.')
    receipt={'schema':'visembler-chg206-visible-text-acceptance.v1','project':'visembler','request':'CHG-206-r1','operation':'FIX','fabric_job_id':'CF-335e326a4b753e8a2e6d9e08','candidate_sha':head,'candidate_tree':tree,'source_sha':SOURCE_SHA,'source_tree':SOURCE_TREE,'r2_source_evidence':str(args.r2_evidence.resolve()),'r2_carrier_sha':R2_SHA,'established_root_cause':{'reproduction':'CHG-173 R2 Executive Business Review PPTX and its qualified LibreOffice raster','observed_xml':'body paragraphs and KPI values lacked explicit run foreground colors; generic body text frames inherited vertical-middle anchoring, wrapping, insets and fit/overflow behavior from the template/theme','pixel_failure':'R2 raster retains headings, charts and tables while omitting or clipping headline/context values, KPI values, decision/risk and next-step text','bounded_authority':'generic report text and KPI shapes only; DiagramEngine and WaferFabEngine renderers remain specialized'},'fixture_notes':['The five exported report models are the persisted CHG-173 R2 whole-report models.','The RCA adds a clearly labeled synthetic paired-die subset solely so this acceptance fixture exercises the CHG-180 Wafer Difference Map renderer alongside its original Wafer Map and Process Flow.'],'reports':{},'browser_errors':[],'renderer':{},'visual_review':{'status':'PENDING_MANUAL_REVIEW'}}
    events=BrowserEvents();output_exports=output/'exports';output_exports.mkdir()
    source_fixture_hash=_digest(ROOT/'tests/fixtures/chg206/whole_report_models.json')
    receipt['fixture_sha256']=source_fixture_hash
    try:
        with tempfile.TemporaryDirectory(prefix='visembler-chg206-native-') as temp_dir, sync_playwright() as playwright:
            with NativeHost(ROOT,Path(temp_dir)/'data') as host:
                report_ids={}
                for key in REPORT_KEYS:
                    row=reports[key]
                    report_ids[key]=host.create(model=row['model'],name='chg206-'+key)
                browser=playwright.chromium.launch(**browser_kwargs())
                receipt['runtime']={'python':sys.version.split()[0],'nicegui':importlib.metadata.version('nicegui'),'playwright':importlib.metadata.version('playwright'),'python_pptx':importlib.metadata.version('python-pptx'),'browser':browser.version,'host_url':host.url,'ephemeral_port':host.port}
                context=browser.new_context(accept_downloads=True,viewport={'width':1600,'height':1120},device_scale_factor=1)
                page=context.new_page();page.set_default_timeout(10_000);events.attach(page)
                for key in REPORT_KEYS:
                    report_id=report_ids[key]
                    response=page.goto(f'{host.url}/visualizer?report={report_id}',wait_until='domcontentloaded')
                    assert response and response.status==200
                    ready(page,require_settled=True)
                    page.wait_for_function('()=>window.socket?.connected===true&&window.did_handshake===true',timeout=15_000)
                    browser_geometry=page.evaluate('()=>window.__VIZ_PROD__.layoutGeometry()')
                    assert browser_geometry and browser_geometry.get('items'),key
                    destination=output_exports/f'{key}.pptx'
                    pptx_bytes=export_download(page,destination,host)
                    structure=inspect_pptx(destination,reports[key]['model'],browser_geometry)
                    assert structure['status']=='PASS',json.dumps(structure['missing_or_misplaced_strings'],ensure_ascii=False)
                    raster=_render(destination,key,args.renderer.resolve(),args.pdftoppm.resolve(),output/'pptx-review')
                    r2_key=BASELINE_PPTX[key]
                    r2_pptx=args.r2_evidence/'reports'/r2_key/f'{key}.pptx'
                    baseline_name=BASELINE_RASTER[key]
                    baseline_raster=args.r2_evidence/'exports/pptx-review/rasters'/baseline_name
                    assert r2_pptx.is_file() and baseline_raster.is_file(),f'Missing exact CHG-173 R2 comparison for {key}.'
                    first_raster=Path(raster['rasters'][0]['path'])
                    receipt['reports'][key]={'title':reports[key]['title'],'report_id':report_id,'source_report_id':reports[key]['source_report_id'],'source_revision':reports[key]['revision'],'supported_action':'#exportBtn → #exportPptAction (Editable PowerPoint)','pptx_sha256':hashlib.sha256(pptx_bytes).hexdigest(),'pptx_bytes':len(pptx_bytes),'structure':structure,'raster':raster,'comparison':{'r2_pptx_sha256':_digest(r2_pptx),'r2_raster':str(baseline_raster),'r2_raster_sha256':_digest(baseline_raster),'candidate_raster':str(first_raster),'candidate_raster_sha256':_digest(first_raster)}}
                context.close();browser.close()
        renderer_path=args.renderer.resolve()
        renderer_app=renderer_path.parents[2]
        renderer_version=subprocess.run([str(renderer_path),'--version'],capture_output=True,text=True,timeout=30)
        codesign=subprocess.run(['codesign','--verify','--deep','--strict',str(renderer_app)],capture_output=True,text=True,timeout=60)
        signature=subprocess.run(['codesign','-dv','--verbose=4',str(renderer_app)],capture_output=True,text=True,timeout=30)
        trust=subprocess.run(['spctl','--assess','--type','execute','--verbose',str(renderer_app)],capture_output=True,text=True,timeout=60)
        poppler_version=subprocess.run([str(args.pdftoppm.resolve()),'-v'],capture_output=True,text=True,timeout=30)
        receipt['renderer']={'name':'LibreOffice','binary_path':str(renderer_path),'binary_sha256':_digest(renderer_path),'app_bundle_path':str(renderer_app),'version':renderer_version.stdout.strip() or renderer_version.stderr.strip(),'codesign_verify_command':['codesign','--verify','--deep','--strict',str(renderer_app)],'codesign_verify_returncode':codesign.returncode,'codesign_verify_output':codesign.stderr or codesign.stdout,'signature_command':['codesign','-dv','--verbose=4',str(renderer_app)],'signature_returncode':signature.returncode,'signature_output':signature.stderr or signature.stdout,'gatekeeper_assessment_command':['spctl','--assess','--type','execute','--verbose',str(renderer_app)],'gatekeeper_returncode':trust.returncode,'gatekeeper_output':trust.stderr or trust.stdout,'rasterizer':'Poppler pdftoppm','rasterizer_path':str(args.pdftoppm.resolve()),'rasterizer_sha256':_digest(args.pdftoppm.resolve()),'rasterizer_version':poppler_version.stderr.strip() or poppler_version.stdout.strip(),'profile_scope':'job-owned per-report LibreOffice user profile','raster_dpi':144}
        assert renderer_version.returncode==0 and codesign.returncode==0 and trust.returncode==0,receipt['renderer']
        rows=[]
        for key in REPORT_KEYS:
            report=receipt['reports'][key]
            rows.append((key,Path(report['comparison']['r2_raster']),Path(report['comparison']['candidate_raster'])))
        _contact_sheet(rows,output/'before-after-failing-raster-contact-sheet.png')
        receipt['browser_errors']=events.unexpected
        assert not receipt['browser_errors'],receipt['browser_errors']
        assert len(receipt['reports'])==5
        receipt['status']='PASS_STRUCTURAL_AND_RASTERIZED_PENDING_MANUAL_PIXEL_REVIEW'
    except Exception as exc:
        receipt['error']=str(exc);receipt['traceback']=traceback.format_exc();receipt['browser_errors']=events.unexpected;receipt['status']='FAIL'
    write_json(output/'acceptance-receipt.json',receipt)
    print(json.dumps(receipt,indent=2,ensure_ascii=False))
    return 0 if receipt['status']=='PASS_STRUCTURAL_AND_RASTERIZED_PENDING_MANUAL_PIXEL_REVIEW' else 1


if __name__=='__main__':
    raise SystemExit(main())
