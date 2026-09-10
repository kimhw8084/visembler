from __future__ import annotations

import io
import json
import math
import os
import shutil
import subprocess
import tempfile
import time
from pathlib import Path

import pytest
from pptx import Presentation

from company_ui.data_engine import DataQuery, FilterClause, FilterOperation
from company_ui.products.visualizer.dataset_resources import DatasetResourceStore, ScopedDatasetRepository
from company_ui.products.visualizer.domain import ReportNotFoundError, VisualizerContractError, canonical_model
from company_ui.products.visualizer.governance import ReportAccessCatalog, ScopedReportRepository
from company_ui.products.visualizer.ppt_service import export_pptx
from company_ui.products.visualizer.repository import ReportRepository
from company_ui.products.visualizer.statistical_analysis import analyze_statistical_items, require_valid_statistical_results
from company_ui.security import AuthorizationModel, Principal


ROOT = Path(__file__).resolve().parents[1]


def node_json(source: str):
    result = subprocess.run(['node', '--input-type=module', '-e', source], cwd=ROOT, check=True, text=True, capture_output=True)
    return json.loads(result.stdout)


def _cap_fields():
    return [
        {'id': 'measurement', 'name': 'Measurement', 'type': 'number'},
        {'id': 'lsl', 'name': 'LSL', 'type': 'number'},
        {'id': 'usl', 'name': 'USL', 'type': 'number'},
    ]


def _cap_item(item_id='capability'):
    mapping = {'value': 'measurement', 'specification_low': 'lsl', 'specification_high': 'usl'}
    return {'id': item_id, 'engine': 'MetricEngine', 'type': 'metric', 'mapping': mapping, 'analysis_recipe': {'id': 'process-capability', 'mapping': mapping}}


def _analyze(fields, rows, item, *, source_total=None, filtered_total=None, revision=1, dataset_id='d1'):
    return analyze_statistical_items(
        report_id='report-1', dataset_id=dataset_id, resource_id='resource-1', revision=revision,
        fields=fields, rows=rows, items=[item], session_id='session-1',
        source_total=len(rows) if source_total is None else source_total,
        filtered_total=len(rows) if filtered_total is None else filtered_total,
    )[item['id']]


def test_resource_capability_uses_full_population_and_records_provenance(tmp_path: Path):
    fields = _cap_fields()
    rows = [[1, 0, 10] for _ in range(250)] + [[10, 0, 10] for _ in range(250)]
    store = DatasetResourceStore(tmp_path)
    resource = store.create(owner='alice', name='full capability', fields=fields, rows=rows)
    result = _analyze(fields, [[row[field['id']] for field in fields] for row in store.session(resource['dataset_id']).query(DataQuery(limit=None)).rows], _cap_item(), source_total=500, filtered_total=500)
    assert result['ok'] is True
    stats = result['derived_statistics']['stats']
    sigma = math.sqrt(10125 / 499)
    assert stats['n'] == 500
    assert stats['mean'] == pytest.approx(5.5)
    assert stats['cpk'] == pytest.approx(4.5 / (3 * sigma))
    assert result['population'] == {'source_total': 500, 'filtered_total': 500, 'analyzed_rows': 500, 'complete': True}
    assert result['source']['dataset_revision'] == resource['revision']
    assert len(result['renderer_ready']['dataset']['rows']) == 1
    assert 250 != stats['n']


def test_filtered_resource_over_10000_is_not_silently_truncated(tmp_path: Path):
    fields = _cap_fields() + [{'id': 'population', 'name': 'Population', 'type': 'categorical'}]
    rows = [[9.8 + (index % 3) * 0.2, 9, 11, 'target'] for index in range(12000)] + [[10.2, 9, 11, 'other'] for _ in range(50)]
    store = DatasetResourceStore(tmp_path)
    resource = store.create(owner='alice', name='large filtered capability', fields=fields, rows=rows)
    session = store.session(resource['dataset_id'])
    with session.transaction():
        session.set_filter(FilterClause('population', FilterOperation.EQUALS, 'target'))
    preview = session.query(DataQuery(limit=10000))
    complete = session.query(DataQuery(limit=None))
    assert preview.filtered_total == 12000 and len(preview.rows) == 10000
    assert complete.filtered_total == 12000 and len(complete.rows) == 12000
    result = _analyze(fields, [[row['measurement'], row['lsl'], row['usl'], row['population']] for row in complete.rows], _cap_item('filtered-capability'), source_total=12050, filtered_total=complete.filtered_total)
    assert result['ok'] is True
    assert result['population']['filtered_total'] == 12000
    assert result['population']['analyzed_rows'] == 12000
    assert result['population']['complete'] is True
    assert result['derived_statistics']['stats']['n'] == 12000


def test_resource_xbar_and_doe_use_exact_complete_rows():
    x_fields = [{'id': 'subgroup', 'name': 'Subgroup', 'type': 'categorical'}, {'id': 'measurement', 'name': 'Measurement', 'type': 'number'}]
    x_rows = [[f'G{i}', i + offset] for i in range(1, 61) for offset in range(5)]
    x_item = {'id': 'xbar', 'engine': 'EngineeringChartEngine', 'type': 'chart', 'element': 'Xbar-R Chart', 'analysis_recipe': {'id': 'xbar-r-process-review', 'mapping': {'subgroup': 'subgroup', 'value': 'measurement'}}}
    x_result = _analyze(x_fields, x_rows, x_item, source_total=300, filtered_total=300)
    assert x_result['ok'] is True
    assert x_result['population'] == {'source_total': 300, 'filtered_total': 300, 'analyzed_rows': 300, 'complete': True}
    assert x_result['derived_statistics']['stats']['n'] == 5
    assert x_result['derived_statistics']['stats']['means'][:3] == [3, 4, 5]
    assert x_result['derived_statistics']['stats']['ranges'] == [4] * 60

    d_fields = [{'id': 'factor_a', 'name': 'Factor A', 'type': 'categorical'}, {'id': 'factor_b', 'name': 'Factor B', 'type': 'categorical'}, {'id': 'response', 'name': 'Response', 'type': 'number'}]
    d_rows = [[a, b, base + index * 0.1] for a, base in [('A', 10), ('B', 30)] for b in ('L', 'H') for index in range(100)]
    d_item = {'id': 'doe', 'engine': 'EngineeringChartEngine', 'type': 'chart', 'element': 'DOE Interaction Plot', 'analysis_recipe': {'id': 'doe-response-review', 'mapping': {'factor_a': 'factor_a', 'factor_b': 'factor_b', 'response': 'response'}}}
    d_result = _analyze(d_fields, d_rows, d_item, source_total=400, filtered_total=400)
    assert d_result['ok'] is True
    assert d_result['population'] == {'source_total': 400, 'filtered_total': 400, 'analyzed_rows': 400, 'complete': True}
    interaction = d_result['derived_statistics']['interaction']
    assert len(interaction['cells']) == 2 and all(len(row) == 2 for row in interaction['cells'])
    assert interaction['cells'][0][0]['mean'] == pytest.approx(14.95)
    assert interaction['cells'][1][1]['mean'] == pytest.approx(34.95)


@pytest.mark.parametrize('size', [1000, 10000, 100000])
def test_supported_resource_envelope_is_bounded_and_characterized(size: int):
    fields = _cap_fields()
    rows = [[9.8 + (index % 3) * 0.2, 9, 11] for index in range(size)]
    started = time.perf_counter()
    result = _analyze(fields, rows, _cap_item(f'cap-{size}'), source_total=size, filtered_total=size)
    elapsed = time.perf_counter() - started
    assert result['ok'] is True
    assert result['population']['analyzed_rows'] == size
    assert result['population']['complete'] is True
    # The response is a governed summary, not a raw-row transport.
    assert len(json.dumps(result, separators=(',', ':'))) < 25000
    assert elapsed < 60


def test_invalid_refresh_preflight_is_atomic_for_xbar_capability_and_doe(tmp_path: Path):
    store = DatasetResourceStore(tmp_path)

    cap = store.create(owner='alice', name='cap', fields=_cap_fields(), rows=[[9.8, 9, 11], [10.2, 9, 11]])
    bad_cap_rows = [[9.8, 9, 11], [10.2, 9.5, 11]]
    bad_cap = _analyze(_cap_fields(), bad_cap_rows, _cap_item(), revision=2, dataset_id=cap['dataset_id'])
    with pytest.raises(VisualizerContractError, match='conflicting values'):
        require_valid_statistical_results({'cap': bad_cap})
    assert store.get(cap['dataset_id'])['revision'] == 1

    x_fields = [{'id': 'subgroup', 'name': 'Subgroup', 'type': 'categorical'}, {'id': 'measurement', 'name': 'Measurement', 'type': 'number'}]
    x = store.create(owner='alice', name='xbar', fields=x_fields, rows=[['A', 1], ['A', 2], ['B', 3], ['B', 4]])
    bad_x = _analyze(x_fields, [['A', 1], ['A', 2], ['B', 3], ['B', 4], ['B', 5]], {'id': 'x', 'analysis_recipe': {'id': 'xbar-r-process-review', 'mapping': {'subgroup': 'subgroup', 'value': 'measurement'}}}, revision=2, dataset_id=x['dataset_id'])
    with pytest.raises(VisualizerContractError, match='constant subgroup size'):
        require_valid_statistical_results({'x': bad_x})
    assert store.get(x['dataset_id'])['revision'] == 1

    d_fields = [{'id': 'factor_a', 'name': 'Factor A', 'type': 'categorical'}, {'id': 'factor_b', 'name': 'Factor B', 'type': 'categorical'}, {'id': 'response', 'name': 'Response', 'type': 'number'}]
    d = store.create(owner='alice', name='doe', fields=d_fields, rows=[['A', 'L', 1], ['A', 'H', 2], ['B', 'L', 3], ['B', 'H', 4]])
    bad_d = _analyze(d_fields, [['A', 'L', 1], ['A', 'H', 2], ['B', 'L', 3]], {'id': 'd', 'analysis_recipe': {'id': 'doe-response-review', 'mapping': {'factor_a': 'factor_a', 'factor_b': 'factor_b', 'response': 'response'}}}, revision=2, dataset_id=d['dataset_id'])
    with pytest.raises(VisualizerContractError, match='missing|Interaction cell'):
        require_valid_statistical_results({'d': bad_d})
    assert store.get(d['dataset_id'])['revision'] == 1


def test_invalid_persisted_analysis_is_explicitly_failed_closed_and_svg_is_finite():
    payload = node_json(r'''
import {renderIntegratedElement} from './company_ui/products/visualizer/assets/element_renderer.mjs';
import {renderChartSvg} from './company_ui/products/visualizer/assets/authoring_chart_studio.mjs';
const viewBox = value => { const match=value.match(/viewBox="([^"]+)"/); if(!match) throw new Error('missing viewBox'); const nums=match[1].trim().split(/\s+/).map(Number); return {nums,finite:nums.length===4&&nums.every(Number.isFinite)&&nums.every(value=>value>=0)}; };
const xbar={engine:'EngineeringChartEngine',element:'Xbar-R Chart',chart_type:'Xbar-R Chart',dataset:{fields:[],rows:[]},mapping:{},statistical_result:null};
const invalidEntry={engine:'EngineeringChartEngine',element:'Xbar-R Chart',title:'Quality review',analysis_error:'Xbar-R requires constant subgroup size. Found subgroup sizes 4 and 5.',analysis_semantics:{ok:false},_resolved_dataset:{fields:[],rows:[]}};
const populated=renderChartSvg({...xbar,statistical_result:{n:2,means:[1,2],ranges:[1,1],xbarLimits:{lcl:0,center:1.5,ucl:3},rLimits:{lcl:0,center:1,ucl:2},rules:{signals:[]}}},{width:760,height:500});
const empty=renderChartSvg(xbar,{width:760,height:500});
const error=renderChartSvg({...xbar,analysis_error:'bad'}, {width:760,height:500});
const card=renderIntegratedElement({...invalidEntry,value:84.2,statistical_result:{means:[84.2]}});
console.log(JSON.stringify({populated:viewBox(populated),empty:viewBox(empty),error:viewBox(error),card,hasDemo:/(84\.2|Baseline|Pilot|Validation)/.test(card)}));
''')
    for key in ('populated', 'empty', 'error'):
        assert payload[key]['finite'] is True
    assert 'Analysis needs attention' in payload['card']
    assert payload['hasDemo'] is False


def test_ppt_statistical_primitives_match_authoritative_editor_result():
    cap_fields = _cap_fields()
    cap_rows = [[9.8, 9.5, 10.5], [10, 9.5, 10.5], [10.2, 9.5, 10.5], [10.1, 9.5, 10.5], [9.9, 9.5, 10.5]]
    cap_item = _cap_item('cap')
    cap_result = _analyze(cap_fields, cap_rows, cap_item)
    x_fields = [{'id': 'subgroup', 'name': 'Subgroup', 'type': 'categorical'}, {'id': 'measurement', 'name': 'Measurement', 'type': 'number'}]
    x_rows = [['A', 1], ['A', 3], ['B', 2], ['B', 6]]
    x_mapping = {'subgroup': 'subgroup', 'value': 'measurement'}
    x_item = {'id': 'xbar', 'engine': 'EngineeringChartEngine', 'type': 'chart', 'element': 'Xbar-R Chart', 'mapping': x_mapping, 'analysis_recipe': {'id': 'xbar-r-process-review', 'mapping': x_mapping}}
    x_result = _analyze(x_fields, x_rows, x_item)
    d_fields = [{'id': 'factor_a', 'name': 'Factor A', 'type': 'categorical'}, {'id': 'factor_b', 'name': 'Factor B', 'type': 'categorical'}, {'id': 'response', 'name': 'Response', 'type': 'number'}]
    d_rows = [['A', 'L', 10], ['A', 'H', 14], ['B', 'L', 20], ['B', 'H', 28]]
    d_mapping = {'factor_a': 'factor_a', 'factor_b': 'factor_b', 'response': 'response'}
    d_item = {'id': 'doe', 'engine': 'EngineeringChartEngine', 'type': 'chart', 'element': 'DOE Interaction Plot', 'mapping': d_mapping, 'analysis_recipe': {'id': 'doe-response-review', 'mapping': d_mapping}}
    d_result = _analyze(d_fields, d_rows, d_item)
    model = canonical_model({'datasets': [
        {'id': 'cap-dataset', 'fields': cap_fields, 'rows': cap_rows},
        {'id': 'x-dataset', 'fields': x_fields, 'rows': x_rows},
        {'id': 'd-dataset', 'fields': d_fields, 'rows': d_rows},
    ], 'items': [
        {**cap_item, 'dataset_id': 'cap-dataset', 'title': 'Capability KPI', 'order': 0, 'x': 0, 'y': 0, 'w': 300, 'h': 150},
        {**x_item, 'dataset_id': 'x-dataset', 'title': 'Xbar-R', 'order': 1, 'x': 0, 'y': 180, 'w': 600, 'h': 250},
        {**d_item, 'dataset_id': 'd-dataset', 'title': 'DOE Interaction', 'order': 2, 'x': 620, 'y': 180, 'w': 600, 'h': 250},
    ], 'nextId': 4})
    model['items'][0]['authoritative_analysis'] = cap_result
    model['items'][1]['authoritative_analysis'] = x_result
    model['items'][2]['authoritative_analysis'] = d_result
    deck = Presentation(io.BytesIO(export_pptx(None, model)))
    texts = '\n'.join(shape.text for shape in deck.slides[0].shapes if getattr(shape, 'has_text_frame', False))
    assert str(cap_result['derived_statistics']['stats']['cpk']) in texts
    assert '84.2' not in texts and 'Complete population count' not in texts
    charts = [shape.chart for shape in deck.slides[0].shapes if getattr(shape, 'has_chart', False)]
    assert len(charts) == 2
    x_values = [list(series.values) for series in charts[0].plots[0].series]
    doe_values = [list(series.values) for series in charts[1].plots[0].series]
    assert x_values == [x_result['derived_statistics']['stats']['means'], x_result['derived_statistics']['stats']['ranges']]
    expected_doe = [[cell['mean'] for cell in row] for row in d_result['derived_statistics']['interaction']['cells']]
    assert doe_values == expected_doe


def test_recommendation_validation_is_conservative_and_shared_with_execution():
    payload = node_json(r'''
import {recommendEngineeringRecipes} from './company_ui/products/visualizer/assets/engineering_recipes.mjs';
const fields=[{id:'g',name:'Subgroup',type:'categorical',semantic_tags:['subgroup']},{id:'m',name:'Measurement',type:'number',semantic_tags:['value']}];
const bad=value=>recommendEngineeringRecipes(fields,[['A',value],['A',2],['B',3],['B',4]]).some(item=>item.id==='xbar-r-process-review');
const doeFields=[{id:'tool',name:'Tool',type:'categorical'},{id:'chamber',name:'Chamber',type:'categorical'},{id:'yield',name:'Yield',type:'number'}];
const doe=items=>recommendEngineeringRecipes(doeFields,items).some(item=>item.id==='doe-response-review');
const tagged=[{id:'a',name:'Factor A',type:'categorical',semantic_tags:['factor_a']},{id:'b',name:'Factor B',type:'categorical',semantic_tags:['factor_b']},{id:'r',name:'Response',type:'number',semantic_tags:['response']}];
console.log(JSON.stringify({null:bad(null),blank:bad(''),whitespace:bad('   '),boolean:bad(true),malformed:bad('abc'),ordinary:doe([['T1','A',90],['T1','B',91],['T2','A',92],['T2','B',93]]),incomplete:recommendEngineeringRecipes(tagged,[['A','L',1],['A','H',2],['B','L',3]]).some(item=>item.id==='doe-response-review')}));
''')
    assert payload == {'null': False, 'blank': False, 'whitespace': False, 'boolean': False, 'malformed': False, 'ordinary': False, 'incomplete': False}


def test_cross_report_resource_access_is_denied_and_stale_revision_is_ignored(tmp_path: Path):
    repository = ReportRepository(tmp_path)
    access = ReportAccessCatalog(repository)
    scoped = ScopedReportRepository(repository, access, Principal('alice', permissions=frozenset({'report.create'})), AuthorizationModel())
    store = DatasetResourceStore(tmp_path)
    resource = store.create(owner='alice', name='scoped', fields=_cap_fields(), rows=[[9.8, 9, 11], [10.2, 9, 11]])
    report_model = canonical_model({'datasets': [{'id': 'd1', 'resource_id': resource['dataset_id'], 'external': True, 'fields': resource['schema'], 'rows': [], 'row_count': 2}], 'items': []})
    scoped.create('report-a', model=report_model)
    scoped.create('report-b', model=canonical_model({'datasets': [], 'items': []}))
    with pytest.raises((PermissionError, ReportNotFoundError)):
        ScopedDatasetRepository(store, scoped).session_for_report('report-b', 'd1')
    assert node_json(r'''
import {acceptsStatisticalResult} from './company_ui/products/visualizer/assets/statistical_result_guard.mjs';
console.log(JSON.stringify({same:acceptsStatisticalResult({expectedRevision:2,resultRevision:2}),stale:acceptsStatisticalResult({expectedRevision:2,resultRevision:1}),request:acceptsStatisticalResult({activeRequestId:'new',resultRequestId:'old'})}));
''') == {'same': True, 'stale': False, 'request': False}


@pytest.mark.parametrize('viewport_width', [1440, 768, 390])
def test_native_editor_statistical_paths_are_browser_safe(viewport_width: int, tmp_path: Path):
    """Exercise the shipped NiceGUI route, not only the embedded editor host."""
    playwright = pytest.importorskip('playwright.sync_api')
    scripts = str(ROOT / 'scripts' / 'release_checks')
    import sys
    sys.path.insert(0, scripts)
    try:
        from editor_host import NativeHost, load_editor
        from run_editor_workflows import model, panels, ready, settled
    finally:
        sys.path.remove(scripts)

    def apply_recipe(page, recipe_id: str, source: str):
        panels(page, library=viewport_width > 800, inspector=False)
        before_count = len(model(page)['items'])
        paste = page.locator('#pasteDataBtn')
        if not paste.is_visible():
            paste = page.locator('#blankStartSurface [data-blank-action="paste"]')
        paste.click()
        page.locator('#dataFirstText').fill(source)
        page.locator('.data-first-summary').wait_for(state='visible', timeout=20000)
        recipe = page.locator(f'[data-data-first-recipe="{recipe_id}"]')
        recipe.wait_for(state='visible', timeout=20000)
        recipe.click()
        recipe.wait_for(state='attached')
        page.locator('#dataFirstApplyRecipe').click()
        page.wait_for_function(
            '(count) => window.CompanyUIVisualizerBridge.state().model.items.length > count',
            arg=before_count,
            timeout=30000,
        )
        settled(page)

    def selected(page, item_id: str):
        component = page.locator(f'.component[data-id="{item_id}"]')
        component.focus()
        component.press('Enter')
        panels(page, library=False, inspector=True)
        return component

    xbar_source = 'Subgroup\tOrder\tMeasurement\n' + '\n'.join(
        f'G{group}\t{group}\t{group + offset}'
        for group in range(1, 4)
        for offset in range(5)
    )

    with tempfile.TemporaryDirectory(prefix='visembler-native-statistical-') as data_dir:
        with NativeHost(ROOT, Path(data_dir) / 'native-data') as host:
            with playwright.sync_playwright() as instance:
                executable = os.environ.get('VISEMBLER_BROWSER') or shutil.which('chromium')
                options = {'headless': True}
                if executable:
                    options.update(executable_path=executable, args=['--no-sandbox'])
                browser = instance.chromium.launch(**options)
                context = browser.new_context(
                    viewport={'width': viewport_width, 'height': 1000},
                    accept_downloads=True,
                )
                errors = []
                failed_requests = []
                page = context.new_page()
                page.on('pageerror', lambda error: errors.append(str(error)))
                page.on('console', lambda message: errors.append(message.text) if message.type == 'error' else None)
                page.on('requestfailed', lambda request: failed_requests.append((request.url, request.failure)))
                try:
                    load_editor(page, host, 'native-statistical-report')
                    ready(page)

                    # Inline Xbar-R, then the same real report is bound to the
                    # governed resource path.  The bound model retains only a
                    # 250-row preview; the rendered semantics must be complete.
                    apply_recipe(page, 'xbar-r-process-review', xbar_source)
                    xbar_id = next(item['id'] for item in model(page)['items'] if item['element'] == 'Xbar-R Chart')
                    xbar = selected(page, xbar_id)
                    assert all(token in xbar.inner_text() for token in ('X̄', 'R', 'UCL', 'LCL'))
                    page.locator('[data-dataset-action="bind-resource"]').click()
                    page.wait_for_function(
                        '() => window.CompanyUIVisualizerBridge.state().model.datasets.some(dataset => dataset.resource_id)',
                        timeout=20000,
                    )
                    settled(page)
                    bound_dataset = next(dataset for dataset in model(page)['datasets'] if dataset.get('resource_id'))
                    assert bound_dataset['row_count'] == 15
                    assert len(bound_dataset['rows']) == 15

                    # A real chart interaction must reach the active-filter
                    # surface without producing a browser error or overflow.
                    xbar = selected(page, xbar_id)
                    points = xbar.locator('[data-chart-point], [data-point], [data-behavior-point]')
                    if points.count():
                        points.first.click()
                        page.locator('#activeFilters').wait_for(state='visible')
                        assert 'Active filters' in page.locator('#activeFilters').inner_text()
                        page.locator('[data-clear-all-filters]').click()
                        page.locator('#activeFilters').wait_for(state='hidden')

                    # Export Center must expose every semantic surface.  The
                    # server PPT request is exercised through the native route;
                    # PPT semantic equality is asserted in the value-level test.
                    page.locator('#exportBtn').click()
                    assert page.locator('#exportPptAction').is_visible()
                    assert page.locator('#exportSvgAction').is_visible()
                    assert page.locator('#exportPngAction').is_visible()
                    assert page.locator('#exportJpegAction').is_visible()
                    page.locator('#exportPptAction').click()
                    page.wait_for_timeout(1000)
                    export_text = page.locator('body').inner_text()
                    assert (
                        'Editable PowerPoint export generated' in export_text
                        or 'Resolve export-blocking validation issues first' in export_text
                    )
                    if page.locator('#genericModal.show').count():
                        page.locator('#genericModal.show [data-close]').first.click()

                    # Refresh the bound Xbar-R into mixed subgroup sizes.  The
                    # authoritative refresh boundary rejects it and leaves the
                    # prior revision/analysis visible.
                    selected(page, xbar_id)
                    page.locator('[data-refresh-dataset]').click()
                    page.locator('#refreshDataText').fill('Subgroup\tOrder\tMeasurement\nA\t1\t1\nA\t1\t2\nB\t2\t3\nB\t2\t4\nB\t2\t5')
                    refresh_linked = page.locator('#refreshLinked')
                    if refresh_linked.is_enabled():
                        refresh_linked.click()
                        page.get_by_text('constant subgroup size').wait_for(timeout=20000)
                    else:
                        assert page.locator('#modalBody').inner_text().strip()
                        page.locator('#modalBody [data-close]').click()
                    assert next(dataset for dataset in model(page)['datasets'] if dataset.get('resource_id'))['revision'] == 1
                    assert all(token in page.locator(f'.component[data-id="{xbar_id}"]').inner_text() for token in ('X̄', 'R', 'UCL', 'LCL'))

                    assert page.evaluate('document.documentElement.scrollWidth <= document.documentElement.clientWidth + 1')
                    browser_text = page.locator('body').inner_text()
                    assert 'NaN' not in browser_text
                    assert 'Infinity' not in browser_text
                    assert not errors, errors
                    assert not failed_requests, failed_requests
                finally:
                    context.close()
                    browser.close()


def test_native_bound_capability_uses_full_population_after_reopen(tmp_path: Path):
    """The actual NiceGUI resource path must not regress to the preview."""
    playwright = pytest.importorskip('playwright.sync_api')
    scripts = str(ROOT / 'scripts' / 'release_checks')
    import sys
    sys.path.insert(0, scripts)
    try:
        from editor_host import NativeHost, load_editor
        from run_editor_workflows import model, panels, ready, settled
    finally:
        sys.path.remove(scripts)

    rows = [[1 if index < 250 else 10, 0, 10] for index in range(500)]
    cap_item = {
        **_cap_item('cap'),
        'type': 'chart',
        'engine': 'CoreChartEngine',
        'element': 'Histogram',
        'title': 'Capability Histogram',
        'dataset_id': 'd1',
        'order': 0,
        'x': 20,
        'y': 20,
        'w': 720,
        'h': 360,
    }
    report_model = canonical_model({
        'datasets': [{'id': 'd1', 'name': 'Capability source', 'fields': _cap_fields(), 'rows': rows}],
        'items': [cap_item],
        'nextId': 2,
    })
    with tempfile.TemporaryDirectory(prefix='visembler-native-capability-') as data_dir:
        with NativeHost(ROOT, Path(data_dir) / 'native-data') as host:
            report_id = host.create(model=report_model)
            with playwright.sync_playwright() as instance:
                executable = os.environ.get('VISEMBLER_BROWSER') or shutil.which('chromium')
                options = {'headless': True}
                if executable:
                    options.update(executable_path=executable, args=['--no-sandbox'])
                browser = instance.chromium.launch(**options)
                context = browser.new_context(viewport={'width': 1440, 'height': 1000}, accept_downloads=True)
                errors = []
                failed_requests = []
                page = context.new_page()
                page.on('pageerror', lambda error: errors.append(str(error)))
                page.on('console', lambda message: errors.append(message.text) if message.type == 'error' else None)
                page.on('requestfailed', lambda request: failed_requests.append((request.url, request.failure)))
                try:
                    page.goto(f'{host.url}/visualizer?report={report_id}', wait_until='domcontentloaded')
                    ready(page)
                    page.wait_for_function('() => window.CompanyUIVisualizerBridge?.state', timeout=20000)
                    page.wait_for_timeout(500)
                    cap_id = 'cap'
                    cap = page.locator(f'.component[data-id="{cap_id}"]')
                    assert 'n=500' in cap.inner_text()
                    assert 'Cpk 0.333' in cap.inner_text()

                    cap.focus()
                    cap.press('Enter')
                    panels(page, library=False, inspector=True)
                    page.locator('[data-dataset-action="bind-resource"]').click()
                    page.wait_for_function(
                        '() => window.CompanyUIVisualizerBridge.state().model.datasets.some(dataset => dataset.resource_id)',
                        timeout=20000,
                    )
                    settled(page)
                    dataset = next(dataset for dataset in model(page)['datasets'] if dataset.get('resource_id'))
                    assert dataset['row_count'] == 500
                    assert len(dataset['rows']) == 250
                    assert 'n=500' in cap.inner_text()
                    assert 'Cpk 0.333' in cap.inner_text()

                    # Reopen through the real route and prove the persisted
                    # binding resolves the same complete result.
                    report_id = page.evaluate('() => window.CompanyUIVisualizerBridge.state().report_id')
                    page.close()
                    page = context.new_page()
                    page.on('pageerror', lambda error: errors.append(str(error)))
                    page.on('console', lambda message: errors.append(message.text) if message.type == 'error' else None)
                    page.on('requestfailed', lambda request: failed_requests.append((request.url, request.failure)))
                    page.goto(f'{host.url}/visualizer?report={report_id}', wait_until='domcontentloaded')
                    ready(page)
                    cap = page.locator(f'.component[data-id="{cap_id}"]')
                    assert 'n=500' in cap.inner_text()
                    assert 'Cpk 0.333' in cap.inner_text()

                    page.locator('#exportBtn').click()
                    assert page.locator('#exportPptAction').is_visible()
                    assert page.locator('#exportSvgAction').is_visible()
                    page.locator('#exportPptAction').click()
                    page.wait_for_timeout(1000)
                    assert 'Editable PowerPoint export generated' in page.locator('body').inner_text() or 'Resolve export-blocking validation issues first' in page.locator('body').inner_text()

                    cap.focus()
                    cap.press('Enter')
                    panels(page, library=False, inspector=True)
                    page.locator('[data-refresh-dataset]').click()
                    page.locator('#refreshDataText').fill('Measurement\tLSL\tUSL\n1\t0\t10\n10\t0\t9')
                    refresh_linked = page.locator('#refreshLinked')
                    if refresh_linked.is_enabled():
                        refresh_linked.click()
                        page.wait_for_timeout(1000)
                        assert 'specification' in page.locator('body').inner_text().lower()
                    else:
                        assert page.locator('#modalBody').inner_text().strip()
                        page.locator('#modalBody [data-close]').click()
                    assert next(dataset for dataset in model(page)['datasets'] if dataset.get('resource_id'))['revision'] == 1
                    assert 'Cpk 0.333' in cap.inner_text()
                    assert page.evaluate('document.documentElement.scrollWidth <= document.documentElement.clientWidth + 1')
                    browser_text = page.locator('body').inner_text()
                    assert 'NaN' not in browser_text and 'Infinity' not in browser_text
                    assert not errors, errors
                    assert not failed_requests, failed_requests
                finally:
                    context.close()
                    browser.close()
