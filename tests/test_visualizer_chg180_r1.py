from __future__ import annotations

import base64
import copy
import io
import json
import shutil
import subprocess
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation

from company_ui.products.visualizer.domain import VisualizerContractError, canonical_model
from company_ui.products.visualizer.ppt_service import _kind, bound_export_items, export_pptx

ROOT = Path(__file__).resolve().parents[1]


def _wafer_model() -> dict:
    fields = [
        {'id': 'die_x', 'name': 'Die X', 'type': 'integer'},
        {'id': 'die_y', 'name': 'Die Y', 'type': 'integer'},
        {'id': 'yield', 'name': 'Yield', 'type': 'number'},
        {'id': 'reference', 'name': 'Reference', 'type': 'number'},
        {'id': 'affected', 'name': 'Affected', 'type': 'number'},
        {'id': 'wafer', 'name': 'Wafer', 'type': 'identifier'},
        {'id': 'lot', 'name': 'Lot', 'type': 'identifier'},
        {'id': 'tool', 'name': 'Tool', 'type': 'categorical'},
        {'id': 'chamber', 'name': 'Chamber', 'type': 'categorical'},
        {'id': 'recipe', 'name': 'Recipe', 'type': 'categorical'},
        {'id': 'process', 'name': 'Process', 'type': 'categorical'},
    ]
    rows = [
        [0, 0, 98.1, 100.0, 100.0, 'W07', 'L2407', 'ETCH-08', 'A', 'REC-17', 'Etch'],
        [1, 0, 96.0, 100.0, 95.0, 'W07', 'L2407', 'ETCH-08', 'A', 'REC-17', 'Etch'],
        [0, 1, 0, 100.0, 102.0, 'W07', 'L2407', 'ETCH-08', 'A', 'REC-17', 'Etch'],
        [1, 1, None, 102.0, None, 'W07', 'L2407', 'ETCH-08', 'A', 'REC-17', 'Etch'],
    ]
    return canonical_model({
        'mode': 'smart', 'canvas': {'width': 1600, 'height': 900}, 'nextId': 3,
        'datasets': [{'id': 'd-wafer', 'name': 'CHG-180 RCA wafer evidence', 'fields': fields, 'rows': rows}],
        'items': [
            {
                'id': 'c1', 'type': 'wafer', 'engine': 'WaferFabEngine', 'element': 'Wafer Map',
                'title': 'Affected Wafer Map', 'order': 0, 'dataset_id': 'd-wafer',
                'mapping': {'die_x': 'die_x', 'die_y': 'die_y', 'value': 'yield', 'wafer_id': 'wafer', 'lot_id': 'lot', 'tool': 'tool', 'chamber': 'chamber', 'recipe': 'recipe', 'process': 'process'},
            },
            {
                'id': 'c2', 'type': 'wafer', 'engine': 'WaferFabEngine', 'element': 'Wafer Difference Map',
                'title': 'Affected − Reference by Die', 'order': 1, 'dataset_id': 'd-wafer',
                'mapping': {'die_x': 'die_x', 'die_y': 'die_y', 'reference_value': 'reference', 'affected_value': 'affected', 'wafer_id': 'wafer', 'lot_id': 'lot', 'tool': 'tool', 'chamber': 'chamber', 'recipe': 'recipe', 'process': 'process'},
            },
        ],
    })


def _two_map_layout() -> dict:
    return {
        'canvas': {'width': 1600, 'height': 900},
        'items': [
            {'id': 'c1', 'x': 20, 'y': 20, 'w': 760, 'h': 860},
            {'id': 'c2', 'x': 820, 'y': 20, 'w': 760, 'h': 860},
        ],
    }


def _descr(shape) -> list[str]:
    return [node.get('descr') or '' for node in shape._element.xpath('.//p:cNvPr')]


def test_chg180_waferfab_map_variants_are_first_class_export_kinds():
    assert _kind({'engine': 'WaferFabEngine', 'element': 'Wafer Map'}) == 'wafer_map'
    assert _kind({'engine': 'WaferFabEngine', 'element': 'Wafer Difference Map'}) == 'wafer_difference'
    assert _kind({'engine': 'WaferFabEngine', 'element': 'Tool × Chamber Matrix'}) == 'fallback'


def test_chg180_bound_wafer_projection_preserves_die_values_identity_and_source_model():
    model = _wafer_model()
    source = copy.deepcopy(model)
    projected = bound_export_items(model)
    by_id = {entry['id']: entry for entry in projected}

    wafer = by_id['c1']
    assert wafer['observations'] == [
        {'x': 0, 'y': 0, 'value': 98.1},
        {'x': 1, 'y': 0, 'value': 96.0},
        {'x': 0, 'y': 1, 'value': 0},
        {'x': 1, 'y': 1, 'value': None},
    ]
    assert [wafer[key] for key in ('wafer_id', 'lot_id', 'tool', 'chamber', 'recipe', 'process')] == ['W07', 'L2407', 'ETCH-08', 'A', 'REC-17', 'Etch']

    difference = by_id['c2']['observations']
    assert [(row['x'], row['y'], row['reference_value'], row['affected_value'], row['delta']) for row in difference] == [
        (0, 0, 100.0, 100.0, 0.0),
        (1, 0, 100.0, 95.0, -5.0),
        (0, 1, 100.0, 102.0, 2.0),
        (1, 1, 102.0, None, None),
    ]
    assert model == source


def test_chg180_powerpoint_has_editable_spatial_cells_not_generic_fallback():
    model = _wafer_model()
    original = copy.deepcopy(model)
    deck = Presentation(io.BytesIO(export_pptx(None, model, layout_geometry=_two_map_layout())))
    slide = deck.slides[0]
    shapes = list(slide.shapes)

    assert len([shape for shape in shapes if shape.name.endswith('wafer-outline')]) == 2
    assert len([shape for shape in shapes if 'orientation-notch' in shape.name]) == 2
    wafer_cells = [shape for shape in shapes if shape.name.startswith('VIZ::Affected Wafer Map::die-')]
    delta_cells = [shape for shape in shapes if shape.name.startswith('VIZ::Affected − Reference by Die::die-')]
    assert len(wafer_cells) == len(delta_cells) == 4
    assert all(shape.left > 0 and shape.top > 0 and shape.width > 0 and shape.height > 0 for shape in wafer_cells + delta_cells)

    def die_map(cells):
        result = {}
        for shape in cells:
            payload = next(value.removeprefix('VisualizerSemanticDie:') for value in _descr(shape) if value.startswith('VisualizerSemanticDie:'))
            observation = json.loads(payload)
            result[(observation['x'], observation['y'])] = {'shape': shape, 'observation': observation}
        return result

    wafer_dies, delta_dies = die_map(wafer_cells), die_map(delta_cells)
    assert wafer_dies[(1, 0)]['shape'].left > wafer_dies[(0, 0)]['shape'].left
    assert wafer_dies[(0, 1)]['shape'].top < wafer_dies[(0, 0)]['shape'].top
    assert delta_dies[(1, 0)]['shape'].left > delta_dies[(0, 0)]['shape'].left
    assert delta_dies[(0, 1)]['shape'].top < delta_dies[(0, 0)]['shape'].top
    negative_color = tuple(delta_dies[(1, 0)]['shape'].fill.fore_color.rgb)
    positive_color = tuple(delta_dies[(0, 1)]['shape'].fill.fore_color.rgb)
    assert negative_color[2] > negative_color[0]
    assert positive_color[0] > positive_color[2]

    wafer_panel = next(shape for shape in shapes if shape.name == 'VIZ::Affected Wafer Map::spatial-panel')
    difference_panel = next(shape for shape in shapes if shape.name == 'VIZ::Affected − Reference by Die::spatial-panel')
    wafer_payload = json.loads(next(value.removeprefix('VisualizerSemantic:') for value in _descr(wafer_panel) if value.startswith('VisualizerSemantic:')))
    difference_payload = json.loads(next(value.removeprefix('VisualizerSemantic:') for value in _descr(difference_panel) if value.startswith('VisualizerSemantic:')))
    assert [row['value'] for row in wafer_payload['observations']] == [98.1, 96.0, 0, None]
    assert [wafer_payload[key] for key in ('wafer_id', 'lot_id', 'tool', 'chamber', 'recipe', 'process')] == ['W07', 'L2407', 'ETCH-08', 'A', 'REC-17', 'Etch']
    assert [row['delta'] for row in difference_payload['observations']] == [0.0, -5.0, 2.0, None]
    assert all(any(value.startswith('VisualizerSemanticDie:') for value in _descr(shape)) for shape in delta_cells)

    text = '\n'.join(shape.text for shape in shapes if getattr(shape, 'has_text_frame', False))
    for expected in ('W07', 'L2407', 'ETCH-08', 'Chamber A', 'REC-17', 'Process Etch', 'SIGNED DELTA · AFFECTED − REFERENCE', '-5', '+2'):
        assert expected in text
    assert not any('Controlled fallback' in shape.text for shape in shapes if getattr(shape, 'has_text_frame', False))
    assert model == original


def test_chg180_geometry_payload_is_rejected_when_incomplete_overlapping_or_out_of_hull():
    model = _wafer_model()
    layout = _two_map_layout()
    layout['items'] = layout['items'][:1]
    with pytest.raises(VisualizerContractError, match='do not match'):
        export_pptx(None, model, layout_geometry=layout)
    layout = _two_map_layout()
    layout['items'][1]['x'] = 700
    with pytest.raises(VisualizerContractError, match='overlapping'):
        export_pptx(None, model, layout_geometry=layout)
    layout = _two_map_layout()
    layout['items'][1]['y'] = 890
    with pytest.raises(VisualizerContractError, match='exceeds'):
        export_pptx(None, model, layout_geometry=layout)


def test_chg180_keeps_existing_chart_table_diagram_image_and_metric_export_kinds():
    expected = {
        'CoreChartEngine': 'chart', 'EngineeringChartEngine': 'chart',
        'TableEngine': 'table', 'MatrixEngine': 'table', 'DiagramEngine': 'diagram',
        'ImageMediaEngine': 'image', 'MetricEngine': 'kpi', 'ComparisonEngine': 'kpi',
    }
    for engine, kind in expected.items():
        assert _kind({'engine': engine}) == kind

    image = io.BytesIO(); Image.new('RGB', (8, 8), (30, 80, 160)).save(image, format='PNG')
    encoded = 'data:image/png;base64,' + base64.b64encode(image.getvalue()).decode('ascii')
    entries = [
        {'id': 'm', 'type': 'metric', 'engine': 'MetricEngine', 'element': 'Hero KPI', 'title': 'Yield', 'value': 0, 'order': 0},
        {'id': 'c', 'type': 'chart', 'engine': 'CoreChartEngine', 'element': 'Line Chart', 'title': 'Trend', 'data': [['W1', 0], ['W2', 5]], 'order': 1},
        {'id': 't', 'type': 'table', 'engine': 'TableEngine', 'element': 'Clean Table', 'title': 'Actions', 'customTable': {'headers': ['Action', 'Status'], 'rows': [['Verify', 'Open']]}, 'order': 2},
        {'id': 'd', 'type': 'diagram', 'engine': 'DiagramEngine', 'element': 'Process Flow', 'title': 'Flow', 'nodes': ['Inspect', 'Release'], 'edges': [['Inspect', 'Release']], 'order': 3},
        {'id': 'i', 'type': 'image', 'engine': 'ImageMediaEngine', 'element': 'Image', 'title': 'Evidence image', 'src': encoded, 'order': 4},
    ]
    geometry = {'canvas': {'width': 1600, 'height': 900}, 'items': [
        {'id': entry['id'], 'x': 20 + index * 310, 'y': 20, 'w': 290, 'h': 300} for index, entry in enumerate(entries)
    ]}
    deck = Presentation(io.BytesIO(export_pptx(None, canonical_model({'items': entries, 'nextId': 6}), layout_geometry=geometry)))
    shapes = list(deck.slides[0].shapes)
    assert any(getattr(shape, 'has_chart', False) for shape in shapes)
    assert any(getattr(shape, 'has_table', False) for shape in shapes)
    assert any(getattr(shape, 'shape_type', None) == 13 for shape in shapes)
    assert any(shape.name.startswith('VIZ::Flow::Inspect') for shape in shapes)
    assert any(getattr(shape, 'has_text_frame', False) and '0' in shape.text for shape in shapes)


def test_chg180_process_flow_geometry_is_rebuilt_from_saved_diagram_content():
    if not shutil.which('node'):
        pytest.skip('Node.js is required for the maintained Visembler renderer checks.')
    script = """
      import {diagramRenderBounds, renderDiagramSvg} from './company_ui/products/visualizer/assets/authoring_diagram_studio.mjs';
      const labels=['Problem definition and scope','Method qualification','Machine chamber verification','Material traceability review','Measurement system confirmation','Environment excursion analysis','Containment effectiveness review','Corrective action verification','Production release approval'];
      const entry={engine:'DiagramEngine',element:'Process Flow',direction:'down',nodes:labels,edges:labels.slice(0,-1).map((label,index)=>[label,labels[index+1]])};
      const bounds=diagramRenderBounds(entry),again=diagramRenderBounds(structuredClone(entry)),svg=renderDiagramSvg(entry);
      console.log(JSON.stringify({bounds,again,svg}));
    """
    result = subprocess.run(['node', '--input-type=module', '-e', script], cwd=ROOT, text=True, capture_output=True, check=True)
    projection = json.loads(result.stdout)
    assert projection['bounds'] == projection['again']
    assert projection['bounds']['height'] > 1400
    view_box = projection['svg'].split('viewBox="', 1)[1].split('"', 1)[0]
    assert float(view_box.split()[3]) == projection['bounds']['height']
    for label in ('Problem definition and scope', 'Machine chamber verification', 'Production release approval'):
        assert label in projection['svg']
