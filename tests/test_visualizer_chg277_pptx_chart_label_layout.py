from __future__ import annotations

import io
import json
import zipfile
import xml.etree.ElementTree as ET

import pytest
from pptx import Presentation
from pptx.util import Pt

from company_ui.products.visualizer.domain import canonical_model
from company_ui.products.visualizer.ppt_service import export_pptx, import_visembler_pptx


NS = {
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart',
}
EMU_PER_POINT = 12700


def _chart_model(*, categories, series, title='Measurements over time', chart_type='Multi-Line', x_axis=None, geometry=(80, 90, 520, 320)):
    fields = [
        {'id': 'period', 'name': 'Period', 'type': 'categorical'},
        {'id': 'group', 'name': 'Group', 'type': 'categorical'},
        {'id': 'reading', 'name': 'Reading', 'type': 'number'},
    ]
    rows = [[category, item['name'], item['values'][index]]
            for index, category in enumerate(categories) for item in series]
    axes_x = {
        'auto': True, 'labelInterval': 1, 'rotation': 0, 'tickCount': 6,
        'title': '', 'format': 'auto', 'show': True, 'position': 'bottom',
    }
    axes_x.update(x_axis or {})
    entry = {
        'id': 'c1', 'type': 'chart', 'engine': 'CoreChartEngine', 'element': chart_type,
        'title': title, 'order': 0, 'dataset_id': 'd1',
        'mapping': {'x': 'period', 'category': 'period', 'series': 'group', 'y': 'reading', 'value': 'reading'},
        'chart_studio': {
            'chart_type': chart_type,
            'mapping': {'x': 'period', 'series': 'group', 'y': 'reading'},
            'axes': {
                'x': axes_x,
                'y': {'title': 'Reading', 'min': None, 'max': None, 'auto': True,
                      'zeroBaseline': False, 'format': 'auto', 'tickCount': 5},
            },
            'legend': {'show': len(series) > 1, 'position': 'bottom', 'order': 'input'},
            'series': [{'id': f's{index + 1}', 'key': item['name'], 'field': 'group',
                        'label': item['name'], 'visible': True, 'legend': True, 'axis': 'primary'}
                       for index, item in enumerate(series)],
        },
        'x': geometry[0], 'y': geometry[1], 'w': geometry[2], 'h': geometry[3],
    }
    return canonical_model({
        'datasets': [{'id': 'd1', 'name': 'Readings', 'fields': fields, 'rows': rows}],
        'items': [entry], 'nextId': 2, 'canvas': {'width': 1600, 'height': 900},
    })


def _compact_geometry(model, rect=None):
    item = model['items'][0]
    x, y, width, height = rect or tuple(item[key] for key in ('x', 'y', 'w', 'h'))
    return {'canvas': model['canvas'], 'items': [{'id': item['id'], 'x': x, 'y': y, 'w': width, 'h': height}]}


def _chart_package(payload):
    package = zipfile.ZipFile(io.BytesIO(payload))
    paths = [name for name in package.namelist() if name.startswith('ppt/charts/chart') and name.endswith('.xml')]
    assert len(paths) == 1
    return package, ET.fromstring(package.read(paths[0]))


def _series_facts(root):
    result = []
    for node in root.findall('.//c:ser', NS):
        name = node.findtext('./c:tx//c:v', namespaces=NS)
        categories = [(int(point.get('idx')), point.findtext('./c:v', namespaces=NS))
                      for point in node.findall('./c:cat//c:pt', NS)]
        values = [(int(point.get('idx')), point.findtext('./c:v', namespaces=NS))
                  for point in node.findall('./c:val//c:pt', NS)]
        result.append({'name': name, 'categories': categories, 'values': values})
    return result


def _line_values(root):
    return [[value for _, value in row['values']] for row in _series_facts(root)]


def _rotation_degrees(root):
    body = root.find('.//c:catAx/c:txPr/a:bodyPr', NS)
    return 0.0 if body is None or body.get('rot') is None else -int(body.get('rot')) / 60000


def _tick_interval(root):
    node = root.find('.//c:catAx/c:tickLblSkip', NS)
    return 1 if node is None else int(node.get('val'))


def _font_size_pt(node):
    sizes = [int(value.get('sz')) / 100 for value in node.findall('.//a:defRPr', NS) if value.get('sz')]
    return min(sizes) if sizes else None


def _assert_no_box_intersections(shapes):
    boxes = [(shape.left, shape.top, shape.left + shape.width, shape.top + shape.height) for shape in shapes]
    for index, first in enumerate(boxes):
        for second in boxes[index + 1:]:
            overlap_x = min(first[2], second[2]) - max(first[0], second[0])
            overlap_y = min(first[3], second[3]) - max(first[1], second[1])
            assert overlap_x <= 0 or overlap_y <= 0


def test_chg277_dense_e_like_chart_keeps_all_chg254_semantics_with_distinct_export_labels():
    categories = ['17:00', '17:30', '18:00', '18:30', '19:00', '19:30', '20:00', '20:30', '21:00']
    series = [
        {'name': 'Battery', 'values': [34, 32, 29, 28, 25, 24, 22, 21, 20]},
        {'name': 'Demand response', 'values': [18, 18, 16, 16, 14, 12, 10, 8, 8]},
        {'name': 'Firm imports', 'values': [41, 40, 40, 39, 37, 37, 36, 34, 33]},
    ]
    source = _chart_model(categories=categories, series=series, title='Delivered MW by dispatch resource',
                          geometry=(424, 400, 534, 259))
    original = json.loads(json.dumps(source))
    payload = export_pptx(None, source, layout_geometry=_compact_geometry(source))
    deck = Presentation(io.BytesIO(payload))
    chart = next(shape.chart for shape in deck.slides[0].shapes if getattr(shape, 'has_chart', False))
    package, root = _chart_package(payload)

    assert chart.chart_type.name == 'LINE_MARKERS'
    assert [series.name for series in chart.series] == ['Battery', 'Demand response', 'Firm imports']
    assert [list(series.values) for series in chart.series] == [item['values'] for item in series]
    assert chart.value_axis.minimum_scale == pytest.approx(6.35)
    assert chart.value_axis.maximum_scale == pytest.approx(42.65)
    facts = _series_facts(root)
    assert [item['name'] for item in facts] == [item['name'] for item in series]
    assert all(item['categories'] == list(enumerate(categories)) for item in facts)
    assert _line_values(root) == [[f'{float(value):.1f}' for value in item['values']] for item in series]
    assert _tick_interval(root) == 2
    assert _rotation_degrees(root) == 0
    assert _font_size_pt(root.find('.//c:catAx/c:txPr', NS)) == pytest.approx(10)
    assert _font_size_pt(root.find('.//c:legend/c:txPr', NS)) >= 7.5
    assert not root.findall('.//c:legend/c:legendEntry/c:delete[@val="1"]', NS)
    assert len([name for name in package.namelist() if name.startswith('ppt/charts/chart') and name.endswith('.xml')]) == 1
    slide = Presentation(io.BytesIO(payload)).slides[0]
    overlays = [shape for shape in slide.shapes if '::category-label-' in shape.name and getattr(shape, 'has_text_frame', False)]
    assert [shape.text for shape in overlays] == categories[::2]
    legend_labels = [shape for shape in slide.shapes if '::legend-label-' in shape.name]
    assert [shape.text for shape in legend_labels] == [item['name'] for item in series]
    _assert_no_box_intersections(overlays)
    _assert_no_box_intersections(legend_labels)
    assert source == original
    assert import_visembler_pptx(payload) == source


def test_chg277_explicit_chart_studio_x_axis_interval_rotation_title_and_format_win():
    categories = [f'Long interval {index:02d} with shift note' for index in range(1, 10)]
    series = [{'name': 'Reference measurements', 'values': list(range(20, 29))},
              {'name': 'Adjusted measurements', 'values': list(range(24, 33))}]
    source = _chart_model(categories=categories, series=series, x_axis={
        'labelInterval': 3, 'rotation': 18, 'tickCount': 4, 'title': 'Observation period', 'format': 'date',
    })
    payload = export_pptx(None, source, layout_geometry=_compact_geometry(source))
    _, root = _chart_package(payload)
    assert _tick_interval(root) == 3
    assert _rotation_degrees(root) == pytest.approx(18)
    assert root.findtext('.//c:catAx/c:title/c:tx//a:t', namespaces=NS) == 'Observation period'
    number_format = root.find('.//c:catAx/c:numFmt', NS)
    assert number_format is not None and number_format.get('formatCode') == 'yyyy-mm-dd hh:mm'
    facts = _series_facts(root)
    assert all([value for _, value in item['categories']] == categories for item in facts)
    assert import_visembler_pptx(payload) == source


def test_chg277_automatic_policy_generalizes_to_dense_labels_and_leaves_sparse_labels_alone():
    dense_categories = [f'Night shift {index:02d}' for index in range(1, 12)]
    dense_series = [
        {'name': 'Baseline tool pressure', 'values': list(range(20, 31))},
        {'name': 'Conditioned chamber', 'values': list(range(22, 33))},
        {'name': 'Adjusted gas flow', 'values': list(range(24, 35))},
    ]
    dense = _chart_model(categories=dense_categories, series=dense_series, title='Independent chamber cycle holdout')
    dense_payload = export_pptx(None, dense, layout_geometry=_compact_geometry(dense))
    _, dense_root = _chart_package(dense_payload)
    assert _tick_interval(dense_root) >= 2
    assert _rotation_degrees(dense_root) == 0
    assert len(_series_facts(dense_root)) == 3
    assert all(item['categories'] == list(enumerate(dense_categories)) for item in _series_facts(dense_root))
    assert [item['name'] for item in _series_facts(dense_root)] == [item['name'] for item in dense_series]
    assert _font_size_pt(dense_root.find('.//c:legend/c:txPr', NS)) >= 7.5
    assert not dense_root.findall('.//c:legend/c:legendEntry/c:delete[@val="1"]', NS)
    dense_slide = Presentation(io.BytesIO(dense_payload)).slides[0]
    dense_interval = _tick_interval(dense_root)
    category_labels = [shape for shape in dense_slide.shapes if '::category-label-' in shape.name
                       and getattr(shape, 'has_text_frame', False)]
    legend_labels = [shape for shape in dense_slide.shapes if '::legend-label-' in shape.name]
    assert [shape.text for shape in category_labels] == dense_categories[::dense_interval]
    assert [shape.text for shape in legend_labels] == [item['name'] for item in dense_series]
    _assert_no_box_intersections(category_labels)
    _assert_no_box_intersections(legend_labels)
    assert all(label.text_frame.paragraphs[0].font.size.pt >= 7.5 for label in legend_labels)

    sparse = _chart_model(categories=['W1', 'W2', 'W3', 'W4'],
                          series=[{'name': 'Control', 'values': [8, 9, 10, 11]},
                                  {'name': 'Treatment', 'values': [9, 10, 12, 13]}],
                          geometry=(30, 40, 720, 380))
    sparse_payload = export_pptx(None, sparse, layout_geometry=_compact_geometry(sparse))
    _, sparse_root = _chart_package(sparse_payload)
    assert _tick_interval(sparse_root) == 1
    assert _rotation_degrees(sparse_root) == 0
    assert sparse_root.find('./c:catAx/c:tickLblSkip', NS) is None
    assert not [shape for shape in Presentation(io.BytesIO(sparse_payload)).slides[0].shapes
                if '::category-label-' in shape.name or '::legend-label-' in shape.name]


def test_chg277_box_plot_bands_separate_title_and_upper_tick_without_changing_statistics():
    fields = [
        {'id': 'cohort', 'name': 'Cohort', 'type': 'categorical'},
        {'id': 'yield', 'name': 'Yield', 'type': 'number', 'unit': '%'},
    ]
    rows = [['Reference', 98.8], ['Reference', 98.5], ['Reference', 98.6],
            ['Affected', 94.1], ['Affected', 94.6], ['Affected', 93.8]]
    item = {
        'id': 'c6', 'type': 'chart', 'engine': 'CoreChartEngine', 'element': 'Box Plot',
        'title': 'Control versus affected distribution', 'order': 0, 'dataset_id': 'd1',
        'mapping': {'category': 'cohort', 'cohort': 'cohort', 'value': 'yield', 'x': 'yield', 'y': 'yield'},
        'rows': [{'label': row[0], 'value': row[1]} for row in rows],
        'x': 803, 'y': 982, 'w': 345, 'h': 175,
    }
    source = canonical_model({'datasets': [{'id': 'd1', 'name': 'Yield observations', 'fields': fields, 'rows': rows}],
                              'items': [item], 'nextId': 7, 'canvas': {'width': 1275, 'height': 1600}})
    geometry = {'canvas': source['canvas'], 'items': [{'id': 'c6', 'x': 803, 'y': 982, 'w': 345, 'h': 175}]}
    payload = export_pptx(None, source, layout_geometry=geometry)
    deck = Presentation(io.BytesIO(payload))
    slide = deck.slides[0]
    package = zipfile.ZipFile(io.BytesIO(payload))
    shapes = {shape.name: shape for shape in slide.shapes}
    title = shapes['VIZ::c6::box-plot-title']
    upper_tick = shapes['VIZ::c6::box-plot-axis-4']
    gap_pt = (upper_tick.top - title.top - title.height) / EMU_PER_POINT

    assert gap_pt >= 2.5
    assert not [name for name in package.namelist() if name.startswith('ppt/charts/chart') and name.endswith('.xml')]
    assert len([shape for shape in slide.shapes if shape.name.startswith('VIZ::c6::box-plot-')]) == 23
    assert sum(shape.name.endswith('-quartiles') for shape in slide.shapes) == 2
    assert sum(shape.name.endswith('-median') for shape in slide.shapes) == 2
    assert sum(shape.name.endswith('-whisker') for shape in slide.shapes) == 2
    categories = {shape.text: json.loads(shape._element.xpath('.//p:cNvPr')[0].get('title'))['quartiles']
                  for shape in slide.shapes if shape.name.endswith('-category')}
    assert set(categories) == {'Reference', 'Affected'}
    expected = {
        'Reference': {'n': 3, 'min': 98.5, 'q1': 98.55, 'median': 98.6, 'q3': 98.7, 'max': 98.8},
        'Affected': {'n': 3, 'min': 93.8, 'q1': 93.95, 'median': 94.1, 'q3': 94.35, 'max': 94.6},
    }
    for category, statistics in expected.items():
        assert categories[category]['n'] == statistics['n']
        for field in ('min', 'q1', 'median', 'q3', 'max'):
            assert categories[category][field] == pytest.approx(statistics[field])
    assert import_visembler_pptx(payload) == source
