from __future__ import annotations

import io
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import MSO_COLOR_TYPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn

from company_ui.products.visualizer.domain import canonical_model
from company_ui.products.visualizer.ppt_service import _report_text_segments, export_pptx, import_visembler_pptx


ROOT = Path(__file__).resolve().parents[1]


def _long_report() -> tuple[dict, dict]:
    headline = (
        'FY26 expansion remains supported by the second consecutive quarter of retention gains, '
        'while the public-sector renewal gap and unsigned procurement cycles keep the capacity '
        'release conditional until the named contracts clear the agreed finance checkpoint.'
    )
    narrative = (
        'The observed movement is consistent across the reporting window and remains bounded by '
        'the matched customer cohort. ' * 7
    ).strip()
    items = [
        {'id':'headline','type':'text','engine':'TextEngine','element':'Executive Statement','title':'Executive Statement','text':headline,'order':0},
        {'id':'context','type':'text','engine':'TextEngine','element':'Report context','title':'Quarterly context','text':'FY26 Q1–Q4 close · enterprise and public-sector portfolios · constant-currency bookings · renewal cohort excludes acquired accounts.','order':1},
        {'id':'narrative','type':'text','engine':'TextEngine','element':'Body Narrative','title':'Narrative interpretation','text':narrative,'order':2},
        {'id':'percent','type':'metric','engine':'MetricEngine','element':'Hero KPI','title':'Treatment completion','value':.276,'metric_format':{'kind':'percent','percent_scale':'ratio','precision':1,'suffix':'%'},'order':3},
        {'id':'currency','type':'metric','engine':'MetricEngine','element':'Hero KPI','title':'Bookings','value':1234567890.12,'metric_format':{'kind':'currency','prefix':'$','precision':2},'order':4},
        {'id':'unit','type':'metric','engine':'MetricEngine','element':'Hero KPI','title':'Support contacts','value':34,'unit':'contacts','metric_format':{'kind':'number','precision':0},'order':5},
        {'id':'negative','type':'metric','engine':'MetricEngine','element':'Hero KPI','title':'Net change','value':-2.35,'metric_format':{'kind':'currency','prefix':'$','precision':2},'order':6},
        {'id':'zero','type':'metric','engine':'MetricEngine','element':'Hero KPI','title':'Zero change','value':0,'metric_format':{'kind':'number','precision':0},'order':7},
        {'id':'missing','type':'metric','engine':'MetricEngine','element':'Hero KPI','title':'Missing input','value':None,'metric_format':{'kind':'number','null_display':'N/A'},'order':8},
        {'id':'comparison','type':'metric','engine':'ComparisonEngine','element':'Before/After KPI','title':'Renewal rate · Q1 to Q4','before':87.9,'after':92.9,'unit':'%','order':9},
        {'id':'risk','type':'text','engine':'DecisionCompositeEngine','element':'Risk Callout','title':'Risk / decision','statement':'Hold public-sector expansion until the renewal gate clears.','detail':'Two delayed procurement cycles remain unsigned; keep the current enterprise plan while finance verifies both contracts.','status':'Watch','order':10},
        {'id':'next','type':'text','engine':'TextEngine','element':'Key Takeaway','title':'Next step','text':'Next step: release the public-sector expansion budget only after the named renewals are signed and the matched cohort is reconciled.','order':11},
    ]
    model=canonical_model({'items':items,'nextId':len(items)+1,'canvas':{'width':1600,'height':1200}})
    rectangles=[
        (20,20,1560,130),(20,170,1560,90),(20,280,1560,210),
        (20,510,370,140),(410,510,370,140),(800,510,370,140),(1190,510,370,140),
        (20,670,370,110),(410,670,370,110),(800,670,370,110),(20,810,1560,110),(20,950,1560,160),
    ]
    geometry={'canvas':{'width':1600,'height':1200},'items':[{'id':item['id'],'x':rect[0],'y':rect[1],'w':rect[2],'h':rect[3]} for item,rect in zip(items,rectangles)]}
    return model,geometry


def _semantic_entry(shape):
    values=shape._element.xpath('.//p:cNvPr')
    if not values:
        return None
    description=values[0].get('descr') or ''
    if not description.startswith('VisualizerSemantic:'):
        return None
    return json.loads(description.removeprefix('VisualizerSemantic:'))


def _contrast_ratio(rgb: tuple[int,int,int], background: tuple[int,int,int]) -> float:
    def luminance(color):
        channels=[]
        for value in color:
            channel=value/255
            channels.append(channel/12.92 if channel<=.04045 else ((channel+.055)/1.055)**2.4)
        return .2126*channels[0]+.7152*channels[1]+.0722*channels[2]
    first,second=sorted((luminance(rgb),luminance(background)),reverse=True)
    return (first+.05)/(second+.05)


def test_chg206_report_text_is_explicit_readable_and_contained():
    model,geometry=_long_report()
    payload=export_pptx(None,model,layout_geometry=geometry)
    deck=Presentation(io.BytesIO(payload))
    slide=deck.slides[0]
    assigned={item['id']:item for item in geometry['items']}
    scale_x=deck.slide_width/geometry['canvas']['width']
    scale_y=deck.slide_height/geometry['canvas']['height']
    found={}
    for shape in slide.shapes:
        entry=_semantic_entry(shape)
        if entry is None or not shape.has_text_frame:
            continue
        item_id=entry['id']
        found[item_id]=shape
        rect=assigned[item_id]
        assert abs(shape.left-round(rect['x']*scale_x))<=1
        assert abs(shape.top-round(rect['y']*scale_y))<=1
        assert abs(shape.width-round(rect['w']*scale_x))<=1
        assert abs(shape.height-round(rect['h']*scale_y))<=1
        assert shape.left>=0 and shape.top>=0
        assert shape.left+shape.width<=deck.slide_width and shape.top+shape.height<=deck.slide_height

        frame=shape.text_frame
        assert frame.word_wrap is True
        assert frame.vertical_anchor==MSO_ANCHOR.TOP
        assert frame.auto_size==MSO_AUTO_SIZE.NONE
        assert (frame.margin_left,frame.margin_right,frame.margin_top,frame.margin_bottom)==(
            round(.03*914400),round(.03*914400),0,0
        )
        assert tuple(shape.fill.fore_color.rgb)==(255,255,255)
        metadata=shape._element.xpath('.//p:cNvPr')[0].get('title') or ''
        assert metadata.startswith('Visembler report text roles: ')
        assert len(shape.text_frame.paragraphs)>0
        for paragraph in frame.paragraphs:
            spacing=paragraph._p.pPr.find(qn('a:lnSpc'))
            assert spacing is not None and spacing.find(qn('a:spcPct')) is not None
            assert paragraph.line_spacing is not None
            for run in paragraph.runs:
                assert run.font.name in {'Arial','Arial Narrow'}
                assert run.font.size is not None and run.font.size.pt>=8.0
                if run.font.name=='Arial Narrow':
                    assert 'context subhead' in metadata
                assert run.font.bold in (True,False)
                assert run.font.italic is False
                assert run.font.color.type==MSO_COLOR_TYPE.RGB
                assert _contrast_ratio(tuple(run.font.color.rgb),(255,255,255))>=7
                rpr=run._r.rPr
                assert rpr is not None
                solid=rpr.find(qn('a:solidFill'))
                assert solid is not None and solid.find(qn('a:srgbClr')) is not None
                latin=rpr.find(qn('a:latin'))
                assert latin is not None and latin.get('typeface')==run.font.name
        body=frame._txBody.bodyPr
        assert body.get('wrap')=='square' and body.get('anchor')=='t'
        assert body.find(qn('a:noAutofit')) is not None

    assert set(found)==set(assigned)
    rectangles=list(geometry['items'])
    for index,first in enumerate(rectangles):
        for second in rectangles[index+1:]:
            overlap_w=min(first['x']+first['w'],second['x']+second['w'])-max(first['x'],second['x'])
            overlap_h=min(first['y']+first['h'],second['y']+second['h'])-max(first['y'],second['y'])
            assert not (overlap_w>1 and overlap_h>1)
    assert 'report headline' in found['headline']._element.xpath('.//p:cNvPr')[0].get('title')
    assert 'narrative interpretation' in found['narrative']._element.xpath('.//p:cNvPr')[0].get('title')
    assert 'risk decision' in found['risk']._element.xpath('.//p:cNvPr')[0].get('title')
    assert 'conclusion next step' in found['next']._element.xpath('.//p:cNvPr')[0].get('title')
    assert '87.9% → 92.9%' in found['comparison'].text
    assert '27.6%' in found['percent'].text
    assert '$1,234,567,890.12' in found['currency'].text
    assert '34 contacts' in found['unit'].text
    assert '-$2.35' in found['negative'].text
    assert '\n0\n' in f'\n{found["zero"].text}\n'
    assert 'N/A' in found['missing'].text

    headline_sizes=[run.font.size.pt for paragraph in found['headline'].text_frame.paragraphs for run in paragraph.runs]
    narrative_sizes=[run.font.size.pt for paragraph in found['narrative'].text_frame.paragraphs for run in paragraph.runs]
    assert min(headline_sizes)<=13.5
    assert min(narrative_sizes)>=9.5
    assert import_visembler_pptx(payload)['items']==model['items']


def test_compact_r2_context_rectangle_fits_at_readable_minimum():
    entry={'id':'context','type':'text','engine':'TextEngine','element':'Body Narrative','title':'Quarterly context','text':'FY26 Q1–Q4 close · enterprise and public sector portfolios · constant-currency bookings · renewal cohort excludes acquired accounts.','order':0}
    model=canonical_model({'items':[entry],'nextId':2,'canvas':{'width':1600,'height':900}})
    geometry={'canvas':{'width':1600,'height':900},'items':[{'id':'context','x':20,'y':20,'w':499,'h':72}]}
    deck=Presentation(io.BytesIO(export_pptx(None,model,layout_geometry=geometry)))
    shape=next(value for value in deck.slides[0].shapes if _semantic_entry(value)['id']=='context')
    assert entry['text'] in shape.text
    assert shape.text_frame.margin_left==round(.03*914400)
    assert shape.text_frame.margin_right==round(.03*914400)
    assert shape.text_frame.margin_top==0
    assert shape.text_frame.margin_bottom==0
    sizes=[run.font.size.pt for paragraph in shape.text_frame.paragraphs for run in paragraph.runs]
    assert min(sizes)>=9.5


def test_dense_long_narrative_uses_bounded_editable_fit_without_clipping():
    reports=json.loads((ROOT/'tests/fixtures/chg206/whole_report_models.json').read_text())['reports']
    report=next(value for value in reports if value['key']=='technical-status-review')
    entry=next(value for value in report['model']['items'] if value['id']=='c12')
    model=canonical_model({'items':[entry],'nextId':2,'canvas':{'width':1600,'height':900}})
    geometry={'canvas':{'width':1600,'height':900},'items':[{'id':entry['id'],'x':10,'y':10,'w':530,'h':82}]}
    deck=Presentation(io.BytesIO(export_pptx(None,model,layout_geometry=geometry)))
    shape=next(value for value in deck.slides[0].shapes if _semantic_entry(value)['id']==entry['id'])
    role_metadata=shape._element.xpath('.//p:cNvPr')[0].get('title') or ''
    assert entry['text'] in shape.text
    assert 'Body Narrative' in shape.text
    assert 'text heading' in role_metadata and 'narrative interpretation' in role_metadata
    assert len(shape.text_frame.paragraphs)==1
    runs=shape.text_frame.paragraphs[0].runs
    assert [run.font.name for run in runs]==['Arial','Arial Narrow']
    assert min(run.font.size.pt for run in runs)>=9.5
    assert shape.text_frame.word_wrap is True
    assert shape.text_frame.vertical_anchor==MSO_ANCHOR.TOP
    assert shape.text_frame.auto_size==MSO_AUTO_SIZE.NONE


def test_tall_slide_headline_uses_readable_condensed_fallback_inside_assigned_box():
    headline='Shift B recovered yield after chamber maintenance. Keep the recipe locked through the next paired-lot check and hand off the evidence to Quality and Equipment Engineering.'
    entry={'id':'headline','type':'text','engine':'TextEngine','element':'Executive Statement','title':'Executive Statement','text':headline,'body':headline,'order':0}
    model=canonical_model({'items':[entry],'nextId':2,'canvas':{'width':1600,'height':2356}})
    geometry={'canvas':{'width':1600,'height':2356},'items':[{'id':'headline','x':14,'y':40,'w':1059.44,'h':120.96}]}
    deck=Presentation(io.BytesIO(export_pptx(None,model,layout_geometry=geometry)))
    shape=next(value for value in deck.slides[0].shapes if _semantic_entry(value)['id']=='headline')
    roles=shape._element.xpath('.//p:cNvPr')[0].get('title') or ''
    assert headline in shape.text
    assert 'Executive Statement' in shape.text
    assert 'report headline' in roles and 'text heading' in roles
    assert len(shape.text_frame.paragraphs)==1
    headline_runs=[run for run in shape.text_frame.paragraphs[0].runs if 'Shift B recovered' in run.text]
    assert len(headline_runs)==1 and headline_runs[0].font.name=='Arial Narrow'
    assert headline_runs[0].font.size.pt>=13.0
    assert abs(shape.width-round(1059.44*(deck.slide_width/1600)))<=1
    assert abs(shape.height-round(120.96*(deck.slide_height/2356)))<=1


def test_decision_and_action_composite_segments_keep_distinct_roles():
    recommendation={'engine':'DecisionCompositeEngine','element':'Risk Callout','title':'Recommendation','statement':'Extend to 25% exposure with a support-contact stop.','detail':'Stop if support-contact change exceeds +0.15 points.','status':'Limited rollout'}
    assert [role for _,role in _report_text_segments(recommendation,'Recommendation')]==[
        'text_heading','conclusion_next_step','risk_decision','action_status'
    ]
    project_action={'engine':'ProjectCompositeEngine','element':'Project Card','title':'Project Card','statement':'Complete chamber verification and update the runbook.','detail':'Owner: Equipment Engineering · due next shift · evidence: paired control lot.','status':'Active'}
    assert [role for _,role in _report_text_segments(project_action,'Project Card')]==[
        'text_heading','conclusion_next_step','evidence_detail','action_status'
    ]
    containment={'engine':'ProjectCompositeEngine','element':'Corrective action','title':'Corrective action','statement':'Contain WQ-7314 pending chamber requalification.','detail':'Verify endpoint calibration and a matched chamber B rerun.','status':'Contain'}
    assert [role for _,role in _report_text_segments(containment,'Corrective action')]==[
        'text_heading','risk_decision','conclusion_next_step','action_status'
    ]
