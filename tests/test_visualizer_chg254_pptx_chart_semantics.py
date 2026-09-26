from __future__ import annotations

import io
import json
import zipfile
import xml.etree.ElementTree as ET

import pytest
from pptx import Presentation

from company_ui.products.visualizer.domain import VisualizerContractError, canonical_model
from company_ui.products.visualizer.ppt_service import export_pptx, import_visembler_pptx


CHART_NS = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart', 'p': 'http://schemas.openxmlformats.org/presentationml/2006/main'}


def _weekly_report(*, chart_type='Line Chart', title='Conversion by week', axes=None):
    fields = [
        {'id': 'week', 'name': 'Week', 'type': 'categorical'},
        {'id': 'cohort', 'name': 'Cohort', 'type': 'categorical'},
        {'id': 'conversion', 'name': 'Conversion', 'type': 'number', 'unit': '%'},
    ]
    rows = [
        ['W1', 'Control', 8], ['W1', 'Treatment', 8.2],
        ['W2', 'Control', 8.1], ['W2', 'Treatment', 8.7],
        ['W3', 'Control', 8.4], ['W3', 'Treatment', 9.4],
        ['W4', 'Control', 8.5], ['W4', 'Treatment', 9.8],
    ]
    chart = {
        'id': 'c1', 'type': 'chart', 'engine': 'CoreChartEngine', 'element': chart_type,
        'title': title, 'order': 0, 'dataset_id': 'experiment-data',
        'mapping': {'x': 'week', 'category': 'week', 'series': 'cohort', 'y': 'conversion', 'value': 'conversion'},
        'chart_studio': {
            'chart_type': chart_type, 'mapping': {'x': 'week', 'series': 'cohort', 'y': 'conversion'},
            'axes': {'x': {'title': 'Week'}, 'y': {'title': 'Conversion', 'min': None, 'max': None, 'auto': True, 'zeroBaseline': False, 'format': 'auto'}},
            'legend': {'show': True, 'position': 'bottom', 'order': 'input'},
            'series': [
                {'id': 's1', 'key': 'Control', 'field': 'cohort', 'label': 'Control', 'visible': True, 'legend': True, 'axis': 'primary'},
                {'id': 's2', 'key': 'Treatment', 'field': 'cohort', 'label': 'Treatment', 'visible': True, 'legend': True, 'axis': 'primary'},
            ],
            'visual': {'markers': True},
        },
        'x': 30, 'y': 40, 'w': 720, 'h': 380,
    }
    if chart_type == 'Box Plot':
        chart['mapping'] = {'category': 'cohort', 'cohort': 'cohort', 'value': 'conversion', 'x': 'conversion', 'y': 'conversion'}
        chart['chart_studio']['mapping'] = {'category': 'cohort', 'value': 'conversion'}
    if axes is not None:
        chart['chart_studio']['axes']['y'].update(axes)
    return canonical_model({'datasets': [{'id': 'experiment-data', 'name': 'Experiment readings', 'fields': fields, 'rows': rows}], 'items': [chart], 'nextId': 2})


def _package(payload: bytes) -> zipfile.ZipFile:
    return zipfile.ZipFile(io.BytesIO(payload))


def _chart_xml(package: zipfile.ZipFile) -> tuple[str, ET.Element]:
    paths = sorted(name for name in package.namelist() if name.startswith('ppt/charts/chart') and name.endswith('.xml'))
    assert len(paths) == 1
    return paths[0], ET.fromstring(package.read(paths[0]))


def _series(package: zipfile.ZipFile) -> list[dict]:
    _, root = _chart_xml(package)
    result = []
    for node in root.findall('.//c:ser', CHART_NS):
        name = node.find('./c:tx//c:v', CHART_NS)
        categories = node.findall('./c:cat//c:pt', CHART_NS)
        values = node.findall('./c:val//c:pt', CHART_NS)
        result.append({
            'name': name.text if name is not None else None,
            'categories': [(int(point.get('idx')), point.findtext('./c:v', namespaces=CHART_NS)) for point in categories],
            'values': [(int(point.get('idx')), point.findtext('./c:v', namespaces=CHART_NS)) for point in values],
        })
    return result


def test_chg254_experiment_line_is_native_grouped_and_keeps_axis_domain():
    source = _weekly_report()
    payload = export_pptx(None, source)
    deck = Presentation(io.BytesIO(payload))
    chart_shape = next(shape for shape in deck.slides[0].shapes if getattr(shape, 'has_chart', False))
    assert chart_shape.chart.chart_type.name == 'LINE_MARKERS'
    assert chart_shape.chart.chart_title.text_frame.text == 'Conversion by week'
    assert chart_shape.chart.has_legend
    assert [category.label for category in chart_shape.chart.plots[0].categories] == ['W1', 'W2', 'W3', 'W4']
    assert [series.name for series in chart_shape.chart.series] == ['Control', 'Treatment']
    assert [list(series.values) for series in chart_shape.chart.series] == [[8, 8.1, 8.4, 8.5], [8.2, 8.7, 9.4, 9.8]]
    assert chart_shape.chart.value_axis.minimum_scale == pytest.approx(7.91)
    assert chart_shape.chart.value_axis.maximum_scale == pytest.approx(9.89)
    path, root = _chart_xml(_package(payload))
    assert root.find('.//c:lineChart', CHART_NS) is not None
    assert len(root.findall('.//c:ser', CHART_NS)) == 2
    assert path.startswith('ppt/charts/chart')
    assert root.find('.//c:legend', CHART_NS) is not None
    assert source['items'][0]['x'] == 30 and chart_shape.left > 0


def test_chg254_box_plot_is_editable_quartile_primitives_with_one_semantic_item():
    source = _weekly_report(chart_type='Box Plot', title='Outcome distribution by cohort')
    payload = export_pptx(None, source)
    package = _package(payload)
    assert not [name for name in package.namelist() if name.startswith('ppt/charts/chart') and name.endswith('.xml')]
    deck = Presentation(io.BytesIO(payload)); slide = deck.slides[0]
    assert not any(getattr(shape, 'has_chart', False) for shape in slide.shapes)
    assert sum(shape.name.endswith('-quartiles') for shape in slide.shapes) == 2
    assert sum(shape.name.endswith('-median') for shape in slide.shapes) == 2
    categories = {shape.text: json.loads(shape._element.xpath('.//p:cNvPr')[0].get('title'))['quartiles']
                  for shape in slide.shapes if shape.name.endswith('-category')}
    assert categories == {
        'Control': {'n': 4, 'min': 8.0, 'q1': 8.075, 'median': 8.25, 'q3': 8.425, 'max': 8.5},
        'Treatment': {'n': 4, 'min': 8.2, 'q1': 8.575, 'median': 9.05, 'q3': 9.5, 'max': 9.8},
    }
    metadata = [node.get('descr', '') for shape in slide.shapes for node in shape._element.xpath('.//p:cNvPr')
                if (node.get('descr') or '').startswith('VisualizerSemantic:')]
    assert len(metadata) == 1 and json.loads(metadata[0].removeprefix('VisualizerSemantic:'))['id'] == 'c1'
    assert any(shape.text == 'Outcome distribution by cohort' for shape in slide.shapes if getattr(shape, 'has_text_frame', False))


def test_chg254_projection_obeys_series_order_visibility_and_sparse_gaps():
    source = _weekly_report(chart_type='Multi-Line', title='Unrelated renamed dataset')
    chart = source['items'][0]
    chart['chart_studio']['series'] = [
        {'id': 's2', 'key': 'Treatment', 'label': 'Treatment', 'visible': True, 'legend': True, 'axis': 'primary'},
        {'id': 's1', 'key': 'Control', 'label': 'Control', 'visible': True, 'legend': True, 'axis': 'primary'},
        {'id': 's3', 'key': 'Hidden', 'label': 'Hidden', 'visible': False, 'legend': True, 'axis': 'primary'},
    ]
    dataset = source['datasets'][0]
    dataset['fields'][1]['id'] = 'cohort'
    dataset['rows'] = [['W1', 'Control', 8], ['W1', 'Treatment', 8.2], ['W2', 'Treatment', 8.7], ['W3', 'Control', 8.4], ['W4', 'Control', 8.5], ['W4', 'Treatment', 9.8], ['W4', 'Hidden', 30]]
    payload = export_pptx(None, source)
    chart = next(shape.chart for shape in Presentation(io.BytesIO(payload)).slides[0].shapes if getattr(shape, 'has_chart', False))
    assert chart.value_axis.minimum_scale == pytest.approx(7.91)
    assert chart.value_axis.maximum_scale == pytest.approx(9.89)
    facts = _series(_package(payload))
    assert [item['name'] for item in facts] == ['Treatment', 'Control']
    assert [item['categories'] for item in facts] == [
        [(0, 'W1'), (1, 'W2'), (2, 'W3'), (3, 'W4')],
        [(0, 'W1'), (1, 'W2'), (2, 'W3'), (3, 'W4')],
    ]
    assert facts[0]['values'] == [(0, '8.2'), (1, '8.7'), (3, '9.8')]
    assert facts[1]['values'] == [(0, '8.0'), (2, '8.4'), (3, '8.5')]


def test_chg254_explicit_axis_wins_and_zero_baseline_bar_keeps_zero():
    explicit = _weekly_report(axes={'min': 7.5, 'max': 10.25, 'auto': False})
    explicit_chart = next(shape.chart for shape in Presentation(io.BytesIO(export_pptx(None, explicit))).slides[0].shapes if getattr(shape, 'has_chart', False))
    assert explicit_chart.value_axis.minimum_scale == 7.5
    assert explicit_chart.value_axis.maximum_scale == 10.25
    bar = _weekly_report(chart_type='Vertical Bar', axes={'min': None, 'max': None, 'auto': True, 'zeroBaseline': True})
    bar_chart = next(shape.chart for shape in Presentation(io.BytesIO(export_pptx(None, bar))).slides[0].shapes if getattr(shape, 'has_chart', False))
    assert bar_chart.chart_type.name == 'COLUMN_CLUSTERED'
    assert bar_chart.value_axis.minimum_scale == 0
    assert bar_chart.value_axis.maximum_scale > 9.8
    area = _weekly_report(chart_type='Area Chart')
    area_chart = next(shape.chart for shape in Presentation(io.BytesIO(export_pptx(None, area))).slides[0].shapes if getattr(shape, 'has_chart', False))
    assert area_chart.chart_type.name == 'AREA'
    horizontal = _weekly_report(chart_type='Horizontal Bar')
    horizontal_chart = next(shape.chart for shape in Presentation(io.BytesIO(export_pptx(None, horizontal))).slides[0].shapes if getattr(shape, 'has_chart', False))
    assert horizontal_chart.chart_type.name == 'BAR_CLUSTERED'


def test_chg254_scatter_preserves_numeric_xy_and_unsupported_family_fails_explicitly():
    model = canonical_model({'datasets': [{'id': 'd', 'fields': [
        {'id': 'x', 'name': 'Temperature', 'type': 'number'}, {'id': 'y', 'name': 'Pressure', 'type': 'number'}],
        'rows': [[20, 1.1], [21, 1.3], [22, 1.25]]}], 'items': [
        {'id': 'c1', 'type': 'chart', 'engine': 'CoreChartEngine', 'element': 'Scatter Plot', 'title': 'Relationship',
         'dataset_id': 'd', 'mapping': {'x': 'x', 'y': 'y'}, 'order': 0}], 'nextId': 2})
    payload = export_pptx(None, model); deck = Presentation(io.BytesIO(payload))
    chart = next(shape.chart for shape in deck.slides[0].shapes if getattr(shape, 'has_chart', False))
    assert chart.chart_type.name == 'XY_SCATTER'
    _, scatter_xml = _chart_xml(_package(payload))
    assert scatter_xml.find('.//c:scatterStyle', CHART_NS).get('val') == 'marker'
    assert [point.findtext('./c:v', namespaces=CHART_NS) for point in scatter_xml.findall('.//c:xVal//c:pt', CHART_NS)] == ['20.0', '21.0', '22.0']
    assert [point.findtext('./c:v', namespaces=CHART_NS) for point in scatter_xml.findall('.//c:yVal//c:pt', CHART_NS)] == ['1.1', '1.3', '1.25']
    for chart_type in ('Pareto', 'I-MR Chart', 'CUSUM Chart', 'EWMA Chart', 'Xbar-R Chart', 'DOE Interaction Plot'):
        unsupported = _weekly_report(chart_type=chart_type)
        with pytest.raises(VisualizerContractError, match=f'does not support chart family {chart_type}'):
            export_pptx(None, unsupported)


def test_chg254_histogram_projects_bound_measurements_to_editable_frequency_bins():
    model = canonical_model({'datasets': [{'id': 'd', 'fields': [
        {'id': 'measurement', 'name': 'Measurement', 'type': 'number'}], 'rows': [[1], [1], [2], [3], [5], [8]]}],
        'items': [{'id': 'c1', 'type': 'chart', 'engine': 'CoreChartEngine', 'element': 'Histogram',
                   'title': 'Measurement distribution', 'dataset_id': 'd', 'mapping': {'value': 'measurement'}, 'order': 0}],
        'nextId': 2})
    payload = export_pptx(None, model)
    deck = Presentation(io.BytesIO(payload))
    chart = next(shape.chart for shape in deck.slides[0].shapes if getattr(shape, 'has_chart', False))
    assert chart.chart_type.name == 'COLUMN_CLUSTERED'
    assert [category.label for category in chart.plots[0].categories] == [
        '1–2.16666666667', '2.16666666667–3.33333333333', '3.33333333333–4.5',
        '4.5–5.66666666667', '5.66666666667–6.83333333333', '6.83333333333–8',
    ]
    assert list(chart.series[0].values) == [3, 1, 0, 1, 0, 1]


def test_chg254_round_trip_preserves_original_report_model_and_geometry():
    source = _weekly_report(chart_type='Box Plot')
    restored = import_visembler_pptx(export_pptx(None, source))
    assert restored == source
