#!/usr/bin/env python3
"""Native Visembler acceptance for CHG-180 R1 Process Flow and wafer exports."""
from __future__ import annotations

import argparse
import copy
import hashlib
import importlib.metadata
import json
import os
import shutil
import subprocess
import sys
import tempfile
import time
import traceback
from pathlib import Path
from urllib.parse import quote

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from PIL import Image, ImageDraw, ImageFont
from playwright.sync_api import sync_playwright

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(Path(__file__).parent))

from company_ui.products.visualizer.domain import canonical_model
from company_ui.products.visualizer.ppt_service import import_visembler_pptx
from company_ui.products.visualizer.templates import REPORT_TEMPLATES
from editor_host import NativeHost
from native_common import BrowserEvents, browser_kwargs, ready, write_json

BASE_SHA = '797f874c5d5c5c7f522674b4a08cdd08bdbf31db'
BASE_TREE = 'e279190e4709e4914054e7098a0d1421723c8a6e'
SOURCE_FILES = (
    'company_ui/products/visualizer/assets/authoring_diagram_studio.mjs',
    'company_ui/products/visualizer/assets/integrated_editor.mjs',
    'company_ui/products/visualizer/assets/integrated_editor.html',
    'company_ui/products/visualizer/page.py',
    'company_ui/products/visualizer/ppt_service.py',
    'tests/test_visualizer_chg180_r1.py',
    'scripts/release_checks/run_chg180_r1_acceptance.py',
)
FLOW_LABELS = [
    'Define affected wafer scope and customer impact',
    'Contain material from the suspect process window',
    'Verify measurement method against qualified standards',
    'Compare affected and reference wafer populations',
    'Inspect etch tool and chamber process history',
    'Correlate recipe settings with die-level yield loss',
    'Confirm the physical defect signature by inspection',
    'Apply corrective action and verify effectiveness',
    'Release production after engineering approval',
]


def operations_model() -> dict:
    model = copy.deepcopy(REPORT_TEMPLATES['operations-review']['model'])
    items = {entry['id']: entry for entry in model['items']}
    items['c1'].update({'value': 94.2, 'unit': '%', 'target': 95.0, 'delta': -1.1})
    items['c2'].update({'data': [['W1', 88.1], ['W2', 90.4], ['W3', 92.7], ['W4', 94.2]], 'rows': [{'label': f'W{i+1}', 'value': value} for i, value in enumerate((88.1, 90.4, 92.7, 94.2))]})
    items['c3']['customTable'] = {'headers': ['Action', 'Owner', 'Due', 'Status'], 'rows': [['Verify etch chamber', 'Process Eng.', 'Sep 30', 'Open'], ['Review inline sampling', 'Quality', 'Oct 02', 'In progress']]}
    items['c4'].update({'statement': 'Etch chamber drift is the leading operational risk', 'detail': 'Hold affected W07 material until die-level deltas return inside the qualified window.', 'status': 'Containment active'})
    items['c5'].update({'statement': 'W07 yield recovery', 'detail': 'Owner · Process Engineering · verify before the next production release.', 'status': 'Verification'})
    model['title'] = 'CHG-180 Dense Operations Review'
    model['description'] = 'Dense operations scorecard and editable causal Process Flow regression.'
    return canonical_model(model)


def rca_model() -> dict:
    fields = [
        {'id': 'die_x', 'name': 'Die X', 'type': 'integer'}, {'id': 'die_y', 'name': 'Die Y', 'type': 'integer'},
        {'id': 'yield', 'name': 'Yield', 'type': 'number'}, {'id': 'reference', 'name': 'Reference', 'type': 'number'},
        {'id': 'affected', 'name': 'Affected', 'type': 'number'}, {'id': 'wafer', 'name': 'Wafer', 'type': 'identifier'},
        {'id': 'lot', 'name': 'Lot', 'type': 'identifier'}, {'id': 'tool', 'name': 'Tool', 'type': 'categorical'},
        {'id': 'chamber', 'name': 'Chamber', 'type': 'categorical'}, {'id': 'recipe', 'name': 'Recipe', 'type': 'categorical'},
        {'id': 'process', 'name': 'Process', 'type': 'categorical'},
    ]
    rows = [
        [0, 0, 98.1, 100.0, 100.0, 'W07', 'L2407', 'ETCH-08', 'A', 'REC-17', 'Etch'],
        [1, 0, 96.0, 100.0, 95.0, 'W07', 'L2407', 'ETCH-08', 'A', 'REC-17', 'Etch'],
        [0, 1, 0, 100.0, 102.0, 'W07', 'L2407', 'ETCH-08', 'A', 'REC-17', 'Etch'],
        [1, 1, None, 102.0, None, 'W07', 'L2407', 'ETCH-08', 'A', 'REC-17', 'Etch'],
    ]
    identity = {'wafer_id': 'wafer', 'lot_id': 'lot', 'tool': 'tool', 'chamber': 'chamber', 'recipe': 'recipe', 'process': 'process'}
    return canonical_model({
        'mode': 'smart', 'layoutPreset': 'technical', 'canvas': {'width': 1600, 'height': 900}, 'nextId': 4,
        'title': 'CHG-180 Semiconductor RCA · W07 / L2407',
        'description': 'Etch chamber yield excursion with die-level affected/reference comparison.',
        'datasets': [{'id': 'd-wafer', 'name': 'W07 etch yield and reference data', 'fields': fields, 'rows': rows}],
        'items': [
            {'id': 'c1', 'type': 'text', 'engine': 'TextEngine', 'element': 'Key Takeaway', 'title': 'Etch chamber RCA', 'text': 'W07 on lot L2407 shows a localized die-level yield loss after ETCH-08 chamber A processing. Negative affected-reference deltas cluster in the upper-right die region; confirm chamber condition and requalify REC-17 before release.', 'order': 0, 'weight': 1.0, 'locked': False, 'z': 1},
            {'id': 'c2', 'type': 'wafer', 'engine': 'WaferFabEngine', 'element': 'Wafer Map', 'title': 'W07 affected yield by die', 'order': 1, 'dataset_id': 'd-wafer', 'mapping': {'die_x': 'die_x', 'die_y': 'die_y', 'value': 'yield', **identity}, 'locked': False, 'z': 2},
            {'id': 'c3', 'type': 'wafer', 'engine': 'WaferFabEngine', 'element': 'Wafer Difference Map', 'title': 'W07 signed affected − reference', 'order': 2, 'dataset_id': 'd-wafer', 'mapping': {'die_x': 'die_x', 'die_y': 'die_y', 'reference_value': 'reference', 'affected_value': 'affected', **identity}, 'locked': False, 'z': 3},
        ],
    })


def editor(page, host: NativeHost, report_id: str) -> None:
    response = page.goto(f'{host.url}/visualizer?report={quote(report_id)}', wait_until='domcontentloaded')
    assert response and response.status == 200, response.status if response else None
    ready(page, require_settled=True)
    page.wait_for_function('()=>window.socket?.connected===true&&window.did_handshake===true', timeout=15_000)


def settled(page) -> None:
    page.wait_for_function('()=>{const s=window.CompanyUIVisualizerBridge?.state?.();return s&&s.pending===0&&!s.inflight}', timeout=20_000)


def repository_record(host: NativeHost, report_id: str, predicate=lambda _: True, timeout: float = 20):
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        record = host.repository.get(report_id)
        if predicate(record):
            return record
        time.sleep(.1)
    raise AssertionError(f'Report {report_id} did not reach the expected persisted state.')


def report_geometry(page, record, flow_id='c6') -> dict:
    browser = page.evaluate('''(flowId)=>{
      const root=document.querySelector('.cui-visualizer-root');
      const component=document.querySelector(`.component[data-id="${flowId}"]`);
      const content=component?.querySelector('.c-content');
      const svg=content?.querySelector('svg.diagram-studio-static');
      const state=window.CompanyUIVisualizerBridge.state();
      const model=state.model;
      const flow=model.items.find(value=>value.id===flowId);
      const measure=node=>node?{clientWidth:node.clientWidth,clientHeight:node.clientHeight,scrollWidth:node.scrollWidth,scrollHeight:node.scrollHeight,offsetWidth:node.offsetWidth,offsetHeight:node.offsetHeight}:null;
      return {revision:state.revision,layout:window.__VIZ_PROD__.layoutGeometry(),preflight:window.__VIZ_PROD__.preflight(),dom:{component:measure(component),content:measure(content),svg:measure(svg)},svgViewBox:svg?.getAttribute('viewBox')||null,model:{mode:model.mode,canvas:model.canvas,flow:flow?{id:flow.id,element:flow.element,direction:flow.direction,nodes:flow.nodes,edges:flow.edges,diagram:flow.diagram||null,itemGeometry:Object.fromEntries(['x','y','w','h','width','height'].filter(key=>flow[key]!==undefined).map(key=>[key,flow[key]]))}:null}};
    }''', flow_id)
    persisted = record.model
    flow = next((entry for entry in persisted.get('items', []) if entry.get('id') == flow_id), None)
    browser['persisted'] = {
        'revision': record.revision,
        'canvas': copy.deepcopy(persisted.get('canvas')),
        'mode': persisted.get('mode'),
        'items': [
            {
                'id': entry.get('id'), 'engine': entry.get('engine'), 'element': entry.get('element'),
                'order': entry.get('order'),
                'geometry': {key: copy.deepcopy(entry[key]) for key in ('x', 'y', 'w', 'h', 'width', 'height') if key in entry},
            }
            for entry in persisted.get('items', [])
        ],
        'flow': None if not flow else {key: copy.deepcopy(flow.get(key)) for key in ('id', 'element', 'direction', 'nodes', 'edges', 'diagram', 'x', 'y', 'w', 'h', 'width', 'height') if key in flow},
    }
    return browser


def save_and_wait(page, host: NativeHost, report_id: str):
    settled(page)
    button = page.locator('#saveBtn')
    page.wait_for_function('()=>document.querySelector("#saveBtn")?.textContent?.trim()==="Autosaved"', timeout=10_000)
    record = host.repository.get(report_id)
    assert record.revision >= 1
    assert button.inner_text().strip() == 'Autosaved'
    return record


def export_download(page, destination: Path) -> bytes:
    page.locator('#exportBtn').click()
    page.get_by_role('button', name='Editable PowerPoint').wait_for(state='visible', timeout=8_000)
    downloads = []
    page.on('download', lambda download: downloads.append(download))
    page.locator('#exportPptAction').click()
    deadline = time.monotonic() + 30
    while time.monotonic() < deadline and not downloads:
        page.wait_for_timeout(100)
    if not downloads:
        diagnostics = page.evaluate('''()=>({debug:window.__VIZ_PROD__?.ui?.debugLog?.slice(0,12),modal:document.querySelector('#genericModal.show #modalBody')?.innerText||null,status:document.querySelector('#saveBtn')?.textContent||null,bridge:window.CompanyUIVisualizerBridge?.state?.()})''')
        page.screenshot(path=str(destination.parent / (destination.stem + '-export-error.png')), full_page=True)
        raise AssertionError('The native Editable PowerPoint action did not produce a download: ' + json.dumps(diagnostics, ensure_ascii=False)[:8000])
    downloads[0].save_as(str(destination))
    return destination.read_bytes()


def rasterize_pptx(paths: list[Path], output: Path) -> dict:
    renderer = shutil.which('libreoffice') or shutil.which('soffice')
    pdf_tool = shutil.which('pdftoppm')
    if not renderer or not pdf_tool:
        return {'status': 'MISSING', 'reason': 'LibreOffice/soffice and pdftoppm are required for trusted headless PowerPoint rasterization.', 'available': {'libreoffice_or_soffice': renderer, 'pdftoppm': pdf_tool}}
    renders = output / 'rasterized-slides'
    renders.mkdir(parents=True, exist_ok=True)
    PDFs = output / 'renderer-pdf'
    PDFs.mkdir(parents=True, exist_ok=True)
    profile = output / 'renderer-profile'
    profile.mkdir(parents=True, exist_ok=True)
    images = []
    for deck in paths:
        prefix = deck.stem
        converted = subprocess.run([renderer, '--headless', f'-env:UserInstallation={profile.as_uri()}', '--convert-to', 'pdf', '--outdir', str(PDFs), str(deck)], capture_output=True, text=True, timeout=120)
        if converted.returncode != 0:
            return {'status': 'FAIL', 'renderer': renderer, 'error': converted.stderr or converted.stdout}
        pdf = PDFs / f'{prefix}.pdf'
        if not pdf.is_file():
            return {'status': 'FAIL', 'renderer': renderer, 'error': 'Headless renderer did not produce the expected PDF.'}
        rendered = subprocess.run([pdf_tool, '-png', '-r', '144', str(pdf), str(renders / prefix)], capture_output=True, text=True, timeout=120)
        if rendered.returncode != 0:
            return {'status': 'FAIL', 'renderer': renderer, 'error': rendered.stderr or rendered.stdout}
        images.extend(sorted(renders.glob(prefix + '-*.png')))
    if not images:
        return {'status': 'FAIL', 'renderer': renderer, 'error': 'PPTX rasterization produced no slide images.'}
    return {'status': 'PASS', 'renderer': renderer, 'images': [str(path.relative_to(output)) for path in images]}


def ui_contact_sheet(paths: list[Path], destination: Path) -> None:
    font = ImageFont.load_default()
    tiles = []
    for path in paths:
        if not path.is_file():
            continue
        source = Image.open(path).convert('RGB')
        source.thumbnail((640, 430))
        tile = Image.new('RGB', (660, 470), '#f3f5f8')
        tile.paste(source, ((660 - source.width) // 2, 28))
        ImageDraw.Draw(tile).text((12, 8), path.name, fill='#14243a', font=font)
        tiles.append(tile)
    if tiles:
        columns = 2
        rows = (len(tiles) + columns - 1) // columns
        sheet = Image.new('RGB', (columns * 660, rows * 470), 'white')
        for index, tile in enumerate(tiles):
            sheet.paste(tile, ((index % columns) * 660, (index // columns) * 470))
        destination.parent.mkdir(parents=True, exist_ok=True)
        sheet.save(destination)


def structural_operations(path: Path) -> dict:
    presentation = Presentation(str(path))
    shapes = [shape for slide in presentation.slides for shape in slide.shapes]
    nodes = [shape for shape in shapes if shape.name.startswith('VIZ::Process Flow::')]
    labels = [shape.text for shape in nodes if getattr(shape, 'has_text_frame', False)]
    connectors = [shape for shape in shapes if shape.shape_type == MSO_SHAPE_TYPE.LINE]
    semantic = []
    for shape in nodes:
        for node in shape._element.xpath('.//p:cNvPr'):
            description = node.get('descr') or ''
            if description.startswith('VisualizerSemantic:'):
                semantic.append(json.loads(description.removeprefix('VisualizerSemantic:')))
    assert len(nodes) == len(FLOW_LABELS), f'Expected {len(FLOW_LABELS)} editable flow nodes, got {len(nodes)}.'
    assert len(connectors) >= len(FLOW_LABELS) - 1, f'Expected causal connectors, got {len(connectors)}.'
    normalized_labels = ' '.join('\n'.join(labels).split())
    assert all(' '.join(label.split()) in normalized_labels for label in FLOW_LABELS), 'PowerPoint flow labels are incomplete.'
    assert semantic and semantic[0].get('direction') == 'down' and len(semantic[0].get('edges') or []) == len(FLOW_LABELS) - 1
    return {'slide_count': len(presentation.slides), 'editable_node_count': len(nodes), 'connector_count': len(connectors), 'node_labels': labels, 'canonical_flow_metadata': semantic[0]}


def structural_rca(path: Path, expected_model: dict) -> dict:
    presentation = Presentation(str(path))
    shapes = [shape for slide in presentation.slides for shape in slide.shapes]
    result = {'slide_count': len(presentation.slides), 'slides': []}
    for title, variant, signed in (
        ('W07 affected yield by die', 'wafer_map', False),
        ('W07 signed affected − reference', 'wafer_difference', True),
    ):
        prefix = f'VIZ::{title}::'
        panel = next(shape for shape in shapes if shape.name == prefix + 'spatial-panel')
        outline = next(shape for shape in shapes if shape.name == prefix + 'wafer-outline')
        notch = next(shape for shape in shapes if shape.name == prefix + 'orientation-notch')
        cells = [shape for shape in shapes if shape.name.startswith(prefix + 'die-')]
        assert len(cells) == 4
        descriptions = [node.get('descr') or '' for node in panel._element.xpath('.//p:cNvPr')]
        payload = json.loads(next(value.removeprefix('VisualizerSemantic:') for value in descriptions if value.startswith('VisualizerSemantic:')))
        assert payload['element'] == ('Wafer Difference Map' if signed else 'Wafer Map')
        expected_ids = {'wafer_id': 'W07', 'lot_id': 'L2407', 'tool': 'ETCH-08', 'chamber': 'A', 'recipe': 'REC-17', 'process': 'Etch'}
        assert all(payload.get(key) == value for key, value in expected_ids.items())
        die_payloads = []
        for cell in cells:
            description = next(node.get('descr') or '' for node in cell._element.xpath('.//p:cNvPr') if (node.get('descr') or '').startswith('VisualizerSemanticDie:'))
            die_payloads.append(json.loads(description.removeprefix('VisualizerSemanticDie:')))
        if signed:
            assert [row.get('delta') for row in sorted(die_payloads, key=lambda row: row['index'])] == [0.0, -5.0, 2.0, None]
            assert 'SIGNED DELTA · AFFECTED − REFERENCE' in '\n'.join(shape.text for shape in shapes if getattr(shape, 'has_text_frame', False))
        else:
            assert [row.get('value') for row in sorted(die_payloads, key=lambda row: row['index'])] == [98.1, 96.0, 0, None]
        expected_item = next(entry for entry in expected_model['items'] if entry['element'] == ('Wafer Difference Map' if signed else 'Wafer Map'))
        result['slides'].append({'kind': variant, 'title': title, 'editable_die_count': len(cells), 'outline': outline.name, 'notch': notch.name, 'panel': panel.name, 'observations': payload['observations'], 'metadata': {key: payload[key] for key in expected_ids}})
    assert not any(getattr(shape, 'has_text_frame', False) and 'Controlled fallback' in shape.text for shape in shapes)
    imported = import_visembler_pptx(path.read_bytes())
    assert imported and {entry['id'] for entry in imported['items']} == {entry['id'] for entry in expected_model['items']}
    result['round_trip_item_ids'] = [entry['id'] for entry in imported['items']]
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', required=True, type=Path)
    parser.add_argument('--candidate-sha', help='Require capture from this committed candidate SHA.')
    parser.add_argument('--candidate-tree', help='Require capture from this committed candidate tree.')
    args = parser.parse_args()
    if bool(args.candidate_sha) != bool(args.candidate_tree):
        parser.error('--candidate-sha and --candidate-tree must be provided together')
    output = args.output.expanduser().resolve()
    shots = output / 'screenshots'
    exports = output / 'powerpoints'
    shots.mkdir(parents=True, exist_ok=True)
    exports.mkdir(parents=True, exist_ok=True)
    head = subprocess.check_output(['git', 'rev-parse', 'HEAD'], cwd=ROOT, text=True).strip()
    tree = subprocess.check_output(['git', 'rev-parse', 'HEAD^{tree}'], cwd=ROOT, text=True).strip()
    source_tree = subprocess.check_output(['git', 'write-tree'], cwd=ROOT, text=True).strip()
    receipt = {
        'schema': 'visembler-chg180-r1-native-acceptance.v1', 'status': 'FAIL',
        'requested_base_sha': BASE_SHA, 'requested_base_tree': BASE_TREE,
        'candidate_sha': args.candidate_sha, 'candidate_tree': args.candidate_tree,
        'capture_start_head': head, 'capture_start_tree': tree,
        'source_tree': source_tree,
        'source_files_sha256': {path: hashlib.sha256((ROOT / path).read_bytes()).hexdigest() for path in SOURCE_FILES},
        'runtime': {'python': sys.version.split()[0], 'nicegui': importlib.metadata.version('nicegui'), 'playwright': importlib.metadata.version('playwright'), 'python_pptx': importlib.metadata.version('python-pptx')},
        'cases': [], 'unexpected_browser_events': [],
        'headless_pptx_renderer': {},
    }
    events = BrowserEvents()
    try:
        if args.candidate_sha:
            assert head == args.candidate_sha, f'Acceptance started on {head}, expected candidate {args.candidate_sha}.'
            assert tree == args.candidate_tree, f'Acceptance started on tree {tree}, expected candidate tree {args.candidate_tree}.'
            assert source_tree == tree, f'Committed candidate tree {tree} differs from captured source tree {source_tree}.'
            dirty = subprocess.check_output(['git', 'status', '--porcelain'], cwd=ROOT, text=True).strip()
            assert not dirty, 'Candidate-mode acceptance requires a clean worktree.'
        else:
            assert head == BASE_SHA, f'Acceptance started on {head}, expected exact base {BASE_SHA}.'
            assert tree == BASE_TREE, f'Acceptance started on tree {tree}, expected exact tree {BASE_TREE}.'
        with tempfile.TemporaryDirectory(prefix='visembler-chg180-r1-') as temporary:
            with NativeHost(ROOT, Path(temporary) / 'data') as host:
                ops_id = host.create(model=operations_model(), name='chg180-r1-dense-operations')
                rca = rca_model()
                original_rca = copy.deepcopy(rca)
                rca_id = host.create(model=rca, name='chg180-r1-semiconductor-rca')
                with sync_playwright() as playwright:
                    browser = playwright.chromium.launch(**browser_kwargs())
                    receipt['runtime']['browser'] = {'name': 'Chromium', 'version': browser.version}
                    context = browser.new_context(accept_downloads=True, viewport={'width': 1600, 'height': 1120}, device_scale_factor=1)
                    page = context.new_page(); events.attach(page); editor(page, host, ops_id)
                    if not page.evaluate('()=>window.__VIZ_PROD__.ui.libraryOpen'):
                        page.locator('#libraryToggle').click()
                    page.locator('#componentSearch').fill('Process Flow')
                    flow_library = page.locator('.library-item[data-element="Process Flow"][data-engine="DiagramEngine"] .library-insert').first
                    flow_library.wait_for(state='visible', timeout=10_000); flow_library.click()
                    page.locator('#iNodes').wait_for(state='visible', timeout=10_000)
                    page.locator('#iNodes').fill('\n'.join(FLOW_LABELS)); page.locator('#iNodes').press('Tab'); settled(page)
                    page.locator('#iEdges').fill('\n'.join(f'{source} -> {target}' for source, target in zip(FLOW_LABELS, FLOW_LABELS[1:]))); page.locator('#iEdges').press('Tab'); settled(page)
                    page.locator('#iDirection').select_option('down'); settled(page)
                    page.wait_for_function('()=>document.querySelectorAll(".component[data-id=\\"c6\\"] .diagram-studio-static [data-diagram-node]").length===9', timeout=15_000)
                    ops_before_record = host.repository.get(ops_id)
                    ops_before = report_geometry(page, ops_before_record)
                    assert ops_before['model']['flow']['nodes'] == FLOW_LABELS
                    assert ops_before['layout']['canvas']['height'] > 900, ops_before['layout']
                    assert ops_before['dom']['content']['scrollHeight'] <= ops_before['dom']['content']['clientHeight'] + 1, ops_before['dom']
                    assert ops_before['preflight']['clipping'] == 0, ops_before['preflight']
                    page.screenshot(path=str(shots / 'operations-authoring-fitted.png'), full_page=True)
                    page.locator('.component[data-id="c6"]').screenshot(path=str(shots / 'operations-authoring-process-flow-closeup.png'))
                    saved_ops = save_and_wait(page, host, ops_id)
                    ops_after_save_record = host.repository.get(ops_id)
                    ops_after_save = report_geometry(page, ops_after_save_record)

                    fresh_context = browser.new_context(accept_downloads=True, viewport={'width': 1600, 'height': 1120}, device_scale_factor=1)
                    fresh = fresh_context.new_page(); events.attach(fresh); editor(fresh, host, ops_id)
                    ops_reloaded_record = host.repository.get(ops_id)
                    ops_reloaded = report_geometry(fresh, ops_reloaded_record)
                    assert ops_before['layout'] == ops_reloaded['layout'], 'Smart Process Flow layout changed after a fresh-context reload.'
                    assert ops_before['persisted'] == ops_reloaded['persisted'], 'Persisted Process Flow content or geometry changed after reload.'
                    assert ops_reloaded['dom']['content']['scrollHeight'] <= ops_reloaded['dom']['content']['clientHeight'] + 1, ops_reloaded['dom']
                    assert ops_reloaded['preflight']['clipping'] == 0, ops_reloaded['preflight']
                    fresh.screenshot(path=str(shots / 'operations-post-save-reload.png'), full_page=True)
                    fresh.locator('.component[data-id="c6"]').screenshot(path=str(shots / 'operations-post-reload-process-flow-closeup.png'))
                    fresh.locator('#preflightBtn').click()
                    fresh.locator('#modalTitle').wait_for()
                    visible_validation = fresh.locator('#modalTitle').inner_text().strip()
                    fresh.screenshot(path=str(shots / 'operations-visible-validate.png'), full_page=True)
                    assert visible_validation == 'Ready to export', fresh.locator('#modalBody').inner_text()
                    assert not ops_reloaded['preflight']['layoutIssues'] and not ops_reloaded['preflight']['dataIssues'], ops_reloaded['preflight']
                    fresh.keyboard.press('Escape')
                    fresh.locator('#previewBtn').click()
                    fresh.locator('.cui-visualizer-root.preview-mode').wait_for(timeout=8_000)
                    fresh.screenshot(path=str(shots / 'operations-preview.png'), full_page=True)
                    fresh.locator('#previewExit').click()
                    fresh.wait_for_function('()=>!document.querySelector(".cui-visualizer-root")?.classList.contains("preview-mode")')
                    ops_pptx = exports / 'dense-operations-editable.pptx'
                    ops_bytes = export_download(fresh, ops_pptx)
                    ops_structure = structural_operations(ops_pptx)
                    assert host.repository.get(ops_id).model == ops_reloaded_record.model, 'Operations export mutated the saved report model.'
                    receipt['cases'].append({'name': 'Dense Operations Process Flow author, save, fresh reload, visible Validate, Preview and editable PowerPoint', 'status': 'PASS', 'report_id': ops_id, 'revision': ops_reloaded_record.revision, 'geometry_before_save': ops_before, 'geometry_after_save': ops_after_save, 'geometry_after_fresh_reload': ops_reloaded, 'visible_validation_title': visible_validation, 'pptx_bytes': len(ops_bytes), 'pptx_sha256': hashlib.sha256(ops_bytes).hexdigest(), 'powerpoint_structure': ops_structure})

                    third_context = browser.new_context(accept_downloads=True, viewport={'width': 1600, 'height': 1120}, device_scale_factor=1)
                    final_page = third_context.new_page(); events.attach(final_page); editor(final_page, host, ops_id)
                    ops_final = report_geometry(final_page, host.repository.get(ops_id))
                    assert ops_reloaded['layout'] == ops_final['layout'], 'Smart Process Flow layout changed after export and another reload.'
                    assert ops_reloaded['persisted'] == ops_final['persisted']
                    receipt['cases'][-1]['geometry_after_export_reload'] = ops_final

                    rca_page = fresh_context.new_page(); events.attach(rca_page); editor(rca_page, host, rca_id)
                    rca_page.locator('#preflightBtn').click(); rca_page.locator('#modalTitle').wait_for()
                    rca_validation = rca_page.locator('#modalTitle').inner_text().strip()
                    rca_page.screenshot(path=str(shots / 'semiconductor-rca-authoring-validate.png'), full_page=True)
                    assert rca_validation == 'Ready to export', rca_page.locator('#modalBody').inner_text()
                    rca_page.keyboard.press('Escape')
                    saved_rca = save_and_wait(rca_page, host, rca_id)
                    rca_pptx = exports / 'semiconductor-rca-wafer-maps-editable.pptx'
                    rca_bytes = export_download(rca_page, rca_pptx)
                    rca_structure = structural_rca(rca_pptx, original_rca)
                    assert host.repository.get(rca_id).model == original_rca, 'Bound wafer export mutated the canonical report or source dataset.'
                    receipt['cases'].append({'name': 'Semiconductor RCA native Validate and supported PowerPoint action for Wafer Map and Wafer Difference Map', 'status': 'PASS', 'report_id': rca_id, 'revision': saved_rca.revision, 'visible_validation_title': rca_validation, 'source_model_unchanged': True, 'pptx_bytes': len(rca_bytes), 'pptx_sha256': hashlib.sha256(rca_bytes).hexdigest(), 'powerpoint_structure': rca_structure})
                    rca_page.screenshot(path=str(shots / 'semiconductor-rca-post-export.png'), full_page=True)
                    receipt['headless_pptx_renderer'] = rasterize_pptx([ops_pptx, rca_pptx], output)
                    receipt['ui_review_contact_sheet'] = 'ui-review-contact-sheet.png'
                    ui_contact_sheet(sorted(shots.glob('*.png')), output / receipt['ui_review_contact_sheet'])
                    receipt['unexpected_browser_events'] = events.unexpected
                    if host.log_path.exists():
                        (output / 'native-server.log').write_text(host.log_path.read_text(encoding='utf-8', errors='replace'), encoding='utf-8')
                    assert not events.unexpected, events.unexpected
                    third_context.close(); fresh_context.close(); context.close(); browser.close()
    except Exception as exc:
        receipt['error'] = str(exc)
        receipt['traceback'] = traceback.format_exc()
        receipt['unexpected_browser_events'] = events.unexpected
        active_host = locals().get('host')
        if active_host is not None and active_host.log_path.exists():
            (output / 'native-server.log').write_text(active_host.log_path.read_text(encoding='utf-8', errors='replace'), encoding='utf-8')
    receipt['passed_cases'] = sum(value['status'] == 'PASS' for value in receipt['cases'])
    receipt['total_cases'] = 2
    receipt['status'] = 'PASS' if receipt['passed_cases'] == receipt['total_cases'] and not receipt['unexpected_browser_events'] and not receipt.get('error') else 'FAIL'
    write_json(output / 'acceptance-receipt.json', receipt)
    print(json.dumps(receipt, indent=2, ensure_ascii=False))
    return 0 if receipt['status'] == 'PASS' else 1


if __name__ == '__main__':
    raise SystemExit(main())
