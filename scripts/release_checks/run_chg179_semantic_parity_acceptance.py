#!/usr/bin/env python3
"""Native browser acceptance for CHG-179 metric and category semantic parity."""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import shutil
import subprocess
import sys
import tempfile
import traceback
import zipfile
from io import BytesIO
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(Path(__file__).parent))

from company_ui.products.visualizer.domain import canonical_model
from editor_host import NativeHost
from native_common import browser_kwargs, write_json
from run_editor_workflows import add, model, panels, ready, settled
from pptx import Presentation
from playwright.sync_api import sync_playwright


METRIC_CASES = [
    {'name': 'compact-multiplier', 'value': '48200000', 'kind': 'compact', 'divisor': '1000000', 'compact_unit': 'M', 'expected': '48.2 M'},
    {'name': 'millions-unit', 'value': '48.2', 'kind': 'number', 'unit': 'M', 'expected': '48.2 M'},
    {'name': 'percentage-points', 'value': '27.6', 'kind': 'percent', 'unit': '%', 'expected': '27.6%'},
    {'name': 'textual-count-unit', 'value': '34', 'kind': 'number', 'unit': 'contacts', 'expected': '34 contacts'},
    {'name': 'zero-unitless', 'value': '0', 'kind': 'number', 'expected': '0'},
    {'name': 'negative-large-currency', 'value': '-1234567.8', 'kind': 'currency', 'prefix': '$', 'decimals': '2', 'expected': '-$1,234,567.80'},
    {'name': 'explicit-precision', 'value': '3.45678', 'kind': 'number', 'decimals': '3', 'expected': '3.457'},
    {'name': 'missing-value', 'value': '', 'kind': 'number', 'expected': '—'},
]


def state(page):
    return page.evaluate('window.CompanyUIVisualizerBridge.state()')


def mount(page, host, report_id):
    page.goto(f'{host.url}/visualizer?report={report_id}', wait_until='domcontentloaded')
    ready(page)
    page.wait_for_function('()=>{const s=window.CompanyUIVisualizerBridge?.state?.();return s&&s.pending===0&&!s.inflight}', timeout=20000)


def select(page, element_id):
    page.locator(f'.component[data-id="{element_id}"]').evaluate('(node)=>node.click()')
    page.wait_for_function('(id)=>document.querySelector(`.component[data-id="${id}"]`)?.getAttribute("aria-selected")==="true"', arg=element_id)
    panels(page, library=False, inspector=True)


def edit_field(page, selector, value):
    control = page.locator(selector)
    control.wait_for(state='visible')
    control.scroll_into_view_if_needed()
    control.fill(str(value))
    control.press('Tab')
    settled(page)


def choose(page, selector, value):
    page.locator(selector).select_option(value)
    settled(page)


def metric_value(page, element_id):
    return page.locator(f'.component[data-id="{element_id}"] .metric-value').inner_text().strip()


def pptx_text(payload: bytes) -> str:
    presentation = Presentation(BytesIO(payload))
    return '\n'.join(
        paragraph.text
        for slide in presentation.slides
        for shape in slide.shapes
        if getattr(shape, 'has_text_frame', False)
        for paragraph in shape.text_frame.paragraphs
    )


def git_text(*args: str) -> str:
    return subprocess.check_output(['git', *args], cwd=ROOT, text=True).strip()


def export_pptx(page, destination: Path) -> tuple[bytes, str]:
    page.locator('#exportBtn').click()
    with page.expect_download(timeout=30000) as download_info:
        page.locator('#exportPptAction').click()
    download = download_info.value
    download.save_as(str(destination))
    payload = destination.read_bytes()
    return payload, pptx_text(payload)


def capture_metric_case(page, element_id, case, output: Path) -> dict:
    edit_field(page, '#iValue', case['value'])
    if 'unit' in case:
        edit_field(page, '#iUnit', case['unit'])
    if 'kind' in case:
        choose(page, '#iValueFormat', case['kind'])
    if 'prefix' in case:
        edit_field(page, '#iMetricPrefix', case['prefix'])
    if 'divisor' in case:
        edit_field(page, '#iCompactDivisor', case['divisor'])
    if 'compact_unit' in case:
        edit_field(page, '#iCompactUnit', case['compact_unit'])
    if 'decimals' in case:
        edit_field(page, '#iDecimals', case['decimals'])

    editor = metric_value(page, element_id)
    if editor != case['expected']:
        raise AssertionError(f"{case['name']}: editor value {editor!r} != {case['expected']!r}")
    page.screenshot(path=str(output / f"metric-{case['name']}-editor.png"), full_page=True)
    page.locator('#previewBtn').click()
    page.locator('.cui-visualizer-root.preview-mode').wait_for(timeout=8000)
    preview = metric_value(page, element_id)
    page.screenshot(path=str(output / f"metric-{case['name']}-preview.png"), full_page=True)
    page.locator('#previewExit').click()
    page.wait_for_function('()=>!document.querySelector(".cui-visualizer-root")?.classList.contains("preview-mode")')

    pptx_path = output / f"metric-{case['name']}.pptx"
    payload, text = export_pptx(page, pptx_path)
    export_matches = [line.strip() for line in text.splitlines() if case['expected'] in line]
    if not export_matches:
        raise AssertionError(f"{case['name']}: PowerPoint text omitted {case['expected']!r}: {text!r}")
    return {
        'name': case['name'],
        'expected': case['expected'],
        'editor': editor,
        'preview': preview,
        'powerpoint': export_matches[0],
        'parity': editor == preview == export_matches[0],
        'pptx_sha256': hashlib.sha256(payload).hexdigest(),
    }


def run_metric_cases(host, browser, output, receipt):
    results = []
    for case in METRIC_CASES:
        context = browser.new_context(viewport={'width': 1440, 'height': 1000}, accept_downloads=True)
        page = context.new_page()
        page.set_default_timeout(10000)
        report_id = host.create(name=f'chg179-metric-{case["name"]}')
        try:
            mount(page, host, report_id)
            element_id = add(page, 'MetricEngine', 'Hero KPI')
            result = capture_metric_case(page, element_id, case, output)
            result['report_id'] = report_id
            result['status'] = 'PASS' if result['parity'] else 'FAIL'
            results.append(result)
            print(f"{result['status']} metric {case['name']}", flush=True)
            if not result['parity']:
                raise AssertionError(f"{case['name']}: editor, preview, and PowerPoint differ")
        except Exception as exc:
            results.append({'name': case['name'], 'report_id': report_id, 'status': 'FAIL', 'error': str(exc), 'traceback': traceback.format_exc()})
            print(f"FAIL metric {case['name']}: {exc}", flush=True)
        finally:
            context.close()
            receipt['metric_cases'] = results
            write_json(output / 'chg179-semantic-parity.json', receipt)
    return results


def create_data_first_chart(page, source_text):
    panels(page, library=True, inspector=True)
    page.locator('#pasteDataBtn').click()
    page.locator('#dataFirstText').fill(source_text)
    page.locator('.data-first-recommendation').first.wait_for(timeout=15000)
    page.locator('[data-data-first-view="bar"]').click()
    page.locator('#dataFirstCreate').click()
    settled(page)
    page.locator('#genericModal.show').wait_for(state='hidden')
    return model(page)['items'][-1]['id']


def chart_studio_model(page):
    return page.evaluate('()=>JSON.parse(JSON.stringify(window.CompanyUIChartStudio.model))')


def open_chart_studio_from_editor(page, element_id):
    page.locator(f'.component[data-id="{element_id}"] [data-action="edit-chart"]').click()
    page.wait_for_url('**/visualizer/chart-studio**', timeout=15000)
    page.locator('#chart-studio[data-studio-ready="true"]').wait_for(timeout=20000)


def category_labels_in_editor(page, element_id):
    svg = page.locator(f'.component[data-id="{element_id}"] svg.cs-chart-svg')
    return svg.locator('.cs-axis-label[data-full-value]').evaluate_all('(nodes)=>nodes.map(node=>node.getAttribute("data-full-value"))')


def export_chart_pptx_categories(page, destination: Path, expected_labels: list[str]) -> tuple[bytes, list[str]]:
    payload, _ = export_pptx(page, destination)
    with zipfile.ZipFile(BytesIO(payload)) as archive:
        chart_xml = '\n'.join(archive.read(name).decode('utf-8', errors='replace') for name in archive.namelist() if name.startswith('ppt/charts/chart') and name.endswith('.xml'))
    missing = [label for label in expected_labels if f'<c:v>{label}</c:v>' not in chart_xml]
    if missing:
        raise AssertionError(f'PowerPoint chart categories missing {missing!r}')
    return payload, expected_labels


def run_category_workflow(host, browser, output):
    context = browser.new_context(viewport={'width': 1440, 'height': 1000}, accept_downloads=True, permissions=['clipboard-read', 'clipboard-write'])
    page = context.new_page()
    page.set_default_timeout(10000)
    report_id = host.create(name='chg179-data-first-categories')
    labels = ['W0', 'W1', 'W2', 'W3']
    try:
        mount(page, host, report_id)
        element_id = create_data_first_chart(page, 'Week\tValue\nW0\t0\nW1\t2\nW2\t3\nW3\t5')
        created = next(item for item in model(page)['items'] if item['id'] == element_id)
        first_dataset = next(ds for ds in model(page)['datasets'] if ds['id'] == created['dataset_id'])
        category_id = next(field['id'] for field in first_dataset['fields'] if field['name'] == 'Week')
        if created['mapping'].get('category') != category_id:
            raise AssertionError(f'Data First did not retain Week as its category binding: {created["mapping"]!r}')

        initial_editor_labels = category_labels_in_editor(page, element_id)
        if not all(label in initial_editor_labels for label in labels):
            raise AssertionError(f'Initial editor axis lost category labels: {initial_editor_labels!r}')
        page.screenshot(path=str(output / 'categories-editor-before-refresh.png'), full_page=True)

        open_chart_studio_from_editor(page, element_id)
        before = chart_studio_model(page)
        if before['mapping'].get('category') != category_id or not all(label in page.locator('#cs-canvas').inner_text() for label in labels):
            raise AssertionError('Chart Studio did not show the bound week categories')
        page.screenshot(path=str(output / 'categories-chart-studio.png'), full_page=True)
        page.locator('[data-action="save"]').click()
        page.wait_for_function('()=>!window.CompanyUIChartStudio.state.pending', timeout=20000)
        page.reload(wait_until='domcontentloaded')
        page.locator('#chart-studio[data-studio-ready="true"]').wait_for(timeout=20000)
        reloaded = chart_studio_model(page)
        if reloaded['mapping'].get('category') != category_id or reloaded['dataset']['rows'][0][0] != 'W0':
            raise AssertionError('Chart Studio save/reload changed category identity')
        page.locator('[data-action="back"]').click()
        page.wait_for_url('**/visualizer**', timeout=15000)
        ready(page)
        settled(page)

        select(page, element_id)
        page.locator('#inspector [data-inspector="duplicate"]').click()
        settled(page)
        items = model(page)['items']
        duplicate = next(item for item in items if item['id'] != element_id and item.get('dataset_id') == created['dataset_id'])
        if duplicate.get('mapping', {}).get('category') != category_id:
            raise AssertionError('Duplicate lost its category binding')
        duplicate_id = duplicate['id']

        select(page, element_id)
        page.locator('[data-refresh-dataset]').click()
        page.locator('#refreshDataText').fill('Value\tWeek\n0\tW0\n2.5\tW1\n3\tW2\n5\tW3')
        page.wait_for_function('()=>!document.querySelector("#refreshLinked")?.disabled', timeout=15000)
        page.locator('#refreshLinked').click()
        settled(page)
        page.locator('#genericModal.show').wait_for(state='hidden')
        refreshed = model(page)
        current = next(item for item in refreshed['items'] if item['id'] == element_id)
        sibling = next(item for item in refreshed['items'] if item['id'] == duplicate_id)
        refreshed_dataset = next(ds for ds in refreshed['datasets'] if ds['id'] == current['dataset_id'])
        new_category_id = next(field['id'] for field in refreshed_dataset['fields'] if field['name'] == 'Week')
        if current['mapping'].get('category') != new_category_id or sibling['mapping'].get('category') != new_category_id:
            raise AssertionError('Shared dataset refresh did not rebind category fields for each chart')
        if refreshed_dataset['rows'] != [[0, 'W0'], [2.5, 'W1'], [3, 'W2'], [5, 'W3']]:
            raise AssertionError(f'Reordered refresh did not preserve week labels and values: {refreshed_dataset["rows"]!r}')

        post_refresh_labels = category_labels_in_editor(page, element_id)
        if not all(label in post_refresh_labels for label in labels):
            raise AssertionError(f'Refreshed editor axis lost labels: {post_refresh_labels!r}')
        page.screenshot(path=str(output / 'categories-editor-after-refresh.png'), full_page=True)
        page.locator('#previewBtn').click()
        page.locator('.cui-visualizer-root.preview-mode').wait_for(timeout=8000)
        preview_labels = category_labels_in_editor(page, element_id)
        if not all(label in preview_labels for label in labels):
            raise AssertionError(f'Preview axis lost labels: {preview_labels!r}')
        page.screenshot(path=str(output / 'categories-preview.png'), full_page=True)
        page.locator('#previewExit').click()

        open_chart_studio_from_editor(page, element_id)
        after_refresh = chart_studio_model(page)
        if after_refresh['dataset']['rows'] != refreshed_dataset['rows'] or after_refresh['mapping'].get('category') != new_category_id:
            raise AssertionError('Chart Studio reopened stale embedded rows or category mapping after refresh')
        if not all(label in page.locator('#cs-canvas').inner_text() for label in labels):
            raise AssertionError('Chart Studio refresh preview lost week labels')
        page.locator('[data-action="save"]').click()
        page.wait_for_function('()=>!window.CompanyUIChartStudio.state.pending', timeout=20000)
        page.reload(wait_until='domcontentloaded')
        page.locator('#chart-studio[data-studio-ready="true"]').wait_for(timeout=20000)
        saved_after_refresh = chart_studio_model(page)
        if saved_after_refresh['dataset']['rows'] != refreshed_dataset['rows'] or saved_after_refresh['mapping'].get('category') != new_category_id:
            raise AssertionError('Refreshed category binding did not survive the second save/reload')
        page.locator('[data-action="back"]').click()
        page.wait_for_url('**/visualizer**', timeout=15000)
        ready(page)

        pptx_path = output / 'categories-refreshed.pptx'
        payload, exported_categories = export_chart_pptx_categories(page, pptx_path, labels)
        page.screenshot(path=str(output / 'categories-export-ready.png'), full_page=True)
        return {
            'status': 'PASS', 'report_id': report_id, 'element_id': element_id, 'duplicate_id': duplicate_id,
            'dataset_revision': refreshed_dataset['revision'], 'initial_category_field': category_id,
            'rebound_category_field': new_category_id, 'labels': labels,
            'editor_axis_labels': post_refresh_labels, 'preview_axis_labels': preview_labels,
            'powerpoint_categories': exported_categories,
            'editor_preview_export_parity': True, 'duplicate_mapping_preserved': True,
            'save_reload_before_refresh': True, 'refresh_rebind_and_save_reload': True,
            'pptx_sha256': hashlib.sha256(payload).hexdigest(),
        }
    finally:
        context.close()


def run_legacy_percent(host, browser, output):
    report = canonical_model({
        'datasets': [{
            'id': 'percent-data', 'name': 'Yield', 'revision': 1,
            'fields': [
                {'id': 'week', 'name': 'Week', 'type': 'categorical'},
                {'id': 'yield', 'name': 'Yield', 'type': 'number', 'unit': '%', 'format': {'kind': 'percent', 'percent_scale': 'ratio'}},
            ],
            'rows': [['W0', .276]],
        }],
        'items': [{
            'id': 'legacy-kpi', 'type': 'metric', 'engine': 'MetricEngine', 'element': 'Hero KPI',
            'title': 'Legacy yield', 'order': 0, 'weight': 1.1, 'locked': False, 'z': 1,
            'value': .276, 'unit': 'M', 'value_format': 'currency', 'dataset_id': 'percent-data', 'mapping': {'value': 'yield'},
        }],
    })
    context = browser.new_context(viewport={'width': 1440, 'height': 1000}, accept_downloads=True)
    page = context.new_page()
    page.set_default_timeout(10000)
    report_id = host.create(model=report, name='chg179-legacy-percent')
    try:
        mount(page, host, report_id)
        select(page, 'legacy-kpi')
        rendered = metric_value(page, 'legacy-kpi')
        if rendered != '27.6%':
            raise AssertionError(f'Bound percent data inherited legacy currency style: {rendered!r}')
        preflight_state = page.evaluate('()=>window.__VIZ_PROD__.preflight()')
        metric_issues = [issue['message'] for issue in preflight_state['dataIssues'] if issue['kind'] == 'metric-format']
        if not any('Percent data is paired with currency formatting' in message for message in metric_issues):
            raise AssertionError(f'Legacy percent/currency contradiction was absent from preflight: {preflight_state!r}')
        page.locator('#exportBtn').click()
        validation_summary = page.locator('#modalBody').inner_text()
        if 'Validation' not in validation_summary:
            raise AssertionError(f'Export validation summary was not opened: {validation_summary!r}')
        page.locator('#exportPptAction').click()
        preflight = page.locator('#modalBody').inner_text()
        if 'Percent data is paired with currency formatting' not in preflight:
            raise AssertionError(f'Preflight did not explain percent/currency contradiction: {preflight!r}')
        page.keyboard.press('Escape')
        page.locator('#genericModal.show').wait_for(state='hidden', timeout=5000)
        if page.locator('#iValueFormat').count() != 1:
            debug = page.evaluate('''()=>({state:window.CompanyUIVisualizerBridge.state(),inspector:document.querySelector('#inspector')?.innerText,selected:document.querySelector('.component[aria-selected="true"]')?.outerHTML.slice(0,500)})''')
            raise AssertionError(f'Legacy MetricEngine inspector did not expose number presentation controls: {debug!r}')
        choose(page, '#iValueFormat', 'percent')
        panels(page, library=True, inspector=True)
        fixed = metric_value(page, 'legacy-kpi')
        if fixed != '27.6%':
            raise AssertionError(f'Explicit percent correction rendered {fixed!r}')
        page.locator('#previewBtn').click()
        page.locator('.cui-visualizer-root.preview-mode').wait_for(timeout=8000)
        preview = metric_value(page, 'legacy-kpi')
        page.locator('#previewExit').click()
        post_repair_preflight = page.evaluate('()=>window.__VIZ_PROD__.preflight()')
        if post_repair_preflight['dataIssues']:
            raise AssertionError(f'Corrected percent semantics still has data blockers: {post_repair_preflight["dataIssues"]!r}')
        if post_repair_preflight['layoutIssues']:
            ui_debug = page.evaluate('''()=>({library:document.querySelector('#libraryToggle')?.getAttribute('aria-pressed'),buttons:[...document.querySelectorAll('.tb,.mini-btn,.library-tab')].filter(button=>button.offsetParent!==null&&button.scrollWidth>button.clientWidth+1).map(button=>({text:button.textContent.trim(),clientWidth:button.clientWidth,scrollWidth:button.scrollWidth,parent:button.parentElement?.className}))})''')
            raise AssertionError(f'Legacy acceptance fixture still has layout blockers: {post_repair_preflight["layoutIssues"]!r}; UI={ui_debug!r}')
        pptx_path = output / 'legacy-percent-fixed.pptx'
        payload, text = export_pptx(page, pptx_path)
        match = next((line.strip() for line in text.splitlines() if '27.6%' in line), None)
        if preview != fixed or match != fixed:
            raise AssertionError(f'Bound percent export parity failed: editor={fixed!r}, preview={preview!r}, export={match!r}')
        return {
            'status': 'PASS', 'report_id': report_id, 'legacy_value_format': 'currency',
            'rendered_before_repair': rendered, 'preflight_issue': 'Percent data is paired with currency formatting',
            'preflight_data_issues': metric_issues,
            'preflight_data_issues_after_repair': post_repair_preflight['dataIssues'],
            'editor': fixed, 'preview': preview, 'powerpoint': match, 'parity': True,
            'pptx_sha256': hashlib.sha256(payload).hexdigest(),
        }
    finally:
        context.close()


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--base-sha', required=True)
    parser.add_argument('--base-tree', required=True)
    args = parser.parse_args()
    output = args.output.expanduser().resolve()
    if output == ROOT or ROOT in output.parents:
        parser.error('--output must be outside the checkout')
    if output.exists() and any(output.iterdir()):
        parser.error('--output must be empty')
    output.mkdir(parents=True, exist_ok=True)
    candidate_sha = git_text('rev-parse', 'HEAD')
    candidate_tree = git_text('rev-parse', 'HEAD^{tree}')
    changed_files = git_text('diff', '--name-only', f'{args.base_sha}..{candidate_sha}').splitlines()
    commits = git_text('rev-list', '--reverse', f'{args.base_sha}..{candidate_sha}').splitlines()
    worktree_status = git_text('status', '--porcelain').splitlines()
    receipt = {
        'scope': 'CHG-179 native browser semantic parity acceptance',
        'host': 'native NiceGUI 3.15.0 with isolated synthetic report repository',
        'base_sha': args.base_sha,
        'base_tree': args.base_tree,
        'candidate_sha': candidate_sha,
        'candidate_tree': candidate_tree,
        'commits': commits,
        'changed_files': changed_files,
        'worktree_status': worktree_status,
        'checks': [],
        'unexpected_browser_events': [],
    }
    try:
        with tempfile.TemporaryDirectory(prefix='visembler-chg179-native-') as temp_dir, sync_playwright() as playwright:
            with NativeHost(ROOT, Path(temp_dir) / 'data') as host:
                browser = playwright.chromium.launch(**browser_kwargs())
                receipt['browser'] = browser.version
                try:
                    metric_results = run_metric_cases(host, browser, output, receipt)
                    receipt['checks'].append({'name': 'metric-editor-preview-powerpoint-parity', 'status': 'PASS' if len(metric_results) == len(METRIC_CASES) and all(row.get('status') == 'PASS' for row in metric_results) else 'FAIL', 'cases': metric_results})
                    try:
                        category_result = run_category_workflow(host, browser, output)
                        receipt['checks'].append({'name': 'data-first-chart-studio-save-refresh-duplicate-category-export', **category_result})
                    except Exception as exc:
                        receipt['checks'].append({'name': 'data-first-chart-studio-save-refresh-duplicate-category-export', 'status': 'FAIL', 'error': str(exc), 'traceback': traceback.format_exc()})
                    try:
                        receipt['checks'].append({'name': 'legacy-bound-percent-preflight-and-parity', **run_legacy_percent(host, browser, output)})
                    except Exception as exc:
                        receipt['checks'].append({'name': 'legacy-bound-percent-preflight-and-parity', 'status': 'FAIL', 'error': str(exc), 'traceback': traceback.format_exc()})
                finally:
                    browser.close()
    except Exception as exc:
        receipt['harness_error'] = str(exc)
        receipt['traceback'] = traceback.format_exc()
    receipt['status'] = 'PASS' if receipt['checks'] and all(check.get('status') == 'PASS' for check in receipt['checks']) else 'FAIL'
    write_json(output / 'chg179-semantic-parity.json', receipt)
    print(json.dumps({'status': receipt['status'], 'checks': [{'name': row.get('name'), 'status': row.get('status')} for row in receipt['checks']]}, ensure_ascii=False))
    return 0 if receipt['status'] == 'PASS' else 1


if __name__ == '__main__':
    raise SystemExit(main())
