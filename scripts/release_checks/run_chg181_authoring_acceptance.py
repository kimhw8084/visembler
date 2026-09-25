#!/usr/bin/env python3
"""Native UI acceptance for CHG-181 whole-report composition and reuse.

Every report in this runner is created with Report Hub controls. Data enters
through Data First and component content is changed through the Inspector or
the table editor. The bridge is read-only and is used only for assertions.
"""
from __future__ import annotations

import argparse
import copy
import hashlib
import json
import re
import sys
import tempfile
import time
import traceback
from pathlib import Path
from urllib.parse import quote

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from playwright.sync_api import sync_playwright

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(Path(__file__).parent))

from editor_host import NativeHost
from native_common import browser_kwargs, write_json


REPORTS = (
    ("executive-business-review", "executive-business-review"),
    ("semiconductor-rca", "semiconductor-rca"),
    ("experiment-decision", "experiment-decision"),
    ("technical-status-review", "technical-status-review"),
)

def semiconductor_wafer_fixture() -> str:
    rows = ["Die X\tDie Y\tYield"]
    for y in range(-5, 6):
        for x in range(-5, 6):
            if x * x + y * y > 25:
                continue
            affected = x in {2, 3} and y in {0, 1}
            value = 88.4 + ((x * 7 + y * 11) % 6) / 10 if affected else 98.1 + ((x * 5 + y * 3) % 8) / 10
            rows.append(f"{x}\t{y}\t{value:.1f}")
    return "\n".join(rows) + "\n"


def exact_numeric_tokens(value: str) -> set[float]:
    return {float(token) for token in re.findall(r"(?<!\d)\d+(?:\.\d+)?(?!\d)", value)}

DATA = {
    "executive-business-review": "Quarter\tRetention\nQ1\t91.8\nQ2\t93.2\nQ3\t94.1\nQ4\t94.6\n",
    "semiconductor-rca": semiconductor_wafer_fixture(),
    "semiconductor-distribution": "Population\tYield\nReference\t98.8\nReference\t98.5\nReference\t98.6\nAffected\t94.1\nAffected\t94.6\nAffected\t93.8\n",
    "experiment-decision": "Week\tCohort\tConversion\nW1\tControl\t8.0\nW1\tTreatment\t8.2\nW2\tControl\t8.1\nW2\tTreatment\t8.7\nW3\tControl\t8.4\nW3\tTreatment\t9.4\nW4\tControl\t8.5\nW4\tTreatment\t9.8\n",
    "experiment-distribution": "Cohort\tConversion\nControl\t8.0\nControl\t8.1\nControl\t8.4\nControl\t8.5\nTreatment\t8.2\nTreatment\t8.7\nTreatment\t9.4\nTreatment\t9.8\n",
    "technical-status-review": "Shift\tYield\nA\t94.1\nB\t94.6\nC\t95.0\n",
    "technical-spc": "Time\tMeasurement\tLSL\tUSL\n1\t94.1\t92\t97\n2\t94.6\t92\t97\n3\t95.0\t92\t97\n4\t94.8\t92\t97\n5\t95.2\t92\t97\n",
    "holdout": "Month\tOrders\tReturns\nJan\t1280\t42\nFeb\t1395\t38\nMar\t1522\t35\nApr\t1611\t31\n",
}

TEXT = {
    "executive-business-review": (
        "Revenue grew across all three regions. Retention improved in the West, "
        "while the South remains below the 95 percent target."
    ),
    "semiconductor-rca": (
        "Affected lot L2407 shows a localized die-level yield loss after "
        "ETCH-08 chamber A processing. Contain the lot until the chamber check "
        "and paired reference run confirm recovery."
    ),
    "experiment-decision": (
        "The treatment cohort improved conversion across four weeks. The "
        "quality guardrail remained inside its agreed range; the sample does "
        "not establish long-term retention."
    ),
    "technical-status-review": (
        "Shift B recovered yield after chamber maintenance. Keep the recipe "
        "locked through the next paired-lot check and hand off the evidence "
        "to Quality and Equipment Engineering."
    ),
    "holdout": (
        "Order volume increased through the spring launch. Returns declined "
        "each month; continue weekly review while the fulfillment mix changes."
    ),
}

INTERPRETATION = {
    "experiment-decision": (
        "Treatment conversion rose from 8.2% in week 1 to 9.8% in week 4; "
        "the control moved from 8.0% to 8.5%. The treatment distribution is "
        "higher across this sample, with all three cohort checks inside the "
        "quality guardrail. Continue only as a limited rollout because these "
        "four weeks do not establish long-term retention."
    ),
    "technical-status-review": (
        "Yield recovered from 94.1% on shift A to 95.0% on shift C, while "
        "the control sample remains within specification. Keep the recipe "
        "locked until the next paired-lot check closes the two open items."
    ),
}


class Actions:
    def __init__(self, name: str):
        self.name = name
        self.events: list[dict] = []
        self.context = "hub"

    def note(self, action: str, context: str | None = None) -> None:
        if context and context != self.context:
            self.events.append({"action": "context switch", "from": self.context, "to": context})
            self.context = context
        self.events.append({"action": action, "context": self.context})

    def summary(self) -> dict:
        return {
            "report": self.name,
            "user_actions": len(self.events),
            "context_switches": sum(row["action"] == "context switch" for row in self.events),
            "modal_openings": sum(row["action"].startswith("open ") for row in self.events),
            "events": self.events,
        }


def state(page) -> dict:
    return page.evaluate("window.CompanyUIVisualizerBridge.state()")


def model(page) -> dict:
    return state(page)["model"]


def settled(page, timeout: int = 20_000) -> None:
    page.wait_for_function(
        "()=>{const s=window.CompanyUIVisualizerBridge?.state?.();return s&&s.pending===0&&!s.inflight}",
        timeout=timeout,
    )


def ready(page) -> None:
    page.locator('.cui-visualizer-root[data-editor-ready="true"]').wait_for(timeout=20_000)
    settled(page)


def create_from_hub(page, host: NativeHost, template_id: str, actions: Actions) -> str:
    page.goto(f"{host.url}/visualizer/reports", wait_until="domcontentloaded")
    hub = page.locator('[data-testid="report-hub"]')
    hub.wait_for(timeout=20_000)
    page.get_by_role("button", name="Create report", exact=True).click()
    actions.note("open Report Hub blueprint chooser")
    chooser = page.locator('[data-cui-overlay="dialog"]:visible')
    chooser.wait_for(timeout=5_000)
    chooser.locator(f'[data-template-id="{template_id}"]').get_by_role(
        "button", name="Create from this blueprint", exact=True
    ).click()
    actions.note(f"create report from {template_id}", "editor")
    page.wait_for_url("**/visualizer?report=**", timeout=20_000)
    ready(page)
    report_id = page.evaluate("new URL(location.href).searchParams.get('report')")
    assert report_id, "Report Hub did not provide a governed report ID."
    return report_id


def select_item(page, item_id: str, *, additive: bool = False) -> None:
    node = page.locator(f'.component[data-id="{item_id}"]')
    node.scroll_into_view_if_needed()
    if additive:
        node.focus()
        node.press("Shift+Enter")
    else:
        node.focus()
        node.press("Enter")
    page.locator("#inspectorToggle").wait_for(state="visible", timeout=5_000)
    if page.locator("#inspectorToggle").get_attribute("aria-pressed") != "true":
        page.locator("#inspectorToggle").click()


def delete_component(page, item_id: str, actions: Actions) -> None:
    select_item(page, item_id)
    page.locator('[data-inspector="delete"]').click()
    settled(page)
    actions.note("remove an unused empty template chart", "editor")


def one_item(page, *, element: str | None = None, engine: str | None = None, role: str | None = None) -> dict:
    candidates = model(page)["items"]
    found = [
        item for item in candidates
        if (element is None or item.get("element") == element)
        and (engine is None or item.get("engine") == engine)
        and (role is None or item.get("composition_role") == role)
    ]
    if not found:
        raise AssertionError(f"No report component matched element={element!r}, engine={engine!r}, role={role!r}.")
    return found[0]


def fill_control(page, selector: str, value: str, actions: Actions, label: str) -> None:
    control = page.locator(selector)
    control.wait_for(state="visible", timeout=5_000)
    control.fill(value)
    control.press("Tab")
    settled(page)
    actions.note(f"edit {label}", "editor")


def set_component_title(page, entry: dict, title: str, actions: Actions) -> None:
    select_item(page, str(entry["id"]))
    if page.locator("#iTitle").count():
        fill_control(page, "#iTitle", title, actions, "component title")
    visible = page.locator("#iShowTitle")
    if visible.count() and not visible.is_checked():
        visible.check()
        settled(page)
        actions.note("show component title", "inspector")


def set_composition_role(page, entry: dict, role: str, actions: Actions) -> None:
    select_item(page, str(entry["id"]))
    control = page.locator("#iCompositionRole")
    if control.count():
        control.select_option(role)
        settled(page)
        actions.note(f"set report intent · {role}", "inspector")


def set_table_headers(page, table_id: str, headers: list[str], actions: Actions) -> None:
    select_item(page, table_id)
    grid = page.locator("#tableEditorGrid")
    while grid.locator("[data-table-header]").count() < len(headers):
        last = grid.locator("[data-table-header]").last
        last.focus()
        page.locator('[data-table-action="insert-column-right"]').click()
        settled(page)
        actions.note("add evidence column", "inspector")
    for index, value in enumerate(headers):
        header = page.locator(f'#tableEditorGrid [data-table-header="{index}"]')
        if not header.count():
            break
        if header.input_value() != value:
            header.fill(value)
            header.press("Tab")
            settled(page)
            actions.note(f"name evidence column · {value}", "inspector")


def fill_text(page, entry: dict, value: str, actions: Actions) -> None:
    select_item(page, str(entry["id"]))
    fill_control(page, "#iText", value, actions, entry.get("title") or entry.get("element") or "narrative")


def fill_metrics_and_decision(page, genre: str, actions: Actions) -> None:
    profiles = {
        "executive-business-review": [
            {"title": "Retention", "value": "94.6", "unit": "%", "target": "95", "delta": "2.8"},
            {"title": "Regional revenue", "value": "551", "unit": "", "target": "", "delta": ""},
        ],
        "semiconductor-rca": [
            {"title": "Affected yield", "value": "94.2", "unit": "%", "target": "95", "delta": "-4.4"},
        ],
        "experiment-decision": [
            {"title": "Treatment conversion", "value": "9.8", "unit": "%", "target": "10", "delta": "1.6"},
        ],
        "technical-status-review": [
            {"title": "First pass yield", "value": "95.0", "unit": "%", "target": "96", "delta": "0.9"},
            {"title": "Unplanned downtime", "value": "8.3", "unit": "h", "target": "6", "delta": "-1.7"},
        ],
        "holdout": [
            {"title": "April orders", "value": "1611", "unit": "", "target": "", "delta": "331"},
            {"title": "Order growth since January", "value": "331", "unit": "orders", "target": "", "delta": ""},
        ],
        "regional-reuse-source": [
            {"title": "Regional revenue total", "value": "175", "unit": "units", "target": "", "delta": ""},
            {"title": "West lead over East", "value": "15", "unit": "units", "target": "", "delta": ""},
        ],
        "regional-reuse-followup": [
            {"title": "Regional revenue total", "value": "385", "unit": "units", "target": "", "delta": ""},
            {"title": "West lead over East", "value": "25", "unit": "units", "target": "", "delta": ""},
        ],
        "blank-rebuild": [
            {"title": "Regional revenue total", "value": "385", "unit": "units", "target": "", "delta": ""},
            {"title": "West lead over East", "value": "25", "unit": "units", "target": "", "delta": ""},
        ],
    }
    comparison_profiles = {
        "executive-business-review": {"title": "Quarterly retention", "before": "91.8", "after": "94.6", "unit": "%"},
        "semiconductor-rca": {"title": "Affected versus reference yield", "before": "98.6", "after": "94.2", "unit": "%"},
        "experiment-decision": {"title": "Treatment conversion · Week 1 to 4", "before": "8.2", "after": "9.8", "unit": "%"},
        "holdout": {"title": "Monthly returns", "before": "42", "after": "31", "unit": ""},
        "regional-reuse-source": {"title": "East to West revenue", "before": "80", "after": "95", "unit": "units"},
        "regional-reuse-followup": {"title": "East to West revenue", "before": "180", "after": "205", "unit": "units"},
        "blank-rebuild": {"title": "East to West revenue", "before": "180", "after": "205", "unit": "units"},
    }
    decision_profiles = {
        "executive-business-review": ("Proceed with monitored revenue recovery", "Close the South retention gap before the next review.", "Review"),
        "semiconductor-rca": ("Contain lot L2407 through chamber requalification", "Release only after chamber A and the paired reference run confirm recovery.", "Contain"),
        "experiment-decision": ("Stage a limited treatment rollout", "Pause expansion if the quality guardrail moves outside its agreed range.", "Review"),
        "technical-status-review": ("Hold added capacity until the paired-lot check closes", "Keep the recipe locked and attach the next-shift control evidence.", "Review"),
        "holdout": ("Review rising orders alongside falling returns", "Check the fulfillment mix before changing the spring capacity plan.", "Review"),
        "regional-reuse-source": ("Validate the West revenue lead before shifting regional targets", "The source review compares East and West; confirm the pattern in the next operating window.", "Review"),
        "regional-reuse-followup": ("Validate the West revenue lead before reallocating regional spend", "The destination extract contains two regions; confirm the pattern in the next operating window.", "Review"),
        "blank-rebuild": ("Validate the West revenue lead before reallocating regional spend", "The destination extract contains two regions; confirm the pattern in the next operating window.", "Review"),
    }
    metrics = [entry for entry in model(page)["items"] if entry.get("engine") == "MetricEngine"]
    for index, entry in enumerate(metrics):
        profile = profiles[genre][index] if index < len(profiles[genre]) else profiles[genre][-1]
        set_component_title(page, entry, profile["title"], actions)
        select_item(page, str(entry["id"]))
        for selector, key in (("#iValue", "value"), ("#iTarget", "target"), ("#iDelta", "delta"), ("#iUnit", "unit")):
            if page.locator(selector).count():
                fill_control(page, selector, profile[key], actions, selector[1:])
    comparisons = [entry for entry in model(page)["items"] if entry.get("engine") == "ComparisonEngine"]
    if comparisons and genre in comparison_profiles:
        entry = comparisons[0]
        profile = comparison_profiles[genre]
        set_component_title(page, entry, profile["title"], actions)
        select_item(page, str(entry["id"]))
        for selector, key in (("#iBefore", "before"), ("#iAfter", "after"), ("#iUnit", "unit")):
            if page.locator(selector).count():
                fill_control(page, selector, profile[key], actions, selector[1:])
    decisions = [
        entry for entry in model(page)["items"]
        if entry.get("engine") in {"DecisionCompositeEngine", "ProjectCompositeEngine"}
    ]
    if decisions:
        entry = decisions[0]
        select_item(page, str(entry["id"]))
        statement, detail, status = decision_profiles[genre]
        for selector, value in (
            ("#iStatement", statement),
            ("#iDetail", detail),
            ("#iStatus", status),
        ):
            if page.locator(selector).count():
                fill_control(page, selector, value, actions, selector[1:])


def fill_table(page, actions: Actions, rows: list[str] | list[list[str]]) -> bool:
    tables = [entry for entry in model(page)["items"] if entry.get("engine") == "TableEngine"]
    if not tables:
        return False
    entry = tables[0]
    row_values = [rows] if rows and isinstance(rows[0], str) else rows
    select_item(page, str(entry["id"]))
    grid = page.locator("#tableEditorGrid")
    if grid.count() == 0:
        return False
    if grid.locator("[data-table-cell]").count() == 0:
        empty_add = page.locator(f'.component[data-id="{entry["id"]}"] [data-empty-action="add-row"]')
        if empty_add.count():
            empty_add.click()
        settled(page)
    for row_index, row in enumerate(row_values):
        existing = len(model(page)["items"][next(i for i, item in enumerate(model(page)["items"]) if item["id"] == entry["id"])].get("customTable", {}).get("rows", []))
        if row_index >= existing:
            if row_index:
                grid.locator(f'[data-table-cell="{row_index - 1}:0"]').focus()
                page.locator('[data-table-action="insert-row-below"]').click()
                settled(page)
            else:
                empty_add = page.locator(f'.component[data-id="{entry["id"]}"] [data-empty-action="add-row"]')
                if empty_add.count():
                    empty_add.click()
                    settled(page)
        for column, value in enumerate(row):
            cell = grid.locator(f'[data-table-cell="{row_index}:{column}"]')
            if not cell.count():
                return False
            cell.fill(value)
            cell.press("Tab")
            settled(page)
    actions.note(f"enter {len(row_values)} evidence rows in table editor", "editor")
    return True


def clear_table_rows(page, table_id: str, actions: Actions) -> None:
    select_item(page, table_id)
    grid = page.locator("#tableEditorGrid")
    cells = grid.locator("[data-table-cell]")
    if not cells.count():
        return
    cells.first.focus()
    page.keyboard.press("Control+a")
    page.locator('[data-table-action="delete-rows"]').click()
    settled(page)
    actions.note("clear source-specific evidence rows", "inspector")


def finish_regional_revenue_report(page, stage: str, actions: Actions) -> dict:
    content = {
        "regional-reuse-source": {
            "headline": "West revenue leads East by 15 units in the regional review.",
            "context": "Regional revenue review · East and West · source extract captured for the first report.",
            "interpretation": "Revenue totals 175 units across the two regions. West is 15 units above East; the report keeps the comparison visible before any regional target changes.",
            "conclusion": "Next step: validate the West lead in the next operating window before shifting regional targets.",
            "rows": [["East", "80", "Comparison baseline", "Source extract"], ["West", "95", "Leads by 15", "Source extract"]],
            "risk": ("Validate the West lead before shifting regional targets", "The current review compares two regions; confirm the pattern in the next operating window.", "Review"),
        },
        "regional-reuse-followup": {
            "headline": "West revenue leads East by 25 units in the newly selected data.",
            "context": "Regional revenue review · East and West · destination data selected through guided remap.",
            "interpretation": "The destination extract totals 385 units. West is 25 units above East; the shared chart pair, metric hierarchy, comparison and evidence table now use the new rows.",
            "conclusion": "Next step: validate the West lead in the next operating window before shifting regional targets.",
            "rows": [["East", "180", "Comparison baseline", "Destination extract"], ["West", "205", "Leads by 25", "Destination extract"]],
            "risk": ("Validate the West lead before shifting regional targets", "The destination contains two regions; confirm the pattern in the next operating window.", "Review"),
        },
        "blank-rebuild": {
            "headline": "West revenue leads East by 25 units in the regional review.",
            "context": "Regional revenue review · East and West · destination source entered into a blank report.",
            "interpretation": "The extract totals 385 units. West is 25 units above East; confirm the pattern before changing regional targets.",
            "conclusion": "Next step: validate the West lead in the next operating window before shifting regional targets.",
            "rows": [["East", "180", "Comparison baseline", "Destination extract"], ["West", "205", "Leads by 25", "Destination extract"]],
            "risk": ("Validate the West lead before shifting regional targets", "The destination contains two regions; confirm the pattern in the next operating window.", "Review"),
        },
    }[stage]
    fill_metrics_and_decision(page, stage, actions)
    roles = {entry.get("composition_role"): entry for entry in model(page)["items"]}
    for role, key in (("report_headline", "headline"), ("context", "context"), ("narrative_interpretation", "interpretation"), ("conclusion", "conclusion")):
        entry = roles.get(role)
        if entry and entry.get("engine") == "TextEngine":
            fill_text(page, entry, content[key], actions)
    decision = next((entry for entry in model(page)["items"] if entry.get("engine") == "DecisionCompositeEngine"), None)
    if decision:
        select_item(page, str(decision["id"]))
        for selector, value in zip(("#iStatement", "#iDetail", "#iStatus"), content["risk"]):
            fill_control(page, selector, value, actions, selector[1:])
    charts = [entry for entry in model(page)["items"] if entry.get("engine") == "CoreChartEngine" and entry.get("dataset_id")]
    for index, chart in enumerate(charts):
        set_component_title(page, chart, "Regional revenue by area" if index == 0 else "Regional revenue comparison", actions)
    table = next(entry for entry in model(page)["items"] if entry.get("engine") == "TableEngine")
    set_table_headers(page, str(table["id"]), ["Region", "Revenue", "Reading", "Source"], actions)
    clear_table_rows(page, str(table["id"]), actions)
    fill_table(page, actions, content["rows"])
    set_component_title(page, table, "Regional revenue evidence", actions)
    saved = model(page)
    return {"content": content, "saved_rows": next(entry for entry in saved["items"] if entry.get("engine") == "TableEngine")["customTable"]["rows"]}


def import_or_replace_data(page, text: str, view: str, actions: Actions, target_id: str | None = None) -> str:
    if target_id:
        select_item(page, target_id)
    if page.locator("#inspectorToggle").get_attribute("aria-pressed") == "true":
        page.locator("#inspectorToggle").click()
    if page.locator("#libraryToggle").get_attribute("aria-pressed") != "true":
        page.locator("#libraryToggle").click()
    if page.locator("#elementsTab").get_attribute("aria-selected") != "true":
        page.locator("#elementsTab").click()
    page.evaluate("()=>document.querySelectorAll('#libraryPane,#elementsView,.quick-section').forEach(node=>{node.scrollTop=0;node.scrollLeft=0})")
    page.wait_for_timeout(180)
    page.locator("#pasteDataBtn").wait_for(state="visible", timeout=5_000)
    page.locator("#pasteDataBtn").click()
    actions.note("open Data First")
    dialog = page.locator("#genericModal.show")
    dialog.wait_for(timeout=5_000)
    page.locator("#dataFirstText").fill(text)
    page.locator(".data-first-summary").wait_for(timeout=10_000)
    choose = page.locator('[data-data-first-stage="choose"]').last
    if choose.count() and choose.is_visible():
        choose.click()
    view_button = page.locator(f'[data-data-first-view="{view}"]')
    try:
        view_button.wait_for(state="visible", timeout=2_000)
    except Exception:
        diagnostics = page.evaluate("""()=>({
          modal:document.querySelector('#modalBody')?.innerText||'',
          views:[...document.querySelectorAll('[data-data-first-view]')].map(n=>({view:n.dataset.dataFirstView,label:n.innerText,visible:!!n.offsetParent,disabled:n.disabled})),
          stages:[...document.querySelectorAll('[data-data-first-stage]')].map(n=>({stage:n.dataset.dataFirstStage,label:n.innerText,visible:!!n.offsetParent,disabled:n.disabled}))
        })""")
        choices = [
            row["view"] for row in diagnostics["views"]
            if row["visible"] and not row["disabled"]
        ]
        if not choices:
            raise AssertionError(f"Data First offered no supported visual: {json.dumps(diagnostics, ensure_ascii=False)[:5000]}")
        selected_view = "bar" if target_id and "bar" in choices else choices[0]
        view_button = page.locator(f'[data-data-first-view="{selected_view}"]')
    else:
        selected_view = view
    if view_button.get_attribute("aria-pressed") != "true":
        view_button.click()
    actions.note(f"choose {selected_view} result")
    map_stage = page.locator('[data-data-first-stage="map"]').last
    if map_stage.count() and map_stage.is_visible():
        map_stage.click()
    commit_stage = page.locator('[data-data-first-stage="commit"]').last
    if commit_stage.count() and commit_stage.is_visible():
        commit_stage.click()
    action_button = page.locator("#dataFirstReplace" if target_id else "#dataFirstCreate")
    action_button.wait_for(state="visible", timeout=10_000)
    if action_button.is_disabled():
        details = page.locator("#modalBody").inner_text()
        raise AssertionError(f"Data First mapping was not compatible: {details[:1800]}")
    action_button.click()
    settled(page)
    page.locator("#genericModal.show").wait_for(state="hidden", timeout=10_000)
    actions.note("commit governed dataset and mapping")
    after = model(page)
    if target_id:
        return target_id
    charts = [entry for entry in after["items"] if entry.get("dataset_id")]
    assert charts, "Data First created no bound visual."
    return str(charts[-1]["id"])


def export_pptx(page, path: Path, actions: Actions, host: NativeHost | None = None) -> list[str]:
    page.locator("#exportBtn").click()
    actions.note("open export menu")
    downloads = []
    listener = lambda download: downloads.append(download)
    page.on("download", listener)
    page.locator("#exportPptAction").click()
    deadline = time.monotonic() + 35
    while time.monotonic() < deadline and not downloads:
        page.wait_for_timeout(100)
    page.remove_listener("download", listener)
    if not downloads:
        details = page.evaluate("""()=>({
          bridge:(()=>{const s=window.CompanyUIVisualizerBridge?.state?.();return s?{report_id:s.report_id,revision:s.revision,pending:s.pending,inflight:s.inflight,recovery:s.recovery}:null})(),
          debug:(window.__VIZ_PROD__?.ui?.debugLog?.slice(0,12)||[]).map(({level,event,detail})=>({level,event,detail})),
          preflight:(()=>{const p=window.__VIZ_PROD__?.preflight?.();return p?{issues:p.issues,layoutIssues:p.layoutIssues,dataIssues:p.dataIssues}:null})(),
          component:(()=>{const p=window.__VIZ_PROD__?.preflight?.(),id=p?.layoutIssues?.[0]?.id,node=id&&document.querySelector('.component[data-id="'+id+'"]'),content=node?.querySelector('.c-content'),rendered=content?.querySelector('.integrated-element-content'),measure=n=>n?{tag:n.tagName,class:n.getAttribute('class'),clientWidth:n.clientWidth,clientHeight:n.clientHeight,scrollWidth:n.scrollWidth,scrollHeight:n.scrollHeight,offsetWidth:n.offsetWidth,offsetHeight:n.offsetHeight,overflow:getComputedStyle(n).overflow,style:(()=>{const s=getComputedStyle(n);return{width:s.width,height:s.height,minHeight:s.minHeight,maxHeight:s.maxHeight,display:s.display,flex:s.flex,gridRow:s.gridRow,gridTemplateColumns:s.gridTemplateColumns,gridTemplateRows:s.gridTemplateRows,alignSelf:s.alignSelf,aspectRatio:s.aspectRatio,cssText:n.getAttribute('style')}})(),box:(()=>{const r=n.getBoundingClientRect();return{x:r.x,y:r.y,w:r.width,h:r.height}})(),viewBox:n.getAttribute('viewBox'),bbox:n.tagName.toLowerCase()==='svg'?(()=>{try{const b=n.getBBox();return{x:b.x,y:b.y,w:b.width,h:b.height}}catch{return null}})():null}:null;return node?{id,label:node.getAttribute('aria-label'),rects:window.__VIZ_PROD__?.layoutRects?.().filter(value=>value.id===id),component:measure(node),content:measure(content),rendered:measure(rendered),contentChildren:[...content.children].map(measure),renderedParent:measure(rendered?.parentElement),children:[...content.querySelectorAll('.chart-wrap,.diagram-svg,.card-body,.cs-static-chart,.cs-chart-svg,svg')].map(measure)}:null})(),
          toast:document.querySelector('[role=status]')?.innerText||'',
          modal:document.querySelector('#genericModal.show #modalBody')?.innerText||''
        })""")
        page.screenshot(path=str(path.parent / (path.stem + "-export-error.png")), full_page=True)
        log_tail = host.log_path.read_text(encoding="utf-8", errors="replace")[-5000:] if host else ""
        raise AssertionError(f"Editable PowerPoint did not download: {json.dumps(details, ensure_ascii=False)[:16000]}\\nNative host log tail:\\n{log_tail}")
    downloads[0].save_as(str(path))
    actions.note("export editable PowerPoint")
    deck = Presentation(str(path))
    texts = [
        shape.text
        for slide in deck.slides
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
    ]
    assert len(deck.slides) >= 1, "Editable PowerPoint has no slides."
    return texts


def export_svg(page, path: Path, actions: Actions) -> None:
    page.locator("#exportBtn").click()
    with page.expect_download(timeout=20_000) as download:
        page.locator("#exportSvgAction").click()
    download.value.save_as(str(path))
    actions.note("export report SVG")
    assert "<svg" in path.read_text(encoding="utf-8")


def pptx_chart_values(path: Path) -> list[dict]:
    deck = Presentation(str(path))
    return [
        {"series": series.name, "values": [float(value) for value in series.values if isinstance(value, (int, float))], "categories": [str(value.label) for value in shape.chart.plots[0].categories]}
        for slide in deck.slides
        for shape in slide.shapes if getattr(shape, "has_chart", False)
        for series in shape.chart.series
    ]


def compose_and_validate(page, actions: Actions) -> dict:
    semantic_before = page.evaluate("""()=>JSON.stringify(window.CompanyUIVisualizerBridge.state().model.items.map(item=>({id:item.id,engine:item.engine,element:item.element,view_type:item.view_type,variant:item.variant,mapping:item.mapping,data:item.data,rows:item.rows,transform_recipe:item.transform_recipe,transform_pipeline:item.transform_pipeline,analysis_recipe:item.analysis_recipe,statistical_recipe:item.statistical_recipe,engineering_recipe:item.engineering_recipe,chart_studio:item.chart_studio})))""")
    page.locator("#auto").click()
    settled(page)
    actions.note("Compose/Reflow report", "editor")
    result = page.evaluate("window.__VIZ_PROD__.preflight()")
    summary = {
        "issues": result["issues"],
        "layout_issue_count": len(result["layoutIssues"]),
        "data_issue_count": len(result["dataIssues"]),
    }
    if summary["layout_issue_count"] or summary["data_issue_count"]:
        summary["rendered_items"] = page.evaluate("""()=>[...document.querySelectorAll('.component')].map(node=>{const c=node.querySelector('.c-content'),content=c?.querySelector('.integrated-element-content')||c||node,box=content.getBoundingClientRect(),all=[c,content,...content.querySelectorAll('.gallery-card,.card-body,.table-wrap,.chart-wrap,.diagram-svg,.viz-svg,.cs-static-chart,.cs-chart-svg,.plot-area,svg')].filter(Boolean),metrics=all.map(child=>{const rect=child.getBoundingClientRect(),css=getComputedStyle(child);return {tag:child.tagName,className:child.getAttribute('class'),box:{width:rect.width,height:rect.height},scroll:{width:child.scrollWidth,height:child.scrollHeight},client:{width:child.clientWidth,height:child.clientHeight},css:{height:css.height,minHeight:css.minHeight,maxHeight:css.maxHeight,overflow:css.overflow}}}),overflow=metrics.slice(1).reduce((max,child)=>({x:Math.max(max.x,child.scroll.width-child.client.width),y:Math.max(max.y,child.scroll.height-child.client.height)}),{x:0,y:0});return {id:node.dataset.id,box:{width:box.width,height:box.height},overflow,metrics}})""")
        raise AssertionError(f"Smart composition is not export-ready after content fit: {json.dumps(summary, ensure_ascii=False)}")
    geometry = page.evaluate("window.__VIZ_PROD__.layoutGeometry()")
    rects = geometry["items"]
    overlaps = [
        [left["id"], right["id"]]
        for index, left in enumerate(rects)
        for right in rects[index + 1:]
        if left["x"] < right["x"] + right["w"] and left["x"] + left["w"] > right["x"] and left["y"] < right["y"] + right["h"] and left["y"] + left["h"] > right["y"]
    ]
    bounded = all(rect["x"] >= 0 and rect["y"] >= 0 and rect["x"] + rect["w"] <= geometry["canvas"]["width"] and rect["y"] + rect["h"] <= geometry["canvas"]["height"] for rect in rects)
    if overlaps or not bounded:
        raise AssertionError(f"Smart report geometry is not a bounded, non-overlapping composition: {json.dumps({'overlaps': overlaps, 'canvas': geometry['canvas'], 'items': rects}, ensure_ascii=False)}")
    serialized = page.evaluate("window.__VIZ_PROD__.serialize()")
    model = json.loads(serialized)
    semantic_after = page.evaluate("""()=>JSON.stringify(window.CompanyUIVisualizerBridge.state().model.items.map(item=>({id:item.id,engine:item.engine,element:item.element,view_type:item.view_type,variant:item.variant,mapping:item.mapping,data:item.data,rows:item.rows,transform_recipe:item.transform_recipe,transform_pipeline:item.transform_pipeline,analysis_recipe:item.analysis_recipe,statistical_recipe:item.statistical_recipe,engineering_recipe:item.engineering_recipe,chart_studio:item.chart_studio})))""")
    if semantic_before != semantic_after:
        raise AssertionError("Whole-report composition must not change chart, data, transform, or analysis meaning.")
    sections_with_multiple_items = {}
    for rect in rects:
        if rect.get("section"):
            sections_with_multiple_items.setdefault(rect["section"], []).append(rect)
    for section_name, section_items in sections_with_multiple_items.items():
        if len(section_items) > 1 and len({item.get("sectionPattern") for item in section_items}) != 1:
            raise AssertionError(f"Every item in {section_name!r} must share one derived section composition pattern.")
        if len(section_items) > 1 and sum(item.get("compositionLevel") == "feature" for item in section_items) != 1:
            raise AssertionError(f"{section_name!r} must expose one feature and its supporting items: {section_items}")
    items_by_id = {str(item.get("id")): item for item in model.get("items", [])}
    summary["geometry"] = {
        "canvas": geometry["canvas"],
        "item_count": len(rects),
        "overlaps": overlaps,
        "bounded": bounded,
        "items": [
            {
                **rect,
                "title": items_by_id.get(str(rect["id"]), {}).get("title"),
                "role": items_by_id.get(str(rect["id"]), {}).get("composition_role"),
                "section": items_by_id.get(str(rect["id"]), {}).get("section_title"),
                "pattern": rect.get("sectionPattern"),
                "level": rect.get("compositionLevel"),
            }
            for rect in rects
        ],
    }
    visible_sections = page.locator("#componentLayer .composition-section-heading").evaluate_all("nodes=>nodes.map(node=>({title:node.textContent.trim(),section:node.dataset.section,top:parseFloat(node.style.top)||0,height:node.getBoundingClientRect().height,role:node.getAttribute('role'),level:node.getAttribute('aria-level')}))")
    if not visible_sections or any(not section["title"] or section["role"] != "heading" or section["level"] != "2" for section in visible_sections):
        raise AssertionError(f"Smart sections must have visible, accessible headings: {visible_sections}")
    for heading in visible_sections:
        section_items = [rect for rect in rects if str(rect.get("section")) == str(heading["section"])]
        if not section_items:
            raise AssertionError(f"Section heading {heading!r} has no matching composition group.")
        first_y = min(float(rect["y"]) for rect in section_items)
        expected_top = max(0.0, float(next((rect["sectionHeadingY"] for rect in section_items if rect.get("sectionStart")), first_y - 26)))
        if abs(float(heading["top"]) - expected_top) > 0.5 or float(heading["top"]) + float(heading["height"]) > first_y + 0.5:
            raise AssertionError(f"Section heading must sit above its first composed row without overlap: {heading!r}; row top={first_y}.")
    summary["geometry"]["visible_sections"] = visible_sections
    summary["geometry"]["chart_data_semantics_unchanged"] = True
    return summary


def capture_editor_canvas(page, destination: Path) -> None:
    page.set_viewport_size({"width": 2300, "height": 3600})
    page.evaluate("window.__VIZ_PROD__.setZoom(1,false)")
    page.wait_for_timeout(150)
    page.locator("#hull").scroll_into_view_if_needed()
    page.locator("#hull").screenshot(path=str(destination))
    page.set_viewport_size({"width": 1440, "height": 1000})
    page.locator("#zoomFit").click()
    page.wait_for_timeout(100)


def capture_report(page, output: Path, name: str, actions: Actions) -> dict:
    settled(page)
    page.set_viewport_size({"width": 1440, "height": 1000})
    page.keyboard.press("Escape")
    for selector in ("#libraryToggle", "#inspectorToggle"):
        if page.locator(selector).get_attribute("aria-pressed") == "true":
            page.locator(selector).click()
    page.locator("#zoomFit").click()
    page.wait_for_timeout(150)
    page.screenshot(path=str(output / f"{name}-editor-workspace-viewport.png"), full_page=True)
    page.locator("#toast").evaluate("node=>{node.classList.remove('show');node.textContent=''}")
    capture_editor_canvas(page, output / f"{name}-editor-desktop.png")
    editor_scroll = page.evaluate("document.documentElement.scrollWidth-innerWidth")
    page.locator("#previewBtn").click()
    page.locator(".cui-visualizer-root.preview-mode").wait_for(timeout=10_000)
    actions.note("open report Preview", "preview")
    preview_height = page.evaluate("Math.ceil(window.__VIZ_PROD__.layoutGeometry().canvas.height*1.2+160)")
    page.set_viewport_size({"width": 1800, "height": max(1600, preview_height)})
    page.locator("#previewFitWidth").click()
    page.locator("#toast").evaluate("node=>{node.classList.remove('show');node.textContent=''}")
    preview_chart_summaries=page.locator("#componentLayer .cs-static-summary").all_inner_texts()
    preview_chart_mark_titles=page.locator("#componentLayer .cs-chart-svg").evaluate_all("svgs=>svgs.map(svg=>({chart:svg.closest('.component')?.dataset.id,titles:[...svg.querySelectorAll('title,[aria-label]')].map(node=>node.getAttribute('aria-label')||node.textContent)}))")
    preview_table_summaries=page.locator('#componentLayer .component[data-engine="TableEngine"]').evaluate_all("nodes=>nodes.map(node=>node.innerText)")
    preview_table_geometry=page.locator('#componentLayer .component[data-engine="TableEngine"]').evaluate_all("nodes=>nodes.map(node=>{const rect=e=>{const r=e?.getBoundingClientRect();return r?{x:r.x,y:r.y,w:r.width,h:r.height}:null},region=node.querySelector('.table-frame'),table=region?.querySelector('table'),rows=[...(table?.querySelectorAll('tr')||[])];return{id:node.dataset.id,item:rect(node),card:rect(node.querySelector('.gallery-card')),body:rect(node.querySelector('.card-body')),region:rect(region),table:rect(table),rows:rows.map(row=>({text:row.innerText,rect:rect(row),color:getComputedStyle(row).color,display:getComputedStyle(row).display,visibility:getComputedStyle(row).visibility}))}})")
    axis_label_overlaps=page.locator("svg.cs-chart-svg").evaluate_all("svgs=>svgs.map(svg=>{const labels=[...svg.querySelectorAll('.cs-axis-label')].map(node=>{const r=node.getBBox();return{text:node.textContent.trim(),x:r.x,y:r.y,w:r.width,h:r.height}}),overlaps=[];labels.forEach((left,index)=>labels.slice(index+1).forEach(right=>{const x=Math.min(left.x+left.w,right.x+right.w)-Math.max(left.x,right.x),y=Math.min(left.y+left.h,right.y+right.h)-Math.max(left.y,right.y);if(x>1&&y>1)overlaps.push({left:left.text,right:right.text})}));return{chart:svg.closest('.component')?.dataset.id,overlaps}})")
    if any(check["overlaps"] for check in axis_label_overlaps):
        raise AssertionError(f"Chart axis labels must remain legible without collisions: {json.dumps(axis_label_overlaps,ensure_ascii=False)}")
    preview_chart_geometry=page.locator("#componentLayer .component:has(.cs-chart-svg)").evaluate_all("nodes=>nodes.map(node=>{const svg=node.querySelector('.cs-chart-svg'),plot=node.querySelector('.cs-static-chart'),component=node.getBoundingClientRect(),svgRect=svg?.getBoundingClientRect(),plotRect=plot?.getBoundingClientRect(),viewBox=svg?.viewBox?.baseVal;return{id:node.dataset.id,element:node.querySelector('.chart-studio-render')?.dataset.element||'',engine:node.dataset.engine||'',component_width:component.width,component_height:component.height,plot_width:plotRect?.width||0,svg_width:svgRect?.width||0,svg_height:svgRect?.height||0,renderer_aspect:viewBox?.width&&viewBox?.height?viewBox.width/viewBox.height:0,preserve_aspect:svg?.getAttribute('preserveAspectRatio')||'xMidYMid meet'}})")
    clipped_table_rows=[{**table,"clipped_rows":[row for row in table["rows"][1:] if row["rect"]["y"]+row["rect"]["h"]>table["region"]["y"]+table["region"]["h"]+1]} for table in preview_table_geometry if len(table["rows"])>1]
    if any(table["clipped_rows"] for table in clipped_table_rows):
        raise AssertionError(f"Preview must show every authored table row without a hidden scroll requirement: {json.dumps(clipped_table_rows,ensure_ascii=False)}")
    spatial_geometry=page.locator("#componentLayer .cs-wafer-svg").evaluate_all("nodes=>nodes.map(svg=>{const group=svg.querySelector('g[clip-path]'),clip=svg.querySelector('clipPath circle'),marks=[...(group?.querySelectorAll('[data-wafer-die]')||[])],cx=Number(clip?.getAttribute('cx')||0),cy=Number(clip?.getAttribute('cy')||0),r=Number(clip?.getAttribute('r')||0),outside=marks.filter(mark=>{const x=Number(mark.getAttribute('x')),y=Number(mark.getAttribute('y')),w=Number(mark.getAttribute('width')||0),h=Number(mark.getAttribute('height')||0);return [[x,y],[x+w,y],[x,y+h],[x+w,y+h]].some(([px,py])=>(px-cx)**2+(py-cy)**2>r**2)}).length;return{label:svg.getAttribute('aria-label'),viewBox:svg.getAttribute('viewBox'),clip_path_applied:!!group?.getAttribute('clip-path'),clip_geometry:clip?{cx,cy,r}:null,die_count:marks.length,boundary_intersections:outside}})")
    if name=="semiconductor-rca":
        assert spatial_geometry and spatial_geometry[0]["die_count"] >= 60, f"The RCA wafer report needs a realistic circular die population: {spatial_geometry}"
        assert all(entry["clip_path_applied"] and entry["clip_geometry"] for entry in spatial_geometry), f"Wafer die marks must stay inside Company-owned SVG clipping geometry: {spatial_geometry}"
    for chart in preview_chart_geometry:
        # The composition policy sizes the whole card to the renderer's viewBox.
        # Its CSS SVG box also reserves the governed summary row, so testing that
        # box alone confuses card geometry with the actual renderer aspect.
        ratio = chart["component_width"] / max(1, chart["component_height"])
        expected = 1.0 if chart["element"] == "Wafer Map" else 680 / 330
        if abs(ratio - expected) > 0.7 or abs(chart["renderer_aspect"] - expected) > 0.02:
            raise AssertionError(f"Smart chart geometry must honor its Visembler renderer aspect: {chart}")
    page.screenshot(path=str(output / f"{name}-preview-desktop-viewport.png"), full_page=True)
    page.locator("#componentLayer").screenshot(path=str(output / f"{name}-preview-desktop.png"))
    page.set_viewport_size({"width": 390, "height": 844})
    page.locator("#previewFitWidth").click()
    page.wait_for_timeout(300)
    mobile_overflow = page.evaluate("document.documentElement.scrollWidth-innerWidth")
    mobile_reading = page.evaluate("""()=>{const root=document.querySelector('.cui-visualizer-root'),hull=document.querySelector('#hull'),layer=document.querySelector('#componentLayer'),items=[...document.querySelectorAll('#componentLayer>.component')],first=items[0]?.getBoundingClientRect(),second=items[1]?.getBoundingClientRect(),body=document.querySelector('.integrated-element-content'),rect=node=>{const r=node.getBoundingClientRect(),s=getComputedStyle(node);return {id:node.dataset.id,x:r.x,y:r.y,w:r.width,h:r.height,position:s.position,display:s.display,order:s.order,text:(node.innerText||'').slice(0,70)}};const textComponents=items.filter(node=>node.dataset.engine==='TextEngine').map(node=>{const card=node.querySelector('.gallery-card'),copy=node.querySelector('.card-body');return{id:node.dataset.id,componentHeight:node.clientHeight,cardScroll:card?.scrollHeight||0,cardClient:card?.clientHeight||0,copyScroll:copy?.scrollHeight||0,copyClient:copy?.clientHeight||0,deadSpace:Math.max(0,node.clientHeight-(card?.clientHeight||0)),clipped:!!card&&card.scrollHeight>card.clientHeight+2||!!copy&&copy.scrollHeight>copy.clientHeight+2}}),comparisonMetrics=items.filter(node=>node.dataset.engine==='ComparisonEngine').map(node=>{const card=node.querySelector('.gallery-card'),copy=node.querySelector('.card-body'),comparison=node.querySelector('.comparison');return{id:node.dataset.id,componentHeight:node.clientHeight,cardScroll:card?.scrollHeight||0,cardClient:card?.clientHeight||0,copyScroll:copy?.scrollHeight||0,copyClient:copy?.clientHeight||0,comparisonScroll:comparison?.scrollHeight||0,comparisonClient:comparison?.clientHeight||0}}),processFlows=items.filter(node=>node.dataset.engine==='DiagramEngine').map(node=>{const r=node.getBoundingClientRect(),svg=node.querySelector('.diagram-svg'),s=svg?.getBoundingClientRect();return{id:node.dataset.id,height:r.height,svgHeight:s?.height||0,svgWidth:s?.width||0}}),tableRegions=[...document.querySelectorAll('#componentLayer [role="region"][aria-label^="Scrollable table:"]')].map(node=>({label:node.getAttribute('aria-label'),tabIndex:node.tabIndex,clientWidth:node.clientWidth,scrollWidth:node.scrollWidth,scrollLeft:node.scrollLeft})),geometry=window.__VIZ_PROD__.layoutGeometry().items,sections=[...document.querySelectorAll('#componentLayer .composition-section-heading')].map(node=>{const members=geometry.filter(item=>String(item.section)===String(node.dataset.section));return{title:node.textContent.trim(),section:node.dataset.section,order:Number(getComputedStyle(node).order),first_component_order:members.length?Math.min(...members.map(item=>Number(item.order)||0)):null,role:node.getAttribute('role'),level:node.getAttribute('aria-level')}});return {hull_width:hull?.getBoundingClientRect().width||0,viewport_width:innerWidth,component_count:items.length,first_font_px:body?parseFloat(getComputedStyle(body).fontSize):0,ordered_flow:!!first&&!!second&&second.top>first.top+1,preview:root?.classList.contains('preview-mode'),layer:layer?{...rect(layer),flexDirection:getComputedStyle(layer).flexDirection}:null,items:items.map(rect).sort((a,b)=>a.y-b.y),reading_order:items.map(rect).sort((a,b)=>a.y-b.y).map(value=>Number(value.order)),text_components:textComponents,comparison_metrics:comparisonMetrics,process_flows:processFlows,table_regions:tableRegions,sections}}""")
    mobile_reading["process_flow_diagnostics"] = page.evaluate("""()=>[...document.querySelectorAll('#componentLayer>.component[data-engine="DiagramEngine"]')].map(node=>{const metrics=[node,node.querySelector('.c-content'),node.querySelector('.integrated-element-content'),node.querySelector('.gallery-card'),node.querySelector('.card-body'),node.querySelector('.diagram-responsive-reading'),...node.querySelectorAll('.diagram-reading-projection,.diagram-wide-reading')].filter(Boolean).map(child=>{const r=child.getBoundingClientRect(),s=getComputedStyle(child);return{tag:child.tagName,className:String(child.className||''),rect:{x:r.x,y:r.y,width:r.width,height:r.height},style:{display:s.display,width:s.width,height:s.height,minHeight:s.minHeight,maxHeight:s.maxHeight,overflow:s.overflow,flex:s.flex},inlineStyle:child.getAttribute('style')}});return{id:node.dataset.id,inlineStyle:node.getAttribute('style'),mobileHeight:node.style.getPropertyValue('--viz-mobile-reader-height'),metrics}})""")
    assert mobile_reading["hull_width"] <= 370 and mobile_reading["hull_width"] >= 340, "Mobile reading should use the narrow viewport width."
    assert mobile_reading["first_font_px"] >= 12 and mobile_reading["ordered_flow"], "Mobile preview must remain readable in composed single-column order."
    assert mobile_reading["sections"] and all(section["role"] == "heading" and section["level"] == "2" for section in mobile_reading["sections"]), "Mobile reading must retain accessible section hierarchy."
    assert mobile_reading["reading_order"] == sorted(mobile_reading["reading_order"]), f"Mobile Preview must follow Smart composition order: {mobile_reading['reading_order']}"
    assert all(section["first_component_order"] is not None and section["order"] == 2 * section["first_component_order"] for section in mobile_reading["sections"]), f"Each mobile section heading must precede its own first composed component: {mobile_reading['sections']}"
    assert not any(entry["clipped"] for entry in mobile_reading["text_components"]), f"Mobile narrative content must not be clipped: {mobile_reading['text_components']}"
    assert all(entry["deadSpace"] <= 72 for entry in mobile_reading["text_components"]), f"Mobile narrative cards must not reserve large empty bands: {mobile_reading['text_components']}"
    assert all(entry["copyScroll"] <= entry["copyClient"] + 2 and entry["comparisonScroll"] <= entry["comparisonClient"] + 2 for entry in mobile_reading["comparison_metrics"]), f"Mobile comparison metrics must fit their authored card height: {mobile_reading['comparison_metrics']}"
    process_flow_ids={str(entry.get("id")) for entry in model(page).get("items",[]) if entry.get("engine")=="DiagramEngine" and "process flow" in str(entry.get("element",entry.get("title",""))).lower()}
    process_flows=[entry for entry in mobile_reading["process_flows"] if entry["id"] in process_flow_ids]
    assert not process_flow_ids or (len(process_flows)==len(process_flow_ids) and all(140 <= entry["height"] <= 205 and entry["svgHeight"] >= 96 for entry in process_flows)), f"Mobile process flow must remain readable without a blank display band: {process_flows}; {mobile_reading['process_flow_diagnostics']}"
    assert all(mobile_reading["items"][index + 1]["y"] >= mobile_reading["items"][index]["y"] + mobile_reading["items"][index]["h"] - 1 for index in range(len(mobile_reading["items"]) - 1)), f"Mobile preview components must not overlap: {mobile_reading['items']}"
    page.locator("#toast").evaluate("node=>{node.classList.remove('show');node.textContent=''}")
    page.locator("#hull").screenshot(path=str(output / f"{name}-preview-mobile.png"))
    page.screenshot(path=str(output / f"{name}-preview-mobile-viewport.png"), full_page=True)
    if mobile_reading["table_regions"]:
        scroll_index=next((index for index,region in enumerate(mobile_reading["table_regions"]) if region["scrollWidth"]>region["clientWidth"]+1),0)
        table_region=page.locator('#componentLayer [role="region"][aria-label^="Scrollable table:"]').nth(scroll_index)
        assert table_region.get_attribute("tabindex")=="0", "Mobile evidence tables must expose a keyboard-focusable scroll region."
        table_region.focus()
        focus_state=table_region.evaluate("node=>({target:{tag:node.tagName,role:node.getAttribute('role'),label:node.getAttribute('aria-label'),tabIndex:node.tabIndex,connected:node.isConnected,rect:(()=>{const r=node.getBoundingClientRect();return{x:r.x,y:r.y,w:r.width,h:r.height}})(),display:getComputedStyle(node).display,visibility:getComputedStyle(node).visibility,parentInert:node.closest('[inert]')?.outerHTML?.slice(0,300)},active:{tag:document.activeElement?.tagName,role:document.activeElement?.getAttribute('role'),label:document.activeElement?.getAttribute('aria-label'),html:document.activeElement?.outerHTML?.slice(0,300)}})")
        assert focus_state.get("active",{}).get("role")=="region", f"The table scroll region must receive keyboard focus: {focus_state}"
        if mobile_reading["table_regions"][scroll_index]["scrollWidth"]>mobile_reading["table_regions"][scroll_index]["clientWidth"]+1:
            page.keyboard.press("ArrowRight")
            mobile_reading["table_regions"][scroll_index]["keyboard_scroll_left"]=table_region.evaluate("node=>node.scrollLeft")
            assert mobile_reading["table_regions"][scroll_index]["keyboard_scroll_left"]>0, "Arrow keys must scroll a focused evidence table horizontally."
        table_region.evaluate("node=>{node.scrollLeft=0;node.blur()}")
    page.set_viewport_size({"width": 1440, "height": 1000})
    page.locator("#previewExit").click()
    page.locator(".cui-visualizer-root.preview-mode").wait_for(state="detached", timeout=10_000)
    page.keyboard.press("Escape")
    for selector in ("#libraryToggle", "#inspectorToggle"):
        if page.locator(selector).get_attribute("aria-pressed") == "true":
            page.locator(selector).click()
    page.locator("#zoomFit").click()
    actions.note("return to editing", "editor")
    return {"editor_horizontal_overflow_px": editor_scroll, "mobile_preview_overflow_px": mobile_overflow, "mobile_reading": mobile_reading,"preview_chart_summaries":preview_chart_summaries,"preview_chart_mark_titles":preview_chart_mark_titles,"preview_table_summaries":preview_table_summaries,"preview_table_geometry":preview_table_geometry,"preview_chart_geometry":preview_chart_geometry,"axis_label_overlaps":axis_label_overlaps,"spatial_geometry":spatial_geometry}


def add_library_component(page, element: str, actions: Actions) -> str:
    if page.locator("#libraryToggle").get_attribute("aria-pressed") != "true":
        page.locator("#libraryToggle").click()
    search = page.locator("#componentSearch")
    search.fill(element)
    choice = page.locator(f'#fullLibrary .library-item[data-element={json.dumps(element)}]')
    choice.wait_for(timeout=8_000)
    choice.locator(".library-insert").click()
    settled(page)
    search.fill("")
    actions.note(f"insert {element} from component library", "editor")
    return str(model(page)["items"][-1]["id"])


def story_template(page, host: NativeHost, output: Path, template_id: str, actions: Actions) -> dict:
    report_id = create_from_hub(page, host, template_id, actions)
    genre = template_id
    narrative = next(
        (entry for entry in model(page)["items"] if entry.get("composition_role") in {"report_headline", "narrative_interpretation"}),
        None,
    )
    if narrative and narrative.get("engine") == "TextEngine":
        fill_text(page, narrative, TEXT[genre], actions)
    context = next((entry for entry in model(page)["items"] if entry.get("composition_role") == "context"), None)
    if context and context.get("engine") == "TextEngine":
        contexts = {
            "executive-business-review": "Quarterly business review · Q1–Q4 · three regions · revenue and retention.",
            "semiconductor-rca": "Lot L2407 · ETCH-08 chamber A · affected die compared with a paired reference · release held for verification.",
            "experiment-decision": "Eligible new-account cohorts · W1–W4 · randomized control and treatment · conversion and quality guardrail.",
            "technical-status-review": "Line 4 · shifts A–C · chamber A controls · capacity release review.",
        }
        fill_text(page, context, contexts[genre], actions)
        if context.get("showTitle"):
            heading=page.locator(f'.component[data-id="{context["id"]}"] .gallery-card header h3')
            assert heading.is_visible() and heading.inner_text().strip()==context.get("title"), "The authored opening context must retain its visible context heading."
    conclusion = next((entry for entry in model(page)["items"] if entry.get("composition_role") == "conclusion"), None)
    if conclusion and conclusion.get("engine") == "TextEngine":
        conclusions = {
            "executive-business-review": "Next step: close the South retention gap before the next quarterly review.",
            "semiconductor-rca": "Next step: verify chamber A against the paired reference lot before releasing L2407.",
            "experiment-decision": "Decision: stage a limited rollout and pause if the quality guardrail leaves its agreed band.",
            "technical-status-review": "Next step: keep the recipe locked through the paired-lot check and share the result with the owners.",
        }
        fill_text(page, conclusion, conclusions[genre], actions)
    interpretation = next((entry for entry in model(page)["items"] if entry.get("composition_role") == "narrative_interpretation" and entry.get("engine") == "TextEngine"), None)
    if interpretation and genre in INTERPRETATION:
        fill_text(page, interpretation, INTERPRETATION[genre], actions)
    fill_metrics_and_decision(page, genre, actions)
    table_rows = {
        "executive-business-review": [["North", "184", "94.1%", "Watch"], ["South", "156", "92.8%", "Below target"], ["West", "211", "95.0%", "On plan"]],
        "semiconductor-rca": [["Yield distribution", "94.2%", "98.6%", "Mean of observed cohorts"], ["Affected lot L2407", "94.2%", "—", "Localized die loss"], ["ETCH-08 / A", "Pressure +2.1%", "Stable", "Verify chamber"]],
        "experiment-decision": [["Control", "1,198", "8.25%", "Within quality guardrail"], ["Treatment", "1,204", "9.03%", "Within quality guardrail"], ["Holdout", "412", "8.45%", "Monitor mix"]],
        "technical-status-review": [["ETCH-08", "A", "Pressure", "Equipment Eng.", "Open"], ["ETCH-07", "C", "Flow", "Process Eng.", "Stable"], ["MET-02", "B", "Yield", "Quality", "Verified"]],
    }[genre]
    table_filled = fill_table(page, actions, table_rows)
    table_item = next((entry for entry in model(page)["items"] if entry.get("engine") == "TableEngine"), None)
    if table_item and genre == "experiment-decision":
        set_table_headers(page, str(table_item["id"]), ["Cohort", "Eligible", "Conversion", "Quality guardrail"], actions)
        set_component_title(page, table_item, "Cohort evidence", actions)
    if genre == "semiconductor-rca":
        wafer = next((entry for entry in model(page)["items"] if entry.get("engine") == "WaferFabEngine"), None)
        if wafer:
            import_or_replace_data(page, DATA[genre], "wafer", actions, str(wafer["id"]))
        distribution = next((entry for entry in model(page)["items"] if entry.get("engine") == "CoreChartEngine" and "Box" in entry.get("element", "")), None)
        if distribution:
            import_or_replace_data(page, DATA["semiconductor-distribution"], "box", actions, str(distribution["id"]))
        chart = next((entry for entry in model(page)["items"] if entry.get("engine") == "EngineeringChartEngine"), None)
        if chart:
            select_item(page, str(chart["id"]))
            if page.locator("#iObservations").count():
                fill_control(page, "#iObservations", "W1\t98.8\nW2\t98.4\nW3\t94.1\nW4\t94.6", actions, "SPC observations")
        finding = next((entry for entry in model(page)["items"] if entry.get("engine") == "EvidenceCompositeEngine"), None)
        if finding:
            select_item(page, str(finding["id"]))
            for selector, value in (
                ("#iStatement", "Finding · chamber A pressure drift aligns with the L2407 yield excursion."),
                ("#iDetail", "Affected yield is 94.2% versus 98.6% reference. Contain L2407 and verify the matched chamber run before release."),
                ("#iStatus", "Verified"),
            ):
                fill_control(page, selector, value, actions, selector[1:])
            set_component_title(page, finding, "Verified finding", actions)
    else:
        line = next((entry for entry in model(page)["items"] if entry.get("engine") == "CoreChartEngine" and "Line" in entry.get("element", "")), None)
        if not line:
            raise AssertionError(f"{genre} template must provide its intended trend visual.")
        import_or_replace_data(page, DATA[genre], "line", actions, str(line["id"]))
        if genre == "experiment-decision":
            distribution = next((entry for entry in model(page)["items"] if entry.get("engine") == "CoreChartEngine" and "Box" in entry.get("element", "")), None)
            if distribution:
                import_or_replace_data(page, DATA["experiment-distribution"], "box", actions, str(distribution["id"]))
        if genre == "technical-status-review":
            spc = next((entry for entry in model(page)["items"] if entry.get("engine") == "EngineeringChartEngine"), None)
            if spc:
                import_or_replace_data(page, DATA["technical-spc"], "engineering", actions, str(spc["id"]))
        actions.note("populate the template's intended trend and distribution visuals")
        if genre == "executive-business-review":
            line = next(entry for entry in model(page)["items"] if entry.get("engine") == "CoreChartEngine" and "Line" in entry.get("element", ""))
            select_item(page, str(line["id"]))
            mapping = page.locator('[data-dataset-role="y"]')
            if mapping.count():
                field = next(value for value in model(page)["datasets"] if value["id"] == line["dataset_id"])
                retention = next(value for value in field["fields"] if value["name"] == "Retention")
                mapping.select_option(str(retention["id"]))
                settled(page)
                actions.note("map Retention as the trend value", "inspector")
    settled(page)
    record = host.repository.get(report_id)
    assert record.model["mode"] == "smart", f"{genre} must use its Smart template defaults."
    assert len(record.model["items"]) >= 5
    if template_id == "executive-business-review":
        assert not any(entry.get("x") is not None for entry in record.model["items"]), "The intended Smart result must not require pixel placement."
    table_record = next((entry for entry in record.model["items"] if entry.get("engine") == "TableEngine"), None)
    saved_table_rows = table_record.get("customTable", {}).get("rows", []) if table_record else []
    assert table_filled and len(saved_table_rows) == len(table_rows), f"{genre} must retain its complete authored evidence rows: {saved_table_rows!r}"
    preflight = compose_and_validate(page, actions)
    screenshots = capture_report(page, output, genre, actions)
    # Recomposition can update the fit after returning from Preview. Capture
    # the final editor state as a whole report after the semantic reflow.
    for selector in ("#libraryToggle", "#inspectorToggle"):
        if page.locator(selector).get_attribute("aria-pressed") == "true":
            page.locator(selector).click()
    page.keyboard.press("Escape")
    page.locator("#zoomFit").click()
    page.locator("#toast").evaluate("node=>{node.classList.remove('show');node.textContent=''}")
    capture_editor_canvas(page, output / f"{genre}-editor-desktop-final.png")
    pptx_texts = export_pptx(page, output / f"{genre}.pptx", actions, host)
    export_svg(page, output / f"{genre}.svg", actions)
    return {
        "report_id": report_id,
        "title": record.title,
        "items": len(record.model["items"]),
        "sections": sorted({str(entry["section_title"]) for entry in record.model["items"] if entry.get("section_title")}),
        "roles": sorted({str(entry["composition_role"]) for entry in record.model["items"] if entry.get("composition_role")}),
        "datasets": len(record.model["datasets"]),
        "table_evidence_filled": table_filled,
        "preflight": preflight,
        "pptx_slides": len(Presentation(str(output / f"{genre}.pptx")).slides),
        "pptx_chart_values": pptx_chart_values(output / f"{genre}.pptx"),
        "pptx_text_samples": pptx_texts[:12],
        **screenshots,
    }


def save_report_preset(page, name: str, actions: Actions, data_mode: str = "structure") -> None:
    if page.locator("#libraryToggle").get_attribute("aria-pressed") != "true":
        page.locator("#libraryToggle").click()
    page.locator("#presetsTab").click()
    actions.note("open reusable preset library", "presets")
    page.locator("#presetSave").click()
    actions.note("open save report preset", "modal")
    page.locator("#presetSaveName").fill(name)
    page.locator(f'input[name="presetDataMode"][value="{data_mode}"]').check()
    page.locator("#presetSaveForm button[type=submit]").click()
    page.locator("#genericModal.show").wait_for(state="hidden", timeout=5_000)
    actions.note("save reusable report structure")


def save_section_preset(page, ids: list[str], name: str, actions: Actions) -> None:
    select_item(page, ids[0])
    select_item(page, ids[1], additive=True)
    if page.locator("#libraryToggle").get_attribute("aria-pressed") != "true":
        page.locator("#libraryToggle").click()
    page.locator("#presetsTab").click()
    actions.note("open reusable preset library", "presets")
    button = page.locator("#presetSaveSelection")
    page.wait_for_function("() => { const button=document.querySelector('#presetSaveSelection'); return !!button && !button.disabled; }", timeout=5_000)
    button.click()
    actions.note("open save section preset", "modal")
    page.locator("#sectionPresetSaveName").fill(name)
    page.locator('input[name="sectionDataMode"][value="structure"]').check()
    page.locator("#sectionPresetSaveForm button[type=submit]").click()
    page.locator("#genericModal.show").wait_for(state="hidden", timeout=5_000)
    actions.note("save reusable section structure")


def open_presets(page, actions: Actions) -> None:
    if page.locator("#libraryToggle").get_attribute("aria-pressed") != "true":
        page.locator("#libraryToggle").click()
    page.locator("#presetsTab").click()
    actions.note("open personal preset library", "presets")


def preset_action(page, preset_name: str, attribute: str):
    cards = page.locator("#presetList .preset")
    for index in range(cards.count()):
        card = cards.nth(index)
        if card.locator(".preset-name-edit").input_value() == preset_name:
            return card.locator(f"[{attribute}]")
    raise AssertionError(f"Saved preset {preset_name!r} was not present in the personal preset library.")


def create_report_reuse_source(page, host: NativeHost, output: Path, actions: Actions) -> tuple[str, str, str, str, dict]:
    report_id = create_from_hub(page, host, "executive-business-review", actions)
    for chart in [entry for entry in model(page)["items"] if entry.get("engine") == "CoreChartEngine"]:
        delete_component(page, str(chart["id"]), actions)
    chart_id = import_or_replace_data(page, "Region\tRevenue\nEast\t80\nWest\t95\n", "line", actions)
    select_item(page, chart_id)
    page.locator('[data-inspector="duplicate"]').click()
    settled(page)
    charts = [entry for entry in model(page)["items"] if entry.get("engine") == "CoreChartEngine" and entry.get("dataset_id")]
    assert len(charts) == 2 and len({entry["dataset_id"] for entry in charts}) == 1
    source_dataset_id = str(charts[0]["dataset_id"])
    for chart in charts:
        set_composition_role(page, chart, "supporting_analysis", actions)
    source_report = finish_regional_revenue_report(page, "regional-reuse-source", actions)
    source_preflight = compose_and_validate(page, actions)
    source_pptx = export_pptx(page, output / "reuse-source-report.pptx", actions, host)
    source_chart_values = pptx_chart_values(output / "reuse-source-report.pptx")
    assert any(80.0 in series["values"] and 95.0 in series["values"] for series in source_chart_values)
    export_svg(page, output / "reuse-source-report.svg", actions)
    source_capture = capture_report(page, output, "reuse-source-report", actions)
    assert any("2 rows" in summary and "Region" in summary and "Revenue" in summary for summary in source_capture["preview_chart_summaries"])
    assert any(all(value in summary for value in ("East", "80", "West", "95")) for summary in source_capture["preview_table_summaries"]), f"The reusable source report must render its evidence rows in Preview: {source_capture['preview_table_summaries']}"
    save_report_preset(page, "Regional Revenue Review", actions)
    save_report_preset(page, "Regional Revenue Source Copy", actions, data_mode="copy")
    save_section_preset(page, [str(charts[0]["id"]), str(charts[1]["id"])], "Regional Comparison Pair", actions)
    settled(page)
    source_before = copy.deepcopy(host.repository.get(report_id).model)
    source_digest = hashlib.sha256(json.dumps(source_before, sort_keys=True).encode()).hexdigest()

    copy_target_id = create_from_hub(page, host, "blank", actions)
    open_presets(page, actions)
    preset_action(page, "Regional Revenue Source Copy", "data-loadpreset").click()
    copy_confirmation = page.locator("#presetSourceConfirmForm")
    copy_confirmation.wait_for(timeout=5_000)
    assert "source rows" in copy_confirmation.inner_text().lower(), "Source-copy confirmation must disclose the copied data source and row count before applying it."
    actions.note("review explicit source-data copy consequences", "modal")
    copy_confirmation.get_by_role("button", name="Apply report", exact=True).click()
    settled(page)
    copy_model = model(page)
    copy_datasets = {str(dataset["id"]): dataset for dataset in copy_model.get("datasets", [])}
    assert source_dataset_id in copy_datasets and copy_datasets[source_dataset_id]["rows"] == [["East", 80], ["West", 95]]
    assert host.repository.get(report_id).model == source_before, "Applying explicit source-data copy must leave its saved source report unchanged."
    actions.note("apply report preset with source data", "editor")
    return report_id, source_dataset_id, source_digest, copy_target_id, {
        "content": source_report["content"],
        "saved_rows": source_report["saved_rows"],
        "preflight": source_preflight,
        "pptx_texts": source_pptx,
        "pptx_chart_values": source_chart_values,
        **source_capture,
    }


def create_blank_with_dataset(page, host: NativeHost, text: str, actions: Actions, name: str) -> tuple[str, str]:
    report_id = create_from_hub(page, host, "blank", actions)
    chart_id = import_or_replace_data(page, text, "line", actions)
    dataset_id = str(next(entry for entry in model(page)["items"] if entry["id"] == chart_id)["dataset_id"])
    return report_id, dataset_id


def remap_report_challenge(page, host: NativeHost, output: Path, actions: Actions) -> dict:
    source_id, source_dataset_id, source_digest, copy_target_id, source_capture = create_report_reuse_source(page, host, output, actions)
    source_model = copy.deepcopy(host.repository.get(source_id).model)
    followup_start = len(actions.events)
    target_id, destination_id = create_blank_with_dataset(
        page, host, "Region\tRevenue\nEast\t180\nWest\t205\n", actions, "reuse-target"
    )
    target_before = copy.deepcopy(model(page))
    open_presets(page, actions)
    trigger = preset_action(page, "Regional Revenue Review", "data-reusepreset")
    trigger.click()
    actions.note("open guided report remap", "modal")
    dialog = page.locator("#reuseRemapForm")
    dialog.wait_for(timeout=5_000)
    page.locator("#genericModal.show").screenshot(path=str(output / "reuse-remap-guided-dialog.png"))
    guided_text = dialog.inner_text()
    for required_label in ("Values to review", "Key metric", "Comparison", "Evidence table"):
        assert required_label.lower() in guided_text.lower(), f"The remap step must disclose manual values that need destination-specific review: {required_label}; dialog={guided_text}"
    assert dialog.get_by_role("button",name="Add a data source",exact=True).is_visible(), "Remap must expose the supported Data First intake path for a newly imported source."
    focused = page.evaluate("()=>document.activeElement?.closest('#reuseRemapForm')!==null")
    assert focused, "The remap workflow must receive keyboard focus."
    field_labels = dialog.locator("[data-remap-role]").count()
    source_slot_text = dialog.inner_text()
    assert "Region" in source_slot_text and "Revenue" in source_slot_text
    assert "used by" in source_slot_text.lower()
    assert "used by 2 visuals" in source_slot_text.lower(), "The shared dataset slot must name both dependent visuals once."
    assert field_labels == 3, f"Shared Category and Measurement requirements should be shown once per distinct source requirement: {field_labels}"
    shared_control = dialog.locator("[data-remap-role]").first
    assignments = json.loads(shared_control.get_attribute("data-remap-assignments") or "[]")
    assert len(assignments) == 2, f"One shared selector must update both visual bindings: {assignments}"
    shared_control.select_option(shared_control.input_value())
    page.wait_for_timeout(80)
    focused_shared = page.evaluate("""()=>({item:document.activeElement?.dataset?.remapItem,role:document.activeElement?.dataset?.role,slot:document.activeElement?.dataset?.remapRole})""")
    assert focused_shared.get("item") == assignments[0]["itemId"], f"Field selection must restore keyboard focus after grouped remap refresh: {focused_shared}"
    grouped_focus_restored = True
    page.keyboard.press("Tab")
    inside_after_tab = page.evaluate("()=>document.activeElement?.closest('#reuseRemapForm')!==null")
    assert inside_after_tab, "Tab must remain in the remap modal."
    page.keyboard.press("Escape")
    dialog.wait_for(state="hidden", timeout=5_000)
    focus_returned = page.evaluate("()=>document.activeElement?.matches('[data-reusepreset]')")
    assert focus_returned, "Escape must return focus to the reuse action."
    actions.note("verify remap keyboard focus and Escape return", "presets")
    preset_action(page, "Regional Revenue Review", "data-reusepreset").click()
    dialog = page.locator("#reuseRemapForm")
    dialog.wait_for(timeout=5_000)
    page.get_by_role("button", name="Add a data source", exact=True).click()
    page.locator("#genericModal.show .data-first-dialog").wait_for(timeout=5_000)
    actions.note("open supported Data First intake from remap", "data-first")
    page.locator("#genericModal.show #modalBody [data-close]").click()
    page.locator("#genericModal.show").wait_for(state="hidden", timeout=5_000)
    preset_action(page, "Regional Revenue Review", "data-reusepreset").click()
    dialog = page.locator("#reuseRemapForm")
    dialog.wait_for(timeout=5_000)
    role_controls = dialog.locator("[data-remap-role]")
    auto_mapping_count = role_controls.count()
    assert dialog.get_by_role("button", name="Apply reusable structure", exact=True).is_enabled()
    dialog.get_by_role("button", name="Apply reusable structure", exact=True).click()
    settled(page)
    dialog.wait_for(state="hidden", timeout=5_000)
    actions.note("apply remapped report preset", "editor")
    remapped = model(page)
    reused_charts = [entry for entry in remapped["items"] if entry.get("engine") == "CoreChartEngine" and entry.get("dataset_id")]
    assert len(reused_charts) >= 2
    assert {str(entry["dataset_id"]) for entry in reused_charts} == {destination_id}
    destinations = {str(dataset["id"]): dataset for dataset in remapped["datasets"]}
    assert source_dataset_id not in destinations
    assert destinations[destination_id]["rows"] == [["East", 180], ["West", 205]]
    serialized = json.dumps(remapped, ensure_ascii=False)
    assert source_dataset_id not in serialized, "A structure-only remap must not retain the source report's dataset identity."
    assert json.dumps([["East", 180], ["West", 205]]) in serialized
    assert json.dumps([["East", 80], ["West", 95]]) not in serialized
    assert not any(
        item.get("data") == [["East", 80], ["West", 95]]
        for item in remapped["items"]
    ), "Cached chart values from the source report must not survive structure-only reuse."
    assert host.repository.get(source_id).model == source_before_model(source_id, host, source_digest)
    editor_chart_summaries=page.locator("#componentLayer .cs-static-summary").all_inner_texts()
    remapped_chart_diagnostics=[{"id":entry.get("id"),"element":entry.get("element"),"view_type":entry.get("view_type"),"mapping":entry.get("mapping"),"chart_type":entry.get("chart_studio",{}).get("chart_type"),"chart_mapping":entry.get("chart_studio",{}).get("mapping"),"dataset_id":entry.get("dataset_id")} for entry in remapped["items"] if entry.get("dataset_id")]
    assert any("2 rows" in summary and "Region" in summary and "Revenue" in summary for summary in editor_chart_summaries), f"The remapped editor must render the destination source's chart projection: {editor_chart_summaries}; mappings={remapped_chart_diagnostics}"
    initial_revision = state(page)["revision"]
    page.locator("#undo").click()
    settled(page)
    assert model(page) == target_before, "Remap must be one atomic, undoable report operation."
    page.locator("#redo").click()
    settled(page)
    assert {str(entry["dataset_id"]) for entry in model(page)["items"] if entry.get("engine") == "CoreChartEngine" and entry.get("dataset_id")} == {destination_id}
    page.reload(wait_until="domcontentloaded")
    ready(page)
    assert state(page)["revision"] >= initial_revision
    assert {str(entry["dataset_id"]) for entry in model(page)["items"] if entry.get("dataset_id")} == {destination_id}
    actions.note("verify undo, redo, autosave and reload", "editor")
    followup_report = finish_regional_revenue_report(page, "regional-reuse-followup", actions)
    followup_model = model(page)
    assert not any(entry.get("value") in {80, 95, "80", "95"} or entry.get("before") in {80, 95, "80", "95"} or entry.get("after") in {80, 95, "80", "95"} for entry in followup_model["items"]), "Destination-specific manual values must replace source values before the report is committed."
    assert followup_report["saved_rows"] == [["East", 180, "Comparison baseline", "Destination extract"], ["West", 205, "Leads by 25", "Destination extract"]], followup_report["saved_rows"]
    assert any(entry.get("value") in {385, "385"} for entry in followup_model["items"] if entry.get("engine") == "MetricEngine")
    assert any(entry.get("before") in {180, "180"} and entry.get("after") in {205, "205"} for entry in followup_model["items"] if entry.get("engine") == "ComparisonEngine")
    followup_assembly_end = len(actions.events)
    save_report_preset(page, "Regional Revenue Review Copy Check", actions)
    history_report = page.evaluate("new URL(location.href).searchParams.get('report')")
    page.goto(f"{host.url}/visualizer/reports?report={quote(history_report)}", wait_until="domcontentloaded")
    card = page.locator(f'[data-testid="report-card"][data-report-id="{history_report}"]')
    card.wait_for(timeout=10_000)
    card.locator('[data-report-action="more"]').click()
    page.locator(".q-menu:visible").get_by_role("button", name="Review history", exact=True).click()
    entries = page.locator('[data-testid="report-history"] [data-history-id]')
    entries.first.wait_for(timeout=8_000)
    assert entries.count() >= 2
    actions.note("verify governed history after reuse", "history")
    page.goto(f"{host.url}/visualizer?report={quote(target_id)}", wait_until="domcontentloaded")
    ready(page)
    remapped_powerpoint_preview = export_pptx(page, output / "reuse-remapped-report.pptx", actions, host)
    report_chart_values = pptx_chart_values(output / "reuse-remapped-report.pptx")
    assert any("Regional revenue" in value or "West revenue" in value for value in remapped_powerpoint_preview)
    report_series_values = [value for series in report_chart_values for value in series["values"]]
    assert 180.0 in report_series_values and 205.0 in report_series_values, f"PowerPoint must carry the selected destination values: {report_chart_values}"
    assert 80.0 not in report_series_values and 95.0 not in report_series_values, f"PowerPoint must omit source values: {report_chart_values}"
    capture = capture_report(page, output, "reuse-remapped-report", actions)
    assert any("2 rows" in summary and "Region" in summary and "Revenue" in summary for summary in capture["preview_chart_summaries"]), f"Preview must show the destination source's chart projection: {capture['preview_chart_summaries']}"
    rendered_mark_text = json.dumps(capture["preview_chart_mark_titles"])
    rendered_mark_values = exact_numeric_tokens(rendered_mark_text)
    assert {180.0, 205.0} <= rendered_mark_values and not ({80.0, 95.0} & rendered_mark_values), f"The remapped Preview must plot destination values and omit source values: {capture['preview_chart_mark_titles']}"
    assert any(all(value in summary for value in ("East", "180", "West", "205")) for summary in capture["preview_table_summaries"]), f"The remapped report must render destination evidence rows in Preview: {capture['preview_table_summaries']}"
    source_after = host.repository.get(source_id).model
    assert hashlib.sha256(json.dumps(source_after, sort_keys=True).encode()).hexdigest() == source_digest

    section_target_id, section_dataset_id = create_blank_with_dataset(
        page, host, "Region\tRevenue\nEast\t280\nWest\t305\n", actions, "section-reuse-target"
    )
    section_before = copy.deepcopy(model(page))
    open_presets(page, actions)
    preset_action(page, "Regional Comparison Pair", "data-reusepreset").click()
    section_dialog = page.locator("#reuseRemapForm")
    section_dialog.wait_for(timeout=5_000)
    assert section_dialog.get_by_role("button", name="Apply reusable structure", exact=True).is_enabled()
    assert "used by" in section_dialog.inner_text().lower()
    section_dialog.get_by_role("button", name="Apply reusable structure", exact=True).click()
    settled(page)
    actions.note("apply remapped section preset", "editor")
    section_charts = [entry for entry in model(page)["items"] if entry.get("engine") == "CoreChartEngine" and entry.get("dataset_id")]
    new_section_charts = [entry for entry in section_charts if str(entry["id"]) not in {str(item["id"]) for item in section_before["items"]}]
    assert len(new_section_charts) == 2, f"The shared section preset must append both bound visuals; existing={section_before['items']}, after={model(page)['items']}"
    assert {str(entry["dataset_id"]) for entry in new_section_charts} == {section_dataset_id}
    assert all(entry.get("data") != [["East", 80], ["West", 95]] for entry in new_section_charts)
    section_revision = state(page)["revision"]
    page.locator("#undo").click()
    settled(page)
    assert model(page) == section_before, "Section remap must also commit as one undoable edit."
    page.locator("#redo").click()
    settled(page)
    assert state(page)["revision"] >= section_revision
    export_pptx(page, output / "reuse-remapped-section.pptx", actions, host)
    section_chart_values = pptx_chart_values(output / "reuse-remapped-section.pptx")
    section_series_values = [value for series in section_chart_values for value in series["values"]]
    assert 280.0 in section_series_values and 305.0 in section_series_values and 80.0 not in section_series_values
    assert any("East" in series["categories"] and "West" in series["categories"] for series in section_chart_values), f"PowerPoint categories must retain the new destination labels: {section_chart_values}"
    section_capture = capture_report(page, output, "reuse-remapped-section", actions)
    section_mark_text = json.dumps(section_capture["preview_chart_mark_titles"])
    section_mark_values = exact_numeric_tokens(section_mark_text)
    assert {280.0, 305.0} <= section_mark_values and not ({80.0, 95.0} & section_mark_values), f"The remapped section Preview must use only its destination data: {section_capture['preview_chart_mark_titles']}"

    page.goto(f"{host.url}/visualizer?report={quote(target_id)}", wait_until="domcontentloaded")
    ready(page)
    conflict_peer = page.context.new_page()
    conflict_peer.set_default_timeout(8_000)
    conflict_peer.goto(f"{host.url}/visualizer?report={quote(target_id)}", wait_until="domcontentloaded")
    ready(conflict_peer)
    winner_headline = next(entry for entry in model(page)["items"] if entry.get("composition_role") == "report_headline" and entry.get("engine") == "TextEngine")
    stale_headline = next(entry for entry in model(conflict_peer)["items"] if entry.get("composition_role") == "report_headline" and entry.get("engine") == "TextEngine")
    fill_text(page, winner_headline, "April review published from the remapped data source.", actions)
    fill_text(conflict_peer, stale_headline, "Stale edit from an older remapped revision.", actions)
    settled(conflict_peer)
    conflict_status = conflict_peer.locator("#saveStatus").inner_text()
    conflict_state = state(conflict_peer)
    authoritative = host.repository.get(target_id).model
    authoritative_headline = next(entry for entry in authoritative["items"] if entry.get("composition_role") == "report_headline" and entry.get("engine") == "TextEngine")
    assert "Local edits need recovery" in conflict_status and conflict_state["recovery"]
    assert authoritative_headline.get("text") == "April review published from the remapped data source."
    remapped_chart = next(entry for entry in authoritative["items"] if entry.get("engine") == "CoreChartEngine" and entry.get("dataset_id"))
    assert str(remapped_chart["dataset_id"]) == destination_id
    conflict_peer.screenshot(path=str(output / "reuse-remap-conflict-recovery.png"), full_page=True)
    conflict_peer.close()
    actions.note("verify conflict recovery after remap", "editor")
    return {
        "source_report_id": source_id,
        "source_dataset_id": source_dataset_id,
        "source_data_copy_target_report_id": copy_target_id,
        "source_report_capture": source_capture,
        "explicit_source_copy_preserved_rows": [["East", 80], ["West", 95]],
        "destination_report_id": target_id,
        "destination_dataset_id": destination_id,
        "source_dataset_identity_absent_from_remapped_model": source_dataset_id not in serialized,
        "editor_destination_dataset_rows": copy.deepcopy(destinations[destination_id]["rows"]),
        "preview_captured_after_remap": True,
        "editor_chart_summaries_after_remap":editor_chart_summaries,
        "preview_chart_summaries_after_remap":capture["preview_chart_summaries"],
        "reused_shared_visual_count": len(reused_charts),
        "automatic_field_control_count": auto_mapping_count,
        "shared_slot_presentation": {
            "dependent_visual_count": 2,
            "distinct_field_controls": field_labels,
            "first_control_assignments": assignments,
            "focus_restored_after_field_change": grouped_focus_restored,
        },
        "section_reuse": {
            "destination_dataset_id": section_dataset_id,
            "remapped_visual_count": len(new_section_charts),
            "shared_dataset_slot_preserved": True,
            "undo_redo": True,
            "new_values": [["East", 280], ["West", 305]],
            "powerpoint_series_values": section_chart_values,
            **section_capture,
        },
        "structure_reused": [
            {"element": entry.get("element"), "role": entry.get("composition_role"), "section": entry.get("section_id")}
            for entry in source_model["items"]
        ],
        "analytical_bindings_reused": [
            {
                "element": entry.get("element"),
                "mapping_roles": sorted((entry.get("mapping") or {}).keys()),
                "field_names": sorted(
                    next((field["name"] for dataset in source_model.get("datasets", []) if str(dataset.get("id")) == str(entry.get("dataset_id")) for field in dataset.get("fields", []) if str(field.get("id")) == str(field_id)), str(field_id))
                    for field_id in (entry.get("mapping") or {}).values()
                ),
                "transforms": copy.deepcopy(entry.get("transform_recipe") or entry.get("transform_pipeline") or []),
                "analysis_recipe": copy.deepcopy(entry.get("analysis_recipe") or {}),
            }
            for entry in source_model["items"] if entry.get("dataset_id")
        ],
        "style_reused": copy.deepcopy(source_model.get("theme") or source_model.get("style") or {}),
        "style_and_composition_reused": {
            "layout_recipe": source_model.get("layoutPreset", "editorial"),
            "theme": source_model.get("theme"),
            "item_styles": [
                {"element": entry.get("element"), "role": entry.get("composition_role"), "emphasis": entry.get("emphasis"), "density": entry.get("contentDensity"), "style": copy.deepcopy(entry.get("style") or {})}
                for entry in source_model["items"]
            ],
        },
        "manual_values_reentered_for_new_data": {
            "metrics": [entry for entry in followup_model["items"] if entry.get("engine") == "MetricEngine"],
            "comparison": [entry for entry in followup_model["items"] if entry.get("engine") == "ComparisonEngine"],
            "evidence_rows": followup_report["saved_rows"],
            "narrative": followup_report["content"],
        },
        "source_headline_reused_then_updated": next((entry.get("text", "") for entry in source_model["items"] if entry.get("composition_role") == "report_headline"), ""),
        "remapped_field_names": {"Category": "Region", "Value": "Revenue"},
        "explicit_field_choices_required": [],
        "automatic_compatible_mapping_steps": 0,
        "followup_assembly_user_actions": len(actions.events[followup_start:followup_assembly_end]),
        "followup_assembly_context_switches": sum(event["action"] == "context switch" for event in actions.events[followup_start:followup_assembly_end]),
        "followup_assembly_modal_openings": sum(event["action"].startswith("open ") for event in actions.events[followup_start:followup_assembly_end]),
        "followup_manual_component_inserts": sum("insert " in event["action"] for event in actions.events[followup_start:followup_assembly_end]),
        "narrative_intentionally_changed_after_reuse": followup_report["content"],
        "guided_manual_value_review_labels": ["Key metric", "Comparison", "Evidence table"],
        "explicit_source_data_copy_mode": True,
        "source_copy_data_rows": [["East", 80], ["West", 95]],
        "stale_manual_values_removed_before_destination_authoring": True,
        "reused_journey_user_actions": actions.summary()["user_actions"],
        "reused_journey_context_switches": actions.summary()["context_switches"],
        "reused_journey_modal_openings": actions.summary()["modal_openings"],
        "source_data_unchanged": True,
        "destination_rows": destinations[destination_id]["rows"],
        "new_value_visible_in_powerpoint": True,
        "stale_value_absent_from_powerpoint": True,
        "undo_redo_reload_history": True,
        "remap_conflict_preserves_authoritative_edit_and_recovers_stale_draft": True,
        "powerpoint_series_values": report_chart_values,
        **capture,
    }


def source_before_model(source_id: str, host: NativeHost, digest: str) -> dict:
    current = host.repository.get(source_id).model
    actual = hashlib.sha256(json.dumps(current, sort_keys=True).encode()).hexdigest()
    assert actual == digest, "The source report changed while its preset was reused."
    return current


def remap_ambiguity_challenge(page, host: NativeHost, actions: Actions, output: Path | None = None) -> dict:
    source_id = create_from_hub(page, host, "blank", actions)
    source_chart = import_or_replace_data(
        page, "Region\tRevenue\nEast\t50\nWest\t70\n", "line", actions
    )
    select_item(page, source_chart)
    open_presets(page, actions)
    page.locator("#presetSave").click()
    page.locator("#presetSaveName").fill("Ambiguity Check")
    page.locator("#presetSaveForm button[type=submit]").click()
    page.locator("#genericModal.show").wait_for(state="hidden")
    actions.note("save ambiguous-schema reusable report", "editor")
    target_id, _ = create_blank_with_dataset(
        page, host, "Region\tREGION\tRevenue\nEast\tEast\t150\nWest\tWest\t170\n", actions, "ambiguous"
    )
    destination_before = copy.deepcopy(model(page))
    open_presets(page, actions)
    preset_action(page, "Ambiguity Check", "data-reusepreset").click()
    dialog = page.locator("#reuseRemapForm")
    dialog.wait_for(timeout=5_000)
    if output:
        page.locator("#genericModal.show").screenshot(path=str(output / "reuse-remap-ambiguous-dialog.png"))
    apply = dialog.get_by_role("button", name="Apply reusable structure", exact=True)
    assert apply.is_disabled(), "Ambiguous normalized names must stop for an explicit field choice."
    assert "ambiguous" in dialog.inner_text().lower() or "choose" in dialog.inner_text().lower()
    unresolved = dialog.locator("[data-remap-role]")
    assert unresolved.count() > 0
    select = unresolved.first
    labels = [option for option in select.locator("option").all_text_contents() if option.strip()]
    assert any("Region" in option for option in labels)
    # Select the field shown as Region explicitly; the duplicate REGION remains
    # a distinct candidate until the author makes this visible choice.
    region_option = select.locator("option").filter(has_text="Region").first
    select.select_option(region_option.get_attribute("value"))
    settled(page)
    assert not apply.is_disabled(), "The explicit compatible choice should make the remap ready."
    actions.note("explicitly choose an ambiguous Category field", "modal")
    apply.click()
    settled(page)
    assert model(page) != destination_before
    return {
        "source_report_id": source_id,
        "target_report_id": target_id,
        "blocked_before_choice": True,
        "field_choice_required": True,
        "apply_ready_after_choice": True,
    }


def incompatible_challenge(page, host: NativeHost, actions: Actions, output: Path | None = None) -> dict:
    source_id = create_from_hub(page, host, "blank", actions)
    source_chart = import_or_replace_data(page, "Region\tRevenue\nEast\t50\nWest\t70\n", "line", actions)
    select_item(page, source_chart)
    open_presets(page, actions)
    page.locator("#presetSave").click()
    page.locator("#presetSaveName").fill("Type Check")
    page.locator("#presetSaveForm button[type=submit]").click()
    page.locator("#genericModal.show").wait_for(state="hidden")
    actions.note("save incompatible-schema reusable report", "editor")
    target_id, _ = create_blank_with_dataset(page, host, "Region\tRevenue\tObservation\nEast\t2026-09-22\t50\nWest\t2026-09-23\t70\n", actions, "incompatible")
    before = copy.deepcopy(model(page))
    open_presets(page, actions)
    preset_action(page, "Type Check", "data-reusepreset").click()
    dialog = page.locator("#reuseRemapForm")
    dialog.wait_for(timeout=5_000)
    incompatible_dataset = (model(page).get("datasets") or [])[-1]
    dialog.locator("[data-remap-dataset]").select_option(str(incompatible_dataset["id"]))
    settled(page)
    if output:
        page.locator("#genericModal.show").screenshot(path=str(output / "reuse-remap-incompatible-dialog.png"))
    apply = dialog.get_by_role("button", name="Apply reusable structure", exact=True)
    assert apply.is_disabled(), "An incompatible value field must block the report remap."
    compatibility_text = dialog.inner_text()
    assert "incompatible" in compatibility_text.lower() or "numeric" in compatibility_text.lower(), compatibility_text
    issue_messages = dialog.locator(".reuse-remap-issues li").all_inner_texts()
    assert issue_messages and len(issue_messages) == len(set(issue_messages)), f"Remap diagnostics must be de-duplicated: {issue_messages}"
    assert not any("incompatible type incompatible type" in message.lower() for message in issue_messages)
    assert any("measurement" in message.lower() or "numeric" in message.lower() for message in issue_messages), issue_messages
    page.keyboard.press("Escape")
    dialog.wait_for(state="hidden", timeout=5_000)
    assert model(page) == before
    actions.note("confirm incompatible data remains uncommitted", "editor")
    return {"source_report_id": source_id, "target_report_id": target_id, "destination_dataset_id": incompatible_dataset["id"], "compatibility_message": compatibility_text, "blocked": True, "report_unchanged": True}


def blank_rebuild_comparison(page, host: NativeHost, output: Path, actions: Actions) -> dict:
    report_id = create_from_hub(page, host, "blank", actions)
    chart_id = import_or_replace_data(
        page, "Region\tRevenue\nEast\t180\nWest\t205\n", "line", actions
    )
    duplicate = model(page)["items"][-1]
    select_item(page, str(duplicate["id"]))
    page.locator('[data-inspector="duplicate"]').click()
    settled(page)
    for chart in [entry for entry in model(page)["items"] if entry.get("engine") == "CoreChartEngine"]:
        set_composition_role(page, chart, "supporting_analysis", actions)
    inserted = [add_library_component(page, element, actions) for element in (
        "Hero Title", "Body Narrative", "Hero KPI", "Hero KPI", "Before/After KPI",
        "Clean Table", "Risk Callout", "Body Narrative", "Key Takeaway",
    )]
    text_roles = {
        inserted[0]: "report_headline",
        inserted[1]: "context",
        inserted[7]: "narrative_interpretation",
        inserted[8]: "conclusion",
    }
    for item_id, role in text_roles.items():
        item = next(entry for entry in model(page)["items"] if str(entry["id"]) == item_id)
        set_composition_role(page, item, role, actions)
    open_presets(page, actions)
    builtin = page.locator('#builtinPresetList [data-built-preset="executive"]')
    if builtin.count() and builtin.is_enabled():
        builtin.click()
        settled(page)
        actions.note("apply Executive composition recipe", "editor")
    else:
        actions.note("Executive composition recipe was already applied", "presets")
    finish_regional_revenue_report(page, "blank-rebuild", actions)
    build_actions=actions.summary()
    preflight=compose_and_validate(page,actions)
    screenshots=capture_report(page,output,"blank-rebuild-executive-review",actions)
    pptx_texts=export_pptx(page,output/"blank-rebuild-executive-review.pptx",actions,host)
    export_svg(page,output/"blank-rebuild-executive-review.svg",actions)
    return {
        "report_id": report_id,
        "source": "Blank canvas + visible component library + Data First + Executive recipe",
        "component_count": len(model(page)["items"]),
        "dataset_id": next(entry["dataset_id"] for entry in model(page)["items"] if entry.get("dataset_id")),
        "created_data_visual_id": chart_id,
        "user_actions": build_actions["user_actions"],
        "context_switches": build_actions["context_switches"],
        "modal_openings": build_actions["modal_openings"],
        "manual_component_inserts": sum("insert " in event["action"] for event in build_actions["events"]),
        "manual_field_selections": sum("map " in event["action"].lower() for event in build_actions["events"]),
        "mapping_steps": sum("choose line result" in event["action"] or "commit governed dataset" in event["action"] for event in build_actions["events"]),
        "style_and_layout_steps": sum("composition recipe" in event["action"] for event in build_actions["events"]),
        "manual_composition_setup_steps": sum(
            "set report intent" in event["action"] or "composition recipe" in event["action"]
            for event in build_actions["events"]
        ),
        "repeated_structure_rebuild_actions": sum(
            "insert " in event["action"] or "set report intent" in event["action"]
            for event in build_actions["events"]
        ),
        "preflight":preflight,
        "pptx_slides":len(Presentation(str(output/"blank-rebuild-executive-review.pptx")).slides),
        "pptx_chart_values":pptx_chart_values(output/"blank-rebuild-executive-review.pptx"),
        **screenshots,
    }


def holdout_from_blank(page, host: NativeHost, output: Path, actions: Actions) -> dict:
    report_id = create_from_hub(page, host, "blank", actions)
    chart_id = add_library_component(page, "Line Chart", actions)
    import_or_replace_data(page, DATA["holdout"], "line", actions, chart_id)
    chart = next(entry for entry in model(page)["items"] if entry["id"] == chart_id)
    select_item(page, str(chart["id"]))
    first_chart = next(entry for entry in model(page)["items"] if entry["id"] == chart_id)
    assert first_chart.get("element") == "Line Chart", "The holdout's first trend must be a supported Line Chart."
    dataset = next(entry for entry in model(page)["datasets"] if entry["id"] == first_chart["dataset_id"])
    orders_field = next(field for field in dataset["fields"] if field["name"] == "Orders")
    assert first_chart["mapping"].get("y") == orders_field["id"], "The first holdout trend must preserve the Orders measure."
    set_component_title(page, first_chart, "Orders by month", actions)
    page.locator('[data-inspector="duplicate"]').click()
    settled(page)
    returns_chart = model(page)["items"][-1]
    select_item(page, str(returns_chart["id"]))
    encoding_tab = page.locator('[data-data-dock-tab="encodings"]').last
    if encoding_tab.count() and encoding_tab.get_attribute("aria-selected") != "true":
        encoding_tab.click()
        settled(page)
    returns_mapping = page.locator('[data-dataset-role="y"]')
    dataset = next(entry for entry in model(page)["datasets"] if entry["id"] == returns_chart["dataset_id"])
    returns_field = next(field for field in dataset["fields"] if field["name"] == "Returns")
    if returns_mapping.count():
        returns_mapping.select_option(str(returns_field["id"]))
        settled(page)
    assert model(page)["items"][-1].get("element") == "Line Chart" and model(page)["items"][-1]["mapping"].get("y") == returns_field["id"], f"The holdout report must carry a distinct Returns line through a supported field remap; item={model(page)['items'][-1]!r}, field={returns_field!r}, role_controls={page.locator('[data-dataset-role]').evaluate_all('nodes=>nodes.map(node=>({role:node.dataset.datasetRole,value:node.value}))')}"
    actions.note("map Returns into a second linked trend", "inspector")
    returns_chart = next(entry for entry in model(page)["items"] if entry["id"] == returns_chart["id"])
    set_component_title(page, returns_chart, "Returns by month", actions)
    set_composition_role(page, returns_chart, "supporting_analysis", actions)
    inserted = [add_library_component(page, element, actions) for element in ("Hero KPI", "Hero KPI", "Before/After KPI", "Clean Table", "Risk Callout", "Process Flow", "Hero Title", "Body Narrative", "Body Narrative", "Key Takeaway")]
    added = {"Hero KPI": inserted[0], "Hero KPI secondary": inserted[1], "Before/After KPI": inserted[2], "Clean Table": inserted[3], "Risk Callout": inserted[4], "Process Flow": inserted[5], "Hero Title": inserted[6], "Body Narrative context": inserted[7], "Body Narrative interpretation": inserted[8], "Key Takeaway": inserted[9]}
    headline = next(entry for entry in model(page)["items"] if entry["id"] == added["Hero Title"])
    fill_text(page, headline, "Order volume rose while returns fell across the spring launch.", actions)
    context = next(entry for entry in model(page)["items"] if entry["id"] == added["Body Narrative context"])
    fill_text(page, context, "Monthly operating review · Jan–Apr · Source: order and returns ledger.", actions)
    set_composition_role(page, context, "context", actions)
    interpretation = next(entry for entry in model(page)["items"] if entry["id"] == added["Body Narrative interpretation"])
    fill_text(page, interpretation, TEXT["holdout"], actions)
    conclusion = next(entry for entry in model(page)["items"] if entry["id"] == added["Key Takeaway"])
    fill_text(page, conclusion, "Next step: review the return mix before changing fulfillment capacity.", actions)
    fill_metrics_and_decision(page, "holdout", actions)
    table_item = next(entry for entry in model(page)["items"] if entry.get("engine") == "TableEngine")
    holdout_rows = [["Jan", "1,280", "42"], ["Feb", "1,395", "38"], ["Mar", "1,522", "35"], ["Apr", "1,611", "31"], ["Change", "+331", "−11"]]
    fill_table(page, actions, holdout_rows)
    set_table_headers(page, str(table_item["id"]), ["Month", "Orders", "Returns"], actions)
    set_component_title(page, table_item, "Order and return evidence", actions)
    flow = next(entry for entry in model(page)["items"] if entry.get("engine") == "DiagramEngine")
    select_item(page, str(flow["id"]))
    fill_control(page, "#iTitle", "Review path", actions, "process-flow title")
    fill_control(page, "#iNodes", "Track order growth\nCheck return trend\nReview fulfillment mix\nAdjust capacity", actions, "review steps")
    fill_control(page, "#iEdges", "Track order growth -> Check return trend\nCheck return trend -> Review fulfillment mix\nReview fulfillment mix -> Adjust capacity", actions, "review path")
    set_component_title(page, flow, "Review path", actions)
    open_presets(page, actions)
    builtin = page.locator('#builtinPresetList [data-built-preset="review"]')
    if builtin.count() and builtin.is_enabled():
        builtin.click()
        settled(page)
        actions.note("apply Review composition recipe", "editor")
    record = host.repository.get(report_id)
    assert len(record.model["items"]) >= 11
    assert {entry.get("composition_role") for entry in record.model["items"]} >= {"report_headline", "context", "hero_metric", "primary_analysis", "supporting_analysis", "narrative_interpretation", "detailed_evidence", "causal_evidence", "decision_risk", "conclusion"}
    assert len(record.model["datasets"]) >= 1
    holdout_table = next(entry for entry in record.model["items"] if entry.get("engine") == "TableEngine")
    saved_rows = holdout_table.get("customTable", {}).get("rows", [])
    assert len(saved_rows) == 5 and saved_rows[:4] == [["Jan", 1280, 42], ["Feb", 1395, 38], ["Mar", 1522, 35], ["Apr", 1611, 31]] and saved_rows[4] == ["Change", 331, "−11"], f"The complete monthly evidence table must remain visible and agree with the linked charts; saved={saved_rows!r}."
    preflight = compose_and_validate(page, actions)
    screenshots = capture_report(page, output, "holdout-order-returns", actions)
    pptx_texts = export_pptx(page, output / "holdout-order-returns.pptx", actions, host)
    holdout_pptx_values = pptx_chart_values(output / "holdout-order-returns.pptx")
    assert any("Y Orders" in value for value in screenshots["preview_chart_summaries"]), "The holdout Preview must retain the Orders trend."
    assert any("Y Returns" in value for value in screenshots["preview_chart_summaries"]), "The holdout Preview must show its separately mapped Returns trend."
    assert any(entry["values"] == [1280.0, 1395.0, 1522.0, 1611.0] and entry["categories"] == ["Jan", "Feb", "Mar", "Apr"] for entry in holdout_pptx_values)
    assert any(entry["values"] == [42.0, 38.0, 35.0, 31.0] and entry["categories"] == ["Jan", "Feb", "Mar", "Apr"] for entry in holdout_pptx_values)
    export_svg(page, output / "holdout-order-returns.svg", actions)
    return {
        "report_id": report_id,
        "title": record.title,
        "source": "Blank canvas + Data First + visible component library",
        "items": len(record.model["items"]),
        "sections": sorted({str(entry["section_title"]) for entry in record.model["items"] if entry.get("section_title")}),
        "roles": sorted({str(entry["composition_role"]) for entry in record.model["items"] if entry.get("composition_role")}),
        "data_rows": len(record.model["datasets"][0]["rows"]),
        "evidence_table_rows": saved_rows,
        "pptx_contains_new_data": any(entry["values"] == [1280.0, 1395.0, 1522.0, 1611.0] for entry in holdout_pptx_values),
        "pptx_chart_values": holdout_pptx_values,
        "preflight": preflight,
        **screenshots,
    }


def contact_sheet(paths: list[Path], destination: Path) -> None:
    tiles = []
    font = ImageFont.load_default()
    for path in paths:
        image = Image.open(path).convert("RGB")
        image.thumbnail((720, 460))
        tile = Image.new("RGB", (740, 500), "#f3f5f8")
        tile.paste(image, ((740 - image.width) // 2, 24))
        ImageDraw.Draw(tile).text((12, 7), path.name, font=font, fill="#14243a")
        tiles.append(tile)
    columns = 2
    rows = (len(tiles) + columns - 1) // columns
    sheet = Image.new("RGB", (columns * 740, rows * 500), "white")
    for index, tile in enumerate(tiles):
        sheet.paste(tile, ((index % columns) * 740, (index // columns) * 500))
    sheet.save(destination)


def whole_report_review_pdf(paths: list[Path], destination: Path) -> None:
    """Write a reproducible multi-page PDF overview without adding dependencies."""
    objects: list[bytes] = []
    def add_object(value: bytes) -> int:
        objects.append(value)
        return len(objects)
    catalog_id=add_object(b"<< /Type /Catalog /Pages 2 0 R >>")
    pages_id=add_object(b"")
    font_id=add_object(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    page_ids=[]
    for path in paths:
        with Image.open(path) as source:
            image=source.convert("RGB")
            image.thumbnail((1200,2400),Image.Resampling.LANCZOS)
            width,height=image.size
            from io import BytesIO
            stream=BytesIO();image.save(stream,format="JPEG",quality=82,subsampling=0,optimize=False,progressive=False)
            jpeg=stream.getvalue()
        page_width=612;image_width=540;image_height=image_width*height/width;page_height=image_height+74
        escaped=path.stem.replace("\\","\\\\").replace("(","\\(").replace(")","\\)").encode("ascii","replace")
        header=f"BT /F1 13 Tf 36 {page_height-26:.2f} Td (Visembler whole-report review - ".encode("ascii")
        image_command=f"q {image_width:.2f} 0 0 {image_height:.2f} 36 30 cm /Im0 Do Q\n".encode("ascii")
        content=header+escaped+b") Tj ET\n"+image_command
        image_id=add_object(f"<< /Type /XObject /Subtype /Image /Width {width} /Height {height} /ColorSpace /DeviceRGB /BitsPerComponent 8 /Filter /DCTDecode /Length {len(jpeg)} >>\nstream\n".encode("ascii")+jpeg+b"\nendstream")
        content_id=add_object(f"<< /Length {len(content)} >>\nstream\n".encode("ascii")+content+b"endstream")
        page_id=add_object(f"<< /Type /Page /Parent {pages_id} 0 R /MediaBox [0 0 {page_width} {page_height:.2f}] /Resources << /Font << /F1 {font_id} 0 R >> /XObject << /Im0 {image_id} 0 R >> >> /Contents {content_id} 0 R >>".encode("ascii"))
        page_ids.append(page_id)
    objects[pages_id-1]=f"<< /Type /Pages /Kids [{' '.join(f'{identifier} 0 R' for identifier in page_ids)}] /Count {len(page_ids)} >>".encode("ascii")
    objects[catalog_id-1]=f"<< /Type /Catalog /Pages {pages_id} 0 R >>".encode("ascii")
    payload=bytearray(b"%PDF-1.4\n%CHG-181 deterministic review\n")
    offsets=[0]
    for index,value in enumerate(objects,1):
        offsets.append(len(payload));payload.extend(f"{index} 0 obj\n".encode("ascii"));payload.extend(value);payload.extend(b"\nendobj\n")
    xref=len(payload);payload.extend(f"xref\n0 {len(objects)+1}\n0000000000 65535 f \n".encode("ascii"))
    for offset in offsets[1:]:payload.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    payload.extend(f"trailer\n<< /Size {len(objects)+1} /Root {catalog_id} 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode("ascii"))
    destination.write_bytes(payload)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--only", choices=[name for name, _ in REPORTS] + ["reuse-remap", "holdout-order-returns"], default="", help="Run one report benchmark, the whole reuse journey, or the reserved holdout while iterating on the native flow.")
    args = parser.parse_args()
    output = args.output.expanduser().resolve()
    if output == ROOT or ROOT in output.parents:
        parser.error("--output must be outside the checkout.")
    if output.exists() and any(output.iterdir()):
        parser.error("--output must be new or empty.")
    output.mkdir(parents=True, exist_ok=True)
    report = {
        "request": "CHG-181-r2",
        "operation": "FIX",
        "host": "native NiceGUI 3.15.0",
        "created_at_utc": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()),
        "candidate_sha": __import__("subprocess").check_output(["git", "rev-parse", "HEAD"], cwd=ROOT, text=True).strip(),
        "candidate_tree": __import__("subprocess").check_output(["git", "rev-parse", "HEAD^{tree}"], cwd=ROOT, text=True).strip(),
        "benchmarks": {},
        "reuse_challenge": {},
        "ambiguity_case": {},
        "incompatible_case": {},
        "blank_rebuild_comparison": {},
        "browser_errors": [],
    }
    actions_by_journey: dict[str, dict] = {}
    screenshot_paths: list[Path] = []
    try:
        with tempfile.TemporaryDirectory(prefix="visembler-chg181-native-") as temp_dir, sync_playwright() as playwright:
            with NativeHost(ROOT, Path(temp_dir) / "data") as host:
                browser = playwright.chromium.launch(**browser_kwargs())
                report["browser"] = browser.version
                context = browser.new_context(viewport={"width": 1440, "height": 1000}, accept_downloads=True)
                page = context.new_page()
                page.set_default_timeout(8_000)
                page.on("pageerror", lambda error: report["browser_errors"].append({"kind":"pageerror","message":str(error),"stack":getattr(error,"stack","")}))
                page.on("console", lambda message: report["browser_errors"].append({"kind":"console","message":message.text,"location":message.location}) if message.type == "error" else None)
                page.on("requestfailed", lambda request: report["browser_errors"].append(str(request.failure)))
                selected_reports = [(key, template_id) for key, template_id in REPORTS if not args.only or key == args.only]
                for key, template_id in selected_reports:
                    actions = Actions(key)
                    row = story_template(page, host, output, template_id, actions)
                    report["benchmarks"][key] = row
                    actions_by_journey[key] = actions.summary()
                if not args.only or args.only == "reuse-remap":
                    actions = Actions("reuse-remap")
                    report["reuse_challenge"] = remap_report_challenge(page, host, output, actions)
                    actions_by_journey["reuse-remap"] = actions.summary()
                if not args.only:
                    actions = Actions("ambiguous-remap")
                    report["ambiguity_case"] = remap_ambiguity_challenge(page, host, actions, output)
                    actions_by_journey["ambiguous-remap"] = actions.summary()
                    actions = Actions("incompatible-remap")
                    report["incompatible_case"] = incompatible_challenge(page, host, actions, output)
                    actions_by_journey["incompatible-remap"] = actions.summary()
                    actions = Actions("blank-rebuild")
                    report["blank_rebuild_comparison"] = blank_rebuild_comparison(page, host, output, actions)
                    actions_by_journey["blank-rebuild"] = actions.summary()
                if not args.only or args.only == "holdout-order-returns":
                    actions = Actions("holdout-order-returns")
                    report["holdout"] = holdout_from_blank(page, host, output, actions)
                    actions_by_journey["holdout"] = actions.summary()
                import importlib.metadata
                report["native_runtime"] = {
                    "nicegui_version": importlib.metadata.version("nicegui"),
                    "base_url": host.url,
                    "ephemeral_port": host.port,
                }
                report["report_count"] = len(host.repository.list())
                context.close()
                browser.close()
    except Exception as error:
        report["harness_error"] = str(error) or type(error).__name__
        report["traceback"] = traceback.format_exc()
    report["journey_action_logs"] = actions_by_journey
    report["journey_action_log_sha256"] = {
        name: hashlib.sha256(json.dumps(value, sort_keys=True).encode()).hexdigest()
        for name, value in actions_by_journey.items()
    }
    screenshot_paths = sorted(output.glob("*.png"))
    if screenshot_paths:
        contact_sheet(screenshot_paths, output / "whole-report-contact-sheet.png")
    review_images=[]
    for prefix in ("executive-business-review","semiconductor-rca","experiment-decision","technical-status-review","holdout-order-returns","reuse-source-report","reuse-remapped-report","reuse-remapped-section","blank-rebuild-executive-review"):
        for suffix in ("-preview-desktop.png","-preview-mobile.png"):
            path=output/f"{prefix}{suffix}"
            if path.is_file():review_images.append(path)
    if review_images:whole_report_review_pdf(review_images,output/"whole-report-review.pdf")
    artifacts = []
    for path in sorted(output.iterdir()):
        if path.is_file() and path.name not in {"chg181-native-acceptance.json", "chg181-native-acceptance.json.sha256"}:
            artifacts.append({
                "path": path.name,
                "bytes": path.stat().st_size,
                "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            })
    report["artifacts"] = artifacts
    passed = bool(
        len(report["benchmarks"]) == (1 if args.only in {name for name, _ in REPORTS} else 4 if not args.only else 0)
        and (bool(args.only) or report.get("holdout", {}).get("items", 0) >= 8)
        and (args.only != "holdout-order-returns" or report.get("holdout", {}).get("items", 0) >= 8)
        and (args.only not in {"", "reuse-remap"} or report.get("reuse_challenge", {}).get("undo_redo_reload_history"))
        and (bool(args.only) or report.get("ambiguity_case", {}).get("blocked_before_choice"))
        and (bool(args.only) or report.get("incompatible_case", {}).get("blocked"))
        and not report["browser_errors"]
        and "harness_error" not in report
    )
    report["status"] = "PASS" if passed else "FAIL"
    target = output / "chg181-native-acceptance.json"
    write_json(target, report)
    target.with_suffix(target.suffix + ".sha256").write_text(
        hashlib.sha256(target.read_bytes()).hexdigest() + "  " + target.name + "\n",
        encoding="utf-8",
    )
    print(json.dumps({"status": report["status"], "benchmarks": list(report["benchmarks"]), "holdout": report.get("holdout", {}).get("items"), "harness_error": report.get("harness_error")}, indent=2))
    return 0 if report["status"] == "PASS" else 1


if __name__ == "__main__":
    raise SystemExit(main())
